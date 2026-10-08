"""Screen stills (T3): the backend half of a capture, and the Captures list.

- a region is validated (empty or absurd sides refused; negative offsets
  taken, since a monitor left of or above the primary has them) and the
  gdigrab argv draws the pointer only when asked;
- a freeze photographs every monitor into one set, replacing the last, and
  ``/frozen/{index}`` serves it or answers 404; the startup clears a set a
  crash left behind, best-effort (a frame held open is left, and the app
  still starts);
- on this machine, with a display, a real gdigrab still through the route is
  a PNG of exactly the asked size;
- the Captures list: take, list (filtered by owner), image, download named
  after the capture, delete; somebody else's is refused;
- Add to deck: the still becomes the deck's LAST slide - the inner project
  gains a slide with empty notes, the .pptx gains a picture slide, the
  slide's image is letterboxed to the deck's slide image size (1920x1080
  when none is rendered: bars, never a stretch) - and a PDF, a busy deck
  and somebody else's deck or capture are refused.
"""

from __future__ import annotations

import io
import os
import threading
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation

import api.routers.capture as capture_router
from api import store as auth_store
from services import capture, captures, jobs, recordings, slides
from services import projects as store
from utils.config import FFMPEG_PATH, config

PASSWORD = "Owner-pass-12345"
LOOPBACK = ("127.0.0.1", 50000)
has_display = pytest.mark.skipif(os.name != "nt" or not FFMPEG_PATH, reason="a Windows desktop and ffmpeg are needed")


@pytest.fixture
def rig(tmp_path, monkeypatch):
    monkeypatch.setattr(store, "PROJECTS_DIR", tmp_path / "projects")
    monkeypatch.setattr(captures, "CAPTURES_DIR", tmp_path / "captures")
    monkeypatch.setattr(captures, "TEMP_DIR", tmp_path / "temp")
    monkeypatch.setattr(capture_router, "SCREEN_DIR", tmp_path / "temp" / "capture")
    monkeypatch.setattr(recordings, "RECORDINGS_DIR", tmp_path / "temp" / "recordings")
    monkeypatch.setattr(config, "_config", {})
    monkeypatch.setattr(config, "save", lambda: None)
    monkeypatch.setattr(auth_store, "DATA_DIR", tmp_path / "auth")
    monkeypatch.setattr(auth_store, "DB_PATH", tmp_path / "auth" / "media_studio.db")
    from api.app import app

    # The desktop shell's window loads the backend at 127.0.0.1: these clients are THIS computer, which
    # the host-screen routes require (TestClient's default address, "testclient", is not).
    with TestClient(app, client=LOOPBACK) as admin:
        seeded = auth_store.authenticate("admin", "admin")
        auth_store.change_password(seeded["id"], "admin", must_change=False)
        assert admin.post("/api/auth/login", json={"username": "admin", "password": "admin"}).status_code == 200
        auth_store.create_user("other", PASSWORD, "Otto Other", role="editor", must_change_password=False)
        other = TestClient(app, client=LOOPBACK)
        assert other.post("/api/auth/login", json={"username": "other", "password": PASSWORD}).status_code == 200
        yield {"admin": admin, "other": other, "tmp": tmp_path, "app": app}


def _fake_grab(colour=(200, 30, 30)):
    """A stand-in for gdigrab: a PNG of the asked size, so these tests need
    no display. The real grab has its own test below."""

    def grab(x, y, width, height, *, cursor, out_dir):
        capture.validate_region(x, y, width, height)
        Path(out_dir).mkdir(parents=True, exist_ok=True)
        out = Path(out_dir) / f"still-fake-{width}x{height}.png"
        Image.new("RGB", (width, height), colour).save(out)
        return out

    return grab


def _deck_bytes(slide_count: int = 2) -> bytes:
    prs = Presentation()
    for i in range(slide_count):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"Slide {i + 1}"
        slide.placeholders[1].text = f"Bullet {i + 1}"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# -- the grab -------------------------------------------------------------------------

def test_a_region_is_validated_and_negative_offsets_are_legal():
    capture.validate_region(-1600, -120, 1600, 900)
    with pytest.raises(ValueError):
        capture.validate_region(0, 0, 0, 10)
    with pytest.raises(ValueError):
        capture.validate_region(0, 0, 10, capture.MAX_DIMENSION + 1)
    with pytest.raises(ValueError):
        capture.validate_region(capture.MAX_DIMENSION + 1, 0, 10, 10)


def test_the_gdigrab_argv_draws_the_pointer_only_when_asked(tmp_path):
    with_cursor = capture.still_command("FF", -1280, 240, 1280, 1024, True, tmp_path / "a.png")
    without = capture.still_command("FF", 0, 0, 10, 10, False, tmp_path / "b.png")
    assert with_cursor[:1] == ["FF"] and "-f" in with_cursor and with_cursor[with_cursor.index("-f") + 1] == "gdigrab"
    assert with_cursor[with_cursor.index("-draw_mouse") + 1] == "1"
    assert without[without.index("-draw_mouse") + 1] == "0"
    assert with_cursor[with_cursor.index("-offset_x") + 1] == "-1280"
    assert with_cursor[with_cursor.index("-video_size") + 1] == "1280x1024"
    assert with_cursor[with_cursor.index("-frames:v") + 1] == "1" and with_cursor[-1] == str(tmp_path / "a.png")


def test_a_freeze_replaces_the_previous_set_and_frozen_serves_it_or_404s(rig, monkeypatch):
    monkeypatch.setattr(capture, "_grab", lambda x, y, w, h, *, cursor, out: (Image.new("RGB", (w, h)).save(out), out)[1])
    c = rig["admin"]
    assert c.get("/api/capture/frozen/0").status_code == 404
    monitors = [{"index": 0, "x": 0, "y": 0, "width": 40, "height": 30}, {"index": 2, "x": 2560, "y": -60, "width": 20, "height": 10}]
    assert c.post("/api/capture/freeze", json={"monitors": monitors}).json() == {"frozen": [0, 2]}
    r = c.get("/api/capture/frozen/2")
    assert r.status_code == 200 and r.headers["content-type"] == "image/png" and Image.open(io.BytesIO(r.content)).size == (20, 10)
    assert c.get("/api/capture/frozen/1").status_code == 404
    # The next freeze names only monitor 1: the old set is gone.
    assert c.post("/api/capture/freeze", json={"monitors": [{"index": 1, "x": -1280, "y": 240, "width": 8, "height": 8}]}).json() == {"frozen": [1]}
    assert c.get("/api/capture/frozen/2").status_code == 404
    assert c.get("/api/capture/frozen/1").status_code == 200
    assert c.post("/api/capture/freeze", json={"monitors": []}).status_code == 422
    # Once it has served, the set is deleted: no picture of the desktop is kept that nobody asked to keep.
    assert c.delete("/api/capture/frozen").status_code == 204
    assert c.get("/api/capture/frozen/1").status_code == 404
    assert not list((rig["tmp"] / "temp" / "capture").glob("frozen-*.png"))
    assert c.delete("/api/capture/frozen").status_code == 204  # idempotent


def test_frozen_frames_a_crash_left_behind_are_cleared_at_startup(rig):
    """A crash between the freeze and the overlay's answer leaves full-resolution pictures of every
    monitor in data\\temp\\capture; the next start deletes them, and nothing else there."""
    screen_dir = rig["tmp"] / "temp" / "capture"
    screen_dir.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (4, 4)).save(screen_dir / "frozen-0.png")
    Image.new("RGB", (4, 4)).save(screen_dir / "frozen-3.png")
    (screen_dir / "unrelated.txt").write_text("kept", encoding="utf-8")
    with TestClient(rig["app"], client=LOOPBACK):
        pass
    assert sorted(p.name for p in screen_dir.iterdir()) == ["unrelated.txt"]


@pytest.mark.skipif(os.name != "nt", reason="an open file refuses a delete on Windows only")
def test_a_frozen_frame_held_open_does_not_stop_the_startup(rig):
    """A frame held open without FILE_SHARE_DELETE (a viewer, a scanner) cannot be deleted (WinError 32):
    the clear logs it and carries on with the others, and the app still starts."""
    screen_dir = rig["tmp"] / "temp" / "capture"
    screen_dir.mkdir(parents=True, exist_ok=True)
    for i in (0, 1, 2):
        Image.new("RGB", (4, 4)).save(screen_dir / f"frozen-{i}.png")
    with open(screen_dir / "frozen-1.png", "rb"):
        assert capture.clear_frozen(screen_dir) == 2
        assert sorted(p.name for p in screen_dir.iterdir()) == ["frozen-1.png"]
        with TestClient(rig["app"], client=LOOPBACK) as again:
            assert again.get("/api/system/health").status_code == 200
    # Released, it goes at the next clear.
    assert capture.clear_frozen(screen_dir) == 1 and not any(screen_dir.iterdir())


def test_there_is_no_route_that_grabs_the_screen_on_a_get(routes):
    """The page never needed one (stills are POST /api/captures, audited); it is gone (review MINOR 4).
    What is left under /api/capture is the overlay's frozen set."""
    assert sorted(f"{r.method} {r.path}" for r in routes if r.path.startswith("/api/capture/")) == [
        "DELETE /api/capture/frozen", "GET /api/capture/frozen/{index}", "POST /api/capture/freeze"]


@has_display
def test_a_real_still_through_the_route_is_a_png_of_the_asked_size(rig):
    """A real gdigrab still - 24 x 16 pixels at the virtual screen's origin, kept only for the length of
    the test - through the route the page uses."""
    c = rig["admin"]
    r = c.post("/api/captures", json={"x": 0, "y": 0, "width": 24, "height": 16, "cursor": False})
    assert r.status_code == 200, r.text
    cid = r.json()["id"]
    try:
        img = c.get(f"/api/captures/{cid}/image")
        assert img.status_code == 200 and img.headers["content-type"] == "image/png"
        assert Image.open(io.BytesIO(img.content)).size == (24, 16)
    finally:
        assert c.delete(f"/api/captures/{cid}").status_code == 204
    assert c.post("/api/captures", json={"x": 0, "y": 0, "width": 0, "height": 16}).status_code in (400, 422)


# -- the Captures list ----------------------------------------------------------------

def test_take_list_image_download_and_delete(rig, monkeypatch):
    monkeypatch.setattr(captures.capture, "take_still", _fake_grab())
    admin, other = rig["admin"], rig["other"]
    r = admin.post("/api/captures", json={"x": -10, "y": 5, "width": 32, "height": 24, "cursor": True, "kind": "window", "name": "Notepad / notes.txt"})
    assert r.status_code == 200, r.text
    rec = r.json()
    assert rec["kind"] == "window" and rec["cursor"] is True and (rec["width"], rec["height"]) == (32, 24)
    assert rec["name"] == "Notepad / notes.txt" and rec["owner_name"] and rec["size_bytes"] > 0
    cid = rec["id"]
    assert [x["id"] for x in admin.get("/api/captures").json()["captures"]] == [cid]
    img = admin.get(f"/api/captures/{cid}/image")
    assert img.status_code == 200 and Image.open(io.BytesIO(img.content)).size == (32, 24)
    dl = admin.get(f"/api/captures/{cid}/download")
    assert dl.status_code == 200 and 'filename="Notepad-notes.txt.png"' in dl.headers["content-disposition"]
    # Not the other editor's: invisible and refused.
    assert other.get("/api/captures").json()["captures"] == []
    assert other.get(f"/api/captures/{cid}/image").status_code == 403
    assert other.delete(f"/api/captures/{cid}").status_code == 403
    assert admin.delete(f"/api/captures/{cid}").status_code == 204
    assert admin.get(f"/api/captures/{cid}/image").status_code == 404
    assert admin.get("/api/captures").json()["captures"] == []


def test_a_capture_without_a_name_is_named_after_the_moment(rig, monkeypatch):
    monkeypatch.setattr(captures.capture, "take_still", _fake_grab())
    rec = rig["admin"].post("/api/captures", json={"x": 0, "y": 0, "width": 4, "height": 4}).json()
    assert rec["name"].startswith("Screen capture 20") and rec["kind"] == "region" and rec["cursor"] is False


# -- letterbox and Add to deck ------------------------------------------------------

def test_letterbox_keeps_the_aspect_and_fills_the_rest_with_the_background():
    square = Image.new("RGB", (400, 400), (0, 0, 255))
    out = slides.letterbox(square, (1920, 1080))
    assert out.size == (1920, 1080)
    assert out.getpixel((10, 540)) == (255, 255, 255)        # a side bar
    assert out.getpixel((960, 540)) == (0, 0, 255)           # the picture, centred
    assert out.getpixel((960 - 540 + 2, 540)) == (0, 0, 255) and out.getpixel((960 - 540 - 2, 540)) == (255, 255, 255)
    wide = Image.new("RGB", (3840, 1080), (0, 255, 0))
    out2 = slides.letterbox(wide, (1920, 1080))
    assert out2.getpixel((960, 10)) == (255, 255, 255) and out2.getpixel((960, 540)) == (0, 255, 0)  # bars above and below
    exact = Image.new("RGB", (192, 108), (9, 9, 9))
    assert slides.letterbox(exact, (1920, 1080)).getpixel((0, 0)) == (9, 9, 9)       # same aspect: fills the frame


def test_add_to_deck_makes_the_still_the_last_slide(rig, monkeypatch):
    monkeypatch.setattr(captures.capture, "take_still", _fake_grab((0, 0, 255)))
    c = rig["admin"]
    deck = c.post("/api/projects/import", files={"file": ("deck.pptx", _deck_bytes(2), "application/octet-stream")}).json()
    cid = c.post("/api/captures", json={"x": 0, "y": 0, "width": 400, "height": 400, "name": "Square"}).json()["id"]

    r = c.post(f"/api/captures/{cid}/add-to-deck", json={"project_id": deck["id"]})
    assert r.status_code == 200, r.text
    assert r.json() == {"project_id": deck["id"], "index": 2, "slide_count": 3}

    listed = c.get(f"/api/projects/{deck['id']}/slides").json()["slides"]
    assert [s["index"] for s in listed] == [0, 1, 2]
    assert listed[2]["has_image"] is True and listed[2]["speaker_notes"] == "" and listed[1]["has_image"] is False
    assert c.get(f"/api/projects/{deck['id']}").json()["slide_count"] == 3
    # The slide's image: the deck's size (nothing rendered -> 1920x1080), the still letterboxed.
    img = c.get(f"/api/projects/{deck['id']}/slides/2/image")
    assert img.status_code == 200
    rendered = Image.open(io.BytesIO(img.content))
    assert rendered.size == (1920, 1080)
    assert rendered.getpixel((960, 540)) == (0, 0, 255) and rendered.getpixel((10, 540)) == (255, 255, 255)
    # The .pptx gained a picture slide at the end.
    source = rig["tmp"] / "projects" / deck["id"] / "deck.pptx"
    prs = Presentation(str(source))
    assert len(prs.slides) == 3
    last = prs.slides[2]
    assert [sh.shape_type for sh in last.shapes] and any(sh.shape_type == 13 for sh in last.shapes)  # 13 = PICTURE
    # The capture itself is still there.
    assert c.get(f"/api/captures/{cid}/image").status_code == 200
    # Twice: a fourth slide, after the third.
    assert c.post(f"/api/captures/{cid}/add-to-deck", json={"project_id": deck["id"]}).json()["index"] == 3


def test_add_to_deck_refuses_a_pdf_a_busy_deck_and_what_is_not_yours(rig, monkeypatch):
    monkeypatch.setattr(captures.capture, "take_still", _fake_grab())
    admin, other = rig["admin"], rig["other"]
    cid = admin.post("/api/captures", json={"x": 0, "y": 0, "width": 8, "height": 8}).json()["id"]
    pdf = store.import_upload("pages.pdf", b"%PDF-1.4 not really", owner_id=auth_store.authenticate("admin", "admin")["id"])
    assert admin.post(f"/api/captures/{cid}/add-to-deck", json={"project_id": pdf["id"]}).status_code == 400
    deck = admin.post("/api/projects/import", files={"file": ("deck.pptx", _deck_bytes(1), "application/octet-stream")}).json()
    # A job holding the deck: refused with 409 until it ends.
    gate = threading.Event()
    jobs.start("render-slides", lambda progress: gate.wait(10), project_id=deck["id"])
    try:
        assert admin.post(f"/api/captures/{cid}/add-to-deck", json={"project_id": deck["id"]}).status_code == 409
    finally:
        gate.set()
    # The other editor owns neither the capture nor the deck.
    assert other.post(f"/api/captures/{cid}/add-to-deck", json={"project_id": deck["id"]}).status_code == 403
    theirs = other.post("/api/captures", json={"x": 0, "y": 0, "width": 8, "height": 8}).json()["id"]
    assert other.post(f"/api/captures/{theirs}/add-to-deck", json={"project_id": deck["id"]}).status_code == 403
    assert admin.post("/api/captures/000000000000/add-to-deck", json={"project_id": deck["id"]}).status_code == 404


# -- the host's screen is this computer's only (S1) ---------------------------------

#: Every route that touches the HOST's screen, with a body that passes validation, so a refusal is the
#: guard's and not FastAPI's 422.
HOST_SCREEN_ROUTES = [
    ("POST", "/api/capture/freeze", {"monitors": [{"index": 0, "x": 0, "y": 0, "width": 8, "height": 8}]}),
    ("GET", "/api/capture/frozen/0", None),
    ("DELETE", "/api/capture/frozen", None),
    ("POST", "/api/captures", {"x": 0, "y": 0, "width": 8, "height": 8}),
    ("POST", "/api/recordings", {"name": "x"}),
    ("PUT", "/api/recordings/000000000000/chunks/0", b"chunk"),
    ("POST", "/api/recordings/000000000000/heartbeat", None),
    ("POST", "/api/recordings/000000000000/release", None),
]


def _signed_in_from(app, client: tuple[str, int], headers: dict | None = None) -> TestClient:
    c = TestClient(app, client=client, headers=headers or {})
    assert c.post("/api/auth/login", json={"username": "admin", "password": "admin"}).status_code == 200
    return c


def _call(c: TestClient, method: str, path: str, body):
    if isinstance(body, bytes):
        return c.request(method, path, content=body)
    return c.request(method, path, json=body) if body is not None else c.request(method, path)


@pytest.mark.parametrize("method, path, body", HOST_SCREEN_ROUTES, ids=[f"{m} {p.split('?')[0]}" for m, p, _ in HOST_SCREEN_ROUTES])
def test_a_client_on_another_machine_may_not_touch_the_host_screen(rig, monkeypatch, method, path, body):
    """Run as a team server (--host 0.0.0.0), the host's desktop is the SERVER's: an editor elsewhere is
    refused with one plain line, before anything is grabbed, and the refusal is audited."""
    grabs = []
    monkeypatch.setattr(capture, "take_still", lambda *a, **k: grabs.append(a) or _fake_grab()(*a, **k))
    monkeypatch.setattr(capture, "_grab", lambda *a, **k: grabs.append(a))
    remote = _signed_in_from(rig["app"], ("10.20.30.40", 50123))
    r = _call(remote, method, path, body)
    assert r.status_code == 403, r.text
    assert r.json()["detail"] == "Screen capture works only in the desktop app on this computer."
    assert grabs == []
    assert recordings.list_unfinished() == []  # no recording was started either
    refused = auth_store.list_audit(action="capture.refused")
    assert refused and refused[0]["detail"] == f"{method} {path.split('?')[0]} from 10.20.30.40"


def test_the_captures_that_do_not_touch_the_screen_still_answer_another_machine(rig, monkeypatch):
    monkeypatch.setattr(captures.capture, "take_still", _fake_grab())
    cid = rig["admin"].post("/api/captures", json={"x": 0, "y": 0, "width": 8, "height": 8}).json()["id"]
    remote = _signed_in_from(rig["app"], ("10.20.30.40", 50123))
    assert [c["id"] for c in remote.get("/api/captures").json()["captures"]] == [cid]
    assert remote.get(f"/api/captures/{cid}/image").status_code == 200
    assert remote.get(f"/api/captures/{cid}/download").status_code == 200
    assert remote.delete(f"/api/captures/{cid}").status_code == 204


@pytest.mark.parametrize("client", [("127.0.0.1", 1), ("127.0.0.2", 1), ("::1", 1), ("::ffff:127.0.0.1", 1)],
                         ids=["127.0.0.1", "127.0.0.2", "::1", "ipv4-mapped"])
def test_every_loopback_spelling_is_this_computer(rig, client):
    c = _signed_in_from(rig["app"], client)
    assert c.get("/api/capture/frozen/0").status_code == 404  # allowed in: there is just no freeze yet


@pytest.mark.parametrize("header", ["X-Forwarded-For", "Forwarded", "X-Real-IP", "CF-Connecting-IP"])
def test_a_loopback_request_that_came_through_a_proxy_is_not_this_computer(rig, header):
    """A tunnel or a reverse proxy on the server connects from 127.0.0.1; the header it adds is what
    says the browser is elsewhere. The shell's own window never sends one."""
    value = "for=203.0.113.9" if header == "Forwarded" else "203.0.113.9"
    c = _signed_in_from(rig["app"], LOOPBACK, headers={header: value})
    assert c.get("/api/capture/frozen/0").status_code == 403


# -- the deck's own background (R2) ----------------------------------------------------

def test_border_colour_is_the_median_of_the_edge():
    img = Image.new("RGB", (400, 200), (30, 40, 50))
    img.paste((255, 255, 255), (100, 50, 300, 150))        # content in the middle, off the edge
    img.paste((250, 0, 0), (0, 0, 60, 4))                  # a logo touching the top edge
    assert slides.border_colour(img) == (30, 40, 50)
    assert slides.border_colour(Image.new("RGB", (3, 3), (9, 8, 7))) == (9, 8, 7)


def test_add_to_deck_letterboxes_onto_the_decks_own_background(rig, monkeypatch):
    """A dark deck: the bars are the first slide's background, not white, and the image takes that
    slide's size."""
    monkeypatch.setattr(captures.capture, "take_still", _fake_grab((0, 0, 255)))
    c = rig["admin"]
    deck = c.post("/api/projects/import", files={"file": ("dark.pptx", _deck_bytes(1), "application/octet-stream")}).json()
    pm = slides.manager(deck["id"])
    first = Image.new("RGB", (1280, 720), (24, 28, 36))
    first.paste((240, 240, 240), (200, 150, 1080, 570))    # the slide's content
    first.save(pm.images_dir / "slide_001.png")
    cid = c.post("/api/captures", json={"x": 0, "y": 0, "width": 400, "height": 400}).json()["id"]
    assert c.post(f"/api/captures/{cid}/add-to-deck", json={"project_id": deck["id"]}).json()["index"] == 1
    added = Image.open(io.BytesIO(c.get(f"/api/projects/{deck['id']}/slides/1/image").content)).convert("RGB")
    assert added.size == (1280, 720)
    assert added.getpixel((20, 360)) == (24, 28, 36) and added.getpixel((1260, 360)) == (24, 28, 36)
    assert added.getpixel((640, 360)) == (0, 0, 255)
