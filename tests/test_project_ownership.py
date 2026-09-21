"""Project ownership (porting vertical 4a): a project belongs to whoever
imported it, and only they (or an admin) may see or touch it.

The rule lives in ONE place - ``api.deps.may_access_project`` /
``require_project`` - so these tests exercise it through the routes that use
it: the project routes, the slide editor and the AI assistant. The sweep in
``test_another_editor_is_refused_every_project_scoped_route`` is enumerated by
hand (each route needs a valid body or FastAPI answers 422 before the endpoint
runs) and a companion test fails when the app grows a project-scoped route the
sweep does not name, so the coverage cannot rot.

Ownership is stored in ``project.json``, not in SQLite: a project stays a
directory you can zip.
"""

import io

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from api import deps
from api import store as auth_store
from services import jobs, processing, projects as store, slides
from utils.config import config

PASSWORD = "Owner-pass-12345"

# Bodies for the 403 sweep: FastAPI validates the body BEFORE the endpoint runs,
# so a route with a required field needs one here or it answers 422 instead of
# the 403 under test. A GET takes None.
PROJECT_SCOPED_ROUTES: dict[tuple[str, str], dict | None] = {
    ("GET", "/api/projects/{pid}"): None,
    ("DELETE", "/api/projects/{pid}"): None,
    ("GET", "/api/projects/{pid}/video"): None,
    ("GET", "/api/projects/{pid}/revoiced-video"): None,
    ("GET", "/api/projects/{pid}/outputs/{kind}"): None,
    ("GET", "/api/projects/{pid}/tracks/{kind}"): None,
    ("POST", "/api/projects/{pid}/transcribe"): {},
    ("PATCH", "/api/projects/{pid}/transcript"): {"transcript": []},
    ("PATCH", "/api/projects/{pid}/transcript/{index}"): {},
    ("GET", "/api/projects/{pid}/transcript/{index}/preview"): None,
    ("GET", "/api/projects/{pid}/transcript/download"): None,
    ("GET", "/api/projects/{pid}/narration/plan"): None,
    ("PATCH", "/api/projects/{pid}/narration/offsets"): {"offsets": [{"index": 0, "offset": 0.1}]},
    ("GET", "/api/projects/{pid}/waveform"): None,
    ("GET", "/api/projects/{pid}/edit"): None,
    ("PUT", "/api/projects/{pid}/edit"): {"keep": [[0.0, 1.0]]},
    ("DELETE", "/api/projects/{pid}/edit"): None,
    ("POST", "/api/projects/{pid}/generate"): {},
    ("POST", "/api/projects/{pid}/revoice"): {},
    ("GET", "/api/projects/{pid}/slides"): None,
    ("PATCH", "/api/projects/{pid}/slides"): {"slides": []},
    ("POST", "/api/projects/{pid}/slides/render"): {},
    ("PATCH", "/api/projects/{pid}/slides/{index}"): {},
    ("POST", "/api/projects/{pid}/slides/{index}/undo"): {},
    ("POST", "/api/projects/{pid}/slides/{index}/reset"): {},
    ("GET", "/api/projects/{pid}/slides/{index}/image"): None,
    ("GET", "/api/projects/{pid}/export/pptx"): None,
    ("GET", "/api/projects/{pid}/export/qa"): None,
    ("GET", "/api/projects/{pid}/ai/status"): None,
    ("POST", "/api/projects/{pid}/ai/notes"): {},
    ("POST", "/api/projects/{pid}/ai/enhance"): {},
    ("POST", "/api/projects/{pid}/ai/qa"): {},
    ("POST", "/api/projects/{pid}/ai/tone"): {"tone": "Technical"},
    ("POST", "/api/projects/{pid}/ai/translate"): {"language": "Spanish"},
    ("POST", "/api/projects/{pid}/ai/pacing"): {},
    ("POST", "/api/projects/{pid}/ai/qa-doc"): {},
    ("POST", "/api/projects/{pid}/ai/analyze"): {},
    ("POST", "/api/projects/{pid}/slides/{index}/ai/enhance"): {},
    ("POST", "/api/projects/{pid}/slides/{index}/ai/qa-fix"): {"criterion": "grammar", "issue": "x"},
}

# Routes that carry no ``{pid}`` ON PURPOSE and so are outside the sweep: the
# music library (api/routers/music.py) is STUDIO-WIDE - shared by every account
# like the voices and the studio settings (spec §12.3) - so every signed-in
# user may list, upload and delete, and the audit row says who. They need a
# session (401 signed out, held by tests/test_music_library.py) but have no
# owner to check. Listed here so the decision is visible beside the table, and
# so the guard below fails if one of them ever grows a project id without
# joining the table.
STUDIO_WIDE_ROUTES: set[tuple[str, str]] = {
    ("GET", "/api/music"),
    ("POST", "/api/music"),
    ("GET", "/api/music/{name}"),
    ("GET", "/api/music/{name}/peaks"),
    ("DELETE", "/api/music/{name}"),
}


@pytest.fixture(autouse=True)
def isolated(tmp_path, monkeypatch):
    monkeypatch.setattr(store, "PROJECTS_DIR", tmp_path / "projects")
    monkeypatch.setattr(auth_store, "DATA_DIR", tmp_path / "auth")
    monkeypatch.setattr(auth_store, "DB_PATH", tmp_path / "auth" / "media_studio.db")
    monkeypatch.setattr(config, "_config", {})
    monkeypatch.setattr(config, "save", lambda: None)
    monkeypatch.setattr(slides, "_deck_cache", {})
    monkeypatch.setattr(slides, "_locks", {})
    # The outer project.json's own per-pid locks (the transcript Save and the
    # narration editor both take them).
    monkeypatch.setattr(store, "_locks", {})
    monkeypatch.setattr(processing, "TEMP_DIR", tmp_path / "temp")
    # The log-once notice is a module-level flag; reset it so a test can assert
    # it fires (and so test order cannot decide whether it does).
    monkeypatch.setattr(deps, "_legacy_notice_logged", False)


@pytest.fixture
def app():
    from api.app import app as fastapi_app

    with TestClient(fastapi_app):
        seeded = auth_store.authenticate("admin", "admin")
        auth_store.change_password(seeded["id"], "admin", must_change=False)
        yield fastapi_app


def _signed_in(app, username: str, password: str) -> TestClient:
    c = TestClient(app)
    r = c.post("/api/auth/login", json={"username": username, "password": password})
    assert r.status_code == 200, r.text
    return c


@pytest.fixture
def cast(app):
    """The admin, the owner (an editor) and another editor, each with their own
    cookie jar, plus the two editors' account rows."""
    owner = auth_store.create_user("owner", PASSWORD, "Olive Owner", role="editor", must_change_password=False)
    other = auth_store.create_user("other", PASSWORD, "Otto Other", role="editor", must_change_password=False)
    return {
        "admin": _signed_in(app, "admin", "admin"),
        "owner": _signed_in(app, "owner", PASSWORD),
        "other": _signed_in(app, "other", PASSWORD),
        "owner_account": owner,
        "other_account": other,
    }


def _deck_bytes(slide_count: int = 2) -> bytes:
    prs = Presentation()
    for i in range(slide_count):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"Slide {i + 1}"
        slide.placeholders[1].text = f"Bullet {i + 1}"
        slide.notes_slide.notes_text_frame.text = f"Notes {i + 1}."
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _import_as(client: TestClient, name: str = "deck.pptx", data: bytes | None = None) -> dict:
    r = client.post("/api/projects/import", files={"file": (name, data or _deck_bytes(), "application/octet-stream")})
    assert r.status_code == 200, r.text
    return r.json()


def _ids(client: TestClient) -> set[str]:
    r = client.get("/api/projects")
    assert r.status_code == 200, r.text
    return {p["id"] for p in r.json()["projects"]}


def _call(client: TestClient, method: str, path: str, body):
    return client.request(method, path, json=body) if body is not None else client.request(method, path)


def _fill(path: str, pid: str) -> str:
    return path.replace("{pid}", pid).replace("{index}", "0").replace("{kind}", "srt")


# ── the record ────────────────────────────────────────────────────────────────

def test_import_records_the_owner_id_and_the_display_name(cast):
    record = _import_as(cast["owner"])
    assert record["owner_id"] == cast["owner_account"]["id"]
    # Denormalised on purpose: the project still says who made it after the
    # account is gone.
    assert record["owner_name"] == "Olive Owner"
    assert store.get_project(record["id"])["owner_name"] == "Olive Owner"


def test_the_owner_name_survives_the_account_being_deactivated(cast):
    record = _import_as(cast["owner"])
    r = cast["admin"].patch(f"/api/users/{cast['owner_account']['id']}", json={"is_active": False})
    assert r.status_code == 200
    assert cast["admin"].get(f"/api/projects/{record['id']}").json()["owner_name"] == "Olive Owner"


# ── the list ──────────────────────────────────────────────────────────────────

def test_the_list_shows_an_editor_only_their_own_and_an_admin_everything(cast):
    mine = _import_as(cast["owner"], "mine.pptx")
    theirs = _import_as(cast["other"], "theirs.pptx")
    admins = _import_as(cast["admin"], "admins.pptx")

    assert _ids(cast["owner"]) == {mine["id"]}
    assert _ids(cast["other"]) == {theirs["id"]}
    assert _ids(cast["admin"]) == {mine["id"], theirs["id"], admins["id"]}
    # The store itself is unfiltered - the filtering is the API's, in one place.
    assert len(store.list_projects()) == 3


# ── legacy records ────────────────────────────────────────────────────────────

def test_a_project_with_no_owner_is_admin_owned_never_public(cast, caplog):
    """The named trap: a record imported before ownership existed must not
    become visible to every editor."""
    import logging

    legacy = store.import_upload("legacy.pptx", _deck_bytes())
    assert legacy["owner_id"] is None

    caplog.set_level(logging.INFO)
    handler_logger = logging.getLogger("mediastudio.AUTH")
    handler_logger.addHandler(caplog.handler)
    try:
        assert legacy["id"] not in _ids(cast["owner"])
        assert legacy["id"] not in _ids(cast["other"])
        assert cast["owner"].get(f"/api/projects/{legacy['id']}").status_code == 403
        assert cast["other"].get(f"/api/projects/{legacy['id']}").status_code == 403
    finally:
        handler_logger.removeHandler(caplog.handler)

    # The admin has it, and can delete it.
    assert legacy["id"] in _ids(cast["admin"])
    assert cast["admin"].get(f"/api/projects/{legacy['id']}").status_code == 200
    assert cast["admin"].delete(f"/api/projects/{legacy['id']}").status_code == 204

    # And the operator was told once - not once per record, not once per request.
    notices = [r for r in caplog.records if "no owner_id" in r.getMessage()]
    assert len(notices) == 1, [r.getMessage() for r in notices]


# ── enforcement per route class ───────────────────────────────────────────────

def test_the_owner_may_use_every_route_class_and_so_may_an_admin(cast):
    pid = _import_as(cast["owner"])["id"]

    for who in ("owner", "admin"):
        client = cast[who]
        # Project routes.
        assert client.get(f"/api/projects/{pid}").status_code == 200, who
        # The slide editor (a real deck, so this is the route working, not just passing the guard).
        r = client.get(f"/api/projects/{pid}/slides")
        assert r.status_code == 200 and len(r.json()["slides"]) == 2, who
        r = client.patch(f"/api/projects/{pid}/slides/0", json={"speaker_notes": f"Edited by {who}."})
        assert r.status_code == 200 and r.json()["speaker_notes"] == f"Edited by {who}.", who
        # The AI assistant: the rules-only pacing call writes without the model.
        assert client.post(f"/api/projects/{pid}/ai/pacing", json={}).status_code == 200, who

    # And the owner can delete their own project.
    assert cast["owner"].delete(f"/api/projects/{pid}").status_code == 204


def test_another_editor_is_refused_every_project_scoped_route(cast):
    pid = _import_as(cast["owner"])["id"]
    other = cast["other"]

    refused = []
    for (method, template), body in PROJECT_SCOPED_ROUTES.items():
        r = _call(other, method, _fill(template, pid), body)
        if r.status_code != 403 or "admin or the project's owner" not in r.json().get("detail", ""):
            refused.append(f"{method} {template} -> {r.status_code} {r.text[:120]}")
    assert refused == [], "these routes did not refuse another editor with a 403:\n  " + "\n  ".join(refused)

    # Nothing was done: the project is untouched and still the owner's.
    assert store.get_project(pid)["owner_id"] == cast["owner_account"]["id"]
    assert jobs.active_for(pid) is None


def test_a_missing_project_is_still_a_404_for_everyone(cast):
    absent = "aabbccddeeff"  # a valid id shape that does not exist
    for who in ("owner", "other", "admin"):
        assert cast[who].get(f"/api/projects/{absent}").status_code == 404, who
        assert cast[who].delete(f"/api/projects/{absent}").status_code == 404, who


def test_the_sweep_covers_every_project_scoped_route_the_app_has(routes):
    """An anti-rot guard on the guard: a new ``{pid}`` route must be added to
    PROJECT_SCOPED_ROUTES above, or this fails."""
    live = {(r.method, r.path) for r in routes if "{pid}" in r.path}
    missing = sorted(live - set(PROJECT_SCOPED_ROUTES))
    stale = sorted(set(PROJECT_SCOPED_ROUTES) - live)
    assert not missing, (
        "these project-scoped routes are not in PROJECT_SCOPED_ROUTES (tests/test_project_ownership.py), "
        f"so nothing checks that they refuse another editor: {missing}"
    )
    assert not stale, f"PROJECT_SCOPED_ROUTES names routes the app no longer has: {stale}"


def test_the_studio_wide_routes_are_live_and_carry_no_project(routes):
    """The exclusion list must rot too: each route it names exists, and none
    of them takes a project id (the moment one did, it would belong in the
    table above, and the guard above would say so)."""
    live = {(r.method, r.path) for r in routes}
    missing = sorted(STUDIO_WIDE_ROUTES - live)
    assert not missing, f"STUDIO_WIDE_ROUTES names routes the app does not have: {missing}"
    assert not any("{pid}" in path for _, path in STUDIO_WIDE_ROUTES)


def test_the_services_a_job_runs_never_re_check_ownership():
    """A job runs on a worker thread with no request and no user, and the route
    that queued it has already authorised the caller. A service that checked
    again would either crash or need a fake user - so the check belongs at the
    route and nowhere else."""
    from pathlib import Path

    root = Path(__file__).resolve().parent.parent
    for name in ("ai_slides", "revoice", "slides", "transcription"):
        source = (root / "services" / f"{name}.py").read_text(encoding="utf-8")
        assert "require_project" not in source, f"services/{name}.py must not authorise; the route already did"
        assert "may_access_project" not in source, f"services/{name}.py must not authorise; the route already did"
