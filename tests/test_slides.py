"""The slide editor (porting vertical 3a): ``services/slides.py`` over the
engine's inner project, its routes, the export lock shared with the generate
pipeline, and the engine's own ``ProjectManager`` fixes (every field saved,
atomic writes, an unreadable file never recreated, no migration on load).

A real small deck is built with python-pptx in every test (three slides, notes
on two) and a real two-page PDF with PyMuPDF, so materialisation, the deck's
title / body text and the .pptx export are exercised for real. The slide
IMAGE export is faked: PowerPoint COM never runs under pytest. The store, the
config and the auth DB are isolated to a tmp dir.
"""

import io
import json
import threading
import time
import types
from dataclasses import dataclass, fields
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pydantic import ValidationError

from api.schemas import SlidesBulkUpdate, SlideUpdate
from core import project_manager as engine
from core.pptx_exporter import ExportedSlide
from core.project_manager import ProjectManager, ProjectStateError
from services import file_item as file_item_module
from services import jobs, processing, slides
from services import projects as store
from utils.config import config

NOTES = ("Welcome to the deck.", "", "Third slide notes.")

# The keys the hand-kept save() of the predecessor wrote, per project and per slide.
LEGACY_PROJECT_KEYS = {
    "pptx_path", "project_dir", "created_at", "last_modified", "voice_id", "output_video_path",
    "transition_pause", "background_music_path", "music_volume", "generation_speed",
    "generation_stability", "generation_similarity_boost", "generation_style", "slides",
    "slide_order", "source_video_path", "revoice_sync_mode",
}
LEGACY_SLIDE_KEYS = {
    "index", "speaker_notes", "audio_path", "image_path", "video_path", "audio_duration", "voice_id",
    "needs_regeneration", "ai_enhanced", "pause_override", "has_animation", "voice_override",
    "alt_text", "notes_history", "original_start_time", "original_end_time", "original_segments",
}


@pytest.fixture(autouse=True)
def isolated(tmp_path, monkeypatch):
    """A throwaway project store, an in-memory config (Edge is the studio's
    provider unless a test says otherwise) and a throwaway auth DB."""
    monkeypatch.setattr(store, "PROJECTS_DIR", tmp_path / "projects")
    monkeypatch.setattr(config, "_config", {})
    monkeypatch.setattr(config, "save", lambda: None)
    monkeypatch.setattr(slides, "_deck_cache", {})
    monkeypatch.setattr(slides, "_locks", {})
    monkeypatch.setattr(processing, "TEMP_DIR", tmp_path / "temp")

    from api import store as auth_store

    monkeypatch.setattr(auth_store, "DATA_DIR", tmp_path / "auth")
    monkeypatch.setattr(auth_store, "DB_PATH", tmp_path / "auth" / "media_studio.db")


@pytest.fixture
def client():
    """A logged-in TestClient (the seeded admin, past its first password change)."""
    from api import store as auth_store
    from api.app import app

    with TestClient(app) as c:
        seeded = auth_store.authenticate("admin", "admin")
        auth_store.change_password(seeded["id"], "admin", must_change=False)
        r = c.post("/api/auth/login", json={"username": "admin", "password": "admin"})
        assert r.status_code == 200
        yield c


def _deck_bytes(notes=NOTES) -> bytes:
    prs = Presentation()
    for i, note in enumerate(notes):
        slide = prs.slides.add_slide(prs.slide_layouts[1])  # title + content
        slide.shapes.title.text = f"Slide {i + 1} title"
        slide.placeholders[1].text = f"Bullet {i + 1}"
        if note:
            slide.notes_slide.notes_text_frame.text = note
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _import_deck(name="deck.pptx") -> str:
    return store.import_upload(name, _deck_bytes())["id"]


def _pdf_bytes(pages=2) -> bytes:
    import fitz

    doc = fitz.open()
    for i in range(pages):
        doc.new_page().insert_text((72, 72), f"Page {i + 1}")
    data = doc.tobytes()
    doc.close()
    return data


def _import_pdf(name="pages.pdf") -> str:
    return store.import_upload(name, _pdf_bytes(2))["id"]


def _png_bytes() -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (4, 4), (200, 30, 30)).save(buf, "PNG")
    return buf.getvalue()


PNG = _png_bytes()


def _inner(pid: str, stem: str = "deck") -> Path:
    return store.PROJECTS_DIR / pid / f"{stem}_project" / "project.json"


class _FakeExporter:
    """Stands in for core.pptx_exporter.PPTXExporter: writes one PNG per deck
    slide (slide 2 "animated", with its clip), records whether the export
    lock was held and when it ran, and reports the backend a test asks for.
    ``entered`` is set on entry and the export blocks on ``gate`` when one is
    given, so tests order threads on events rather than sleeps."""

    backend_to_report = "powerpoint"
    runs: list = []  # (deck name, lock held, start, end)
    entered: threading.Event | None = None
    gate: threading.Event | None = None

    def __init__(self, pptx_path, output_dir):
        self.pptx_path = Path(pptx_path)
        self.output_dir = Path(output_dir)
        self.backend = None

    def export_slides_as_images(self, progress_callback=None):
        start = time.monotonic()
        held = slides.SLIDE_EXPORT_LOCK.locked()
        cls = type(self)
        if cls.entered is not None:
            cls.entered.set()
        if cls.gate is not None:
            assert cls.gate.wait(10), "the test never opened the gate"
        count = len(Presentation(str(self.pptx_path)).slides)
        out = []
        for i in range(count):
            image = self.output_dir / f"slide_{i + 1:03d}.png"
            image.write_bytes(PNG)
            video = None
            if i == 1:  # an animated slide comes with its clip, as PowerPoint exports it
                video = self.output_dir / "animations" / f"slide_{i + 1:03d}_anim.mp4"
                video.parent.mkdir(exist_ok=True)
                video.write_bytes(b"mp4")
            out.append(ExportedSlide(index=i, image_path=image, video_path=video, has_animation=video is not None))
            if progress_callback:
                progress_callback(i + 1, count)
        self.backend = cls.backend_to_report
        cls.runs.append((self.pptx_path.name, held, start, time.monotonic()))
        return out


def _fake_exporter(monkeypatch, backend="powerpoint", gate=None, entered=None):
    import core.pptx_exporter as exporter_module

    _FakeExporter.backend_to_report = backend
    _FakeExporter.runs = []
    _FakeExporter.gate = gate
    _FakeExporter.entered = entered
    monkeypatch.setattr(exporter_module, "PPTXExporter", _FakeExporter)


def _wait_job(client, job_id, timeout=10.0):
    deadline = time.time() + timeout
    while time.time() < deadline:
        job = client.get(f"/api/jobs/{job_id}").json()
        if job["status"] in ("done", "error"):
            return job
        time.sleep(0.02)
    raise AssertionError("job did not finish in time")


def _wait_until(condition, timeout=10.0, what="the condition"):
    """Poll ``condition`` until it holds (an observable event, not a fixed sleep)."""
    deadline = time.time() + timeout
    while time.time() < deadline:
        if condition():
            return
        time.sleep(0.01)
    raise AssertionError(f"timed out waiting for {what}")


class _BlockingJob:
    """A job attached to a project that runs until the test releases it."""

    def __init__(self, pid, kind="generate"):
        self.started, self.release = threading.Event(), threading.Event()

        def work(progress):
            self.started.set()
            self.release.wait(10)
            return "ok"

        self.id = jobs.submit(kind, work, project_id=pid)
        assert self.started.wait(5)


# -- materialisation and listing -------------------------------------------------

def test_list_slides_materialises_the_inner_project_from_the_deck_without_images():
    pid = _import_deck()
    record = store.get_project(pid)
    record["slide_count"] = None  # as an import whose reader failed would leave it
    store.save_project(record)
    assert not _inner(pid).exists(), "nothing before the first read"

    listed = slides.list_slides(pid)

    assert _inner(pid).is_file(), "the engine's project.json, where generate would have put it"
    assert listed[0] == {
        "index": 0, "speaker_notes": "Welcome to the deck.", "original_notes": "Welcome to the deck.",
        "title": "Slide 1 title", "body_text": "Bullet 1",
        "has_image": False, "has_audio": False, "ai_enhanced": False, "needs_regeneration": True,
        "voice_override": None, "pause_override": None, "alt_text": None,
        "notes_history_depth": 0, "has_animation": False,
    }
    assert [s["speaker_notes"] for s in listed] == list(NOTES)
    assert [s["title"] for s in listed] == ["Slide 1 title", "Slide 2 title", "Slide 3 title"]
    assert list((store.PROJECTS_DIR / pid / "deck_project" / "images").iterdir()) == [], "no image export"

    saved = store.get_project(pid)
    assert saved["slide_count"] == 3, "reconciled from the engine project"
    assert saved["engine_project_dir"] == "deck_project", "the directory name, never an absolute path"

    for slide in listed:
        for value in slide.values():
            if isinstance(value, str):
                assert not Path(value).is_absolute() and "_project" not in value, value

    # A second read opens the same project (no recreation, the edits below survive),
    # and a plain read never rewrites the file.
    slides.update_slide(pid, 2, speaker_notes="Kept.")
    stamp = json.loads(_inner(pid).read_text())["last_modified"]
    assert slides.list_slides(pid)[2]["speaker_notes"] == "Kept."
    slides.list_slides(pid)
    assert json.loads(_inner(pid).read_text())["last_modified"] == stamp


def test_slides_payload_reports_readiness_and_the_image_source():
    pid = _import_deck()
    payload = slides.slides_payload(pid)
    assert payload["slides_ready"] is False
    assert payload["images_source"] is None and payload["images_rendered_at"] is None
    assert len(payload["slides"]) == 3
    assert slides.images_ready(pid) is False

    record = store.get_project(pid)
    record["images_source"] = "libreoffice"  # a value this edition does not know
    store.save_project(record)
    assert slides.slides_payload(pid)["images_source"] is None, "unknown is None, not a stray string"


def test_manager_refuses_a_video_project_and_an_unknown_one():
    video = store.import_upload("clip.mp4", b"v")["id"]
    with pytest.raises(ValueError):
        slides.manager(video)
    with pytest.raises(slides.ProjectNotFound):
        slides.manager("aabbccddeeff")
    with pytest.raises(slides.ProjectNotFound):
        slides.list_slides("../../etc")
    assert issubclass(slides.ProjectNotFound, LookupError)

    # A malformed record is an error of its own, never "not found".
    pid = _import_deck()
    record = store.get_project(pid)
    del record["source_filename"]
    store.save_project(record)
    with pytest.raises(KeyError):
        slides.list_slides(pid)


def test_deleting_a_project_forgets_its_lock_and_deck_cache(client):
    pid = _import_deck()
    slides.list_slides(pid)
    assert pid in slides._locks and pid in slides._deck_cache
    assert client.delete(f"/api/projects/{pid}").status_code == 204
    assert pid not in slides._locks and pid not in slides._deck_cache
    assert store.delete_project("aabbccddeeff") is False


# -- editing ---------------------------------------------------------------------------

def test_update_undo_reset_round_trip_through_the_inner_project_json():
    pid = _import_deck()

    updated = slides.update_slide(
        pid, 0, speaker_notes="Edited notes.", pause_override=2.5,
        voice_override="en-GB-SoniaNeural", alt_text=" A chart ",
    )
    assert updated["speaker_notes"] == "Edited notes." and updated["original_notes"] == "Welcome to the deck."
    assert updated["needs_regeneration"] is True and updated["notes_history_depth"] == 1
    assert updated["pause_override"] == 2.5 and updated["voice_override"] == "en-GB-SoniaNeural"
    assert updated["alt_text"] == "A chart"

    on_disk = json.loads(_inner(pid).read_text())["slides"][0]
    assert on_disk["speaker_notes"] == "Edited notes."
    assert on_disk["notes_history"] == ["Welcome to the deck."], "the manager's own undo history"
    assert on_disk["pause_override"] == 2.5 and on_disk["voice_override"] == "en-GB-SoniaNeural"
    assert on_disk["alt_text"] == "A chart" and on_disk["needs_regeneration"] is True

    # None clears an override; a keyword left out is left alone.
    cleared = slides.update_slide(pid, 0, pause_override=None, voice_override=None, alt_text=None)
    assert (cleared["pause_override"], cleared["voice_override"], cleared["alt_text"]) == (None, None, None)
    assert cleared["speaker_notes"] == "Edited notes." and cleared["notes_history_depth"] == 1
    assert slides.update_slide(pid, 0, speaker_notes="Edited notes.")["notes_history_depth"] == 1, "no change, no history entry"

    undone = slides.undo_slide(pid, 0)
    assert undone["speaker_notes"] == "Welcome to the deck." and undone["notes_history_depth"] == 0
    assert slides.undo_slide(pid, 0) is None, "nothing left to undo"

    slides.update_slide(pid, 0, speaker_notes="Second edit.")
    reset = slides.reset_slide(pid, 0)
    assert reset["speaker_notes"] == "Welcome to the deck."
    assert reset["notes_history_depth"] == 2, "reset is an edit: the text it replaced can be undone"
    assert slides.undo_slide(pid, 0)["speaker_notes"] == "Second edit."


def test_bulk_update_saves_every_slide_or_none():
    pid = _import_deck()
    out = slides.bulk_update(pid, [{"index": 0, "speaker_notes": "A"}, {"index": 2, "speaker_notes": "C"}])
    assert [s["speaker_notes"] for s in out] == ["A", "", "C"]
    assert [s["notes_history_depth"] for s in out] == [1, 0, 1]

    with pytest.raises(ValueError):
        slides.bulk_update(pid, [{"index": 1, "speaker_notes": "B"}, {"index": 9, "speaker_notes": "X"}])
    assert [s["speaker_notes"] for s in slides.list_slides(pid)] == ["A", "", "C"], "the bad batch wrote nothing"


def test_update_slide_refuses_bad_values_without_writing():
    pid = _import_deck()
    slides.list_slides(pid)
    before = _inner(pid).read_text()

    for kwargs in (
        {"speaker_notes": "x" * (slides.MAX_NOTES_CHARS + 1)},
        {"speaker_notes": 42},
        {"pause_override": 31}, {"pause_override": -1}, {"pause_override": "soon"}, {"pause_override": True},
        {"voice_override": "af_heart"},  # a Kokoro id while the studio narrates with Edge
        {"voice_override": 5},
        {"alt_text": "x" * 2001},
    ):
        with pytest.raises(ValueError):
            slides.update_slide(pid, 0, **kwargs)
    for index in (3, -1, "0", True):
        with pytest.raises(ValueError):
            slides.update_slide(pid, index, speaker_notes="x")
        with pytest.raises(ValueError):
            slides.undo_slide(pid, index)
        with pytest.raises(ValueError):
            slides.reset_slide(pid, index)
    assert _inner(pid).read_text() == before

    # The override is checked against the provider it is for, like generate's voice.
    assert slides.update_slide(pid, 0, voice_override="af_heart", provider="kokoro")["voice_override"] == "af_heart"
    with pytest.raises(ValueError):
        slides.update_slide(pid, 0, voice_override="en-US-AriaNeural", provider="kokoro")
    with pytest.raises(ValueError):
        slides.update_slide(pid, 0, voice_override="af_heart", provider="polly")
    config._config["tts_provider"] = "kokoro"  # Settings › Studio: no provider named = the studio's
    assert slides.update_slide(pid, 0, voice_override="am_adam")["voice_override"] == "am_adam"
    with pytest.raises(ValueError):
        slides.update_slide(pid, 0, voice_override="en-US-AriaNeural")
    assert slides.update_slide(pid, 0, voice_override="")["voice_override"] is None, '"" clears it too'


def test_a_new_voice_override_marks_the_slide_for_regeneration():
    """The engine's own change detection knows nothing about overrides, so a
    changed override must set the flag or the next generate keeps the old audio."""
    pid = _import_deck()
    pm = slides.manager(pid)
    pm.state.slides[0].needs_regeneration = False
    pm.save()
    assert slides.update_slide(pid, 0, voice_override="en-GB-RyanNeural")["needs_regeneration"] is True


def _narrated(pid: str, voice="en-US-AriaNeural") -> ProjectManager:
    """The project after a generate: cached audio on every slide and the
    generation settings saved, so nothing needs regenerating."""
    pm = slides.manager(pid)
    for slide in pm.state.slides:
        audio = pm.audio_dir / f"slide_{slide.index + 1:03d}_audio.mp3"
        audio.write_bytes(b"mp3")
        pm.update_slide_audio(slide.index, audio, 2.0, voice)
    pm.update_generation_settings(speed=1.0, voice_id=voice)
    return pm


def _needing(pm: ProjectManager, voice="en-US-AriaNeural") -> list:
    return pm.get_slides_needing_regeneration(
        current_speed=1.0, current_voice_id=voice, current_stability=0.5,
        current_similarity_boost=0.75, current_style=0.0,
    )


def test_edits_stay_flagged_for_regeneration_across_loads():
    """load() used to clear needs_regeneration on any slide with cached audio
    (a predecessor migration that assumed one load per session); here it runs
    on every request, so an edit was un-flagged before generate could see it."""
    pid = _import_deck()
    source = store.PROJECTS_DIR / pid / "deck.pptx"
    assert _needing(_narrated(pid)) == []

    slides.update_slide(pid, 1, speaker_notes="Edited after the first generate.")

    # What the generate route does at the start of every job.
    item = file_item_module.FileItem(source, projects_base=store.PROJECTS_DIR / pid)
    assert item.load() is True
    assert _needing(item.project_manager) == [1], "the edited slide is re-narrated"
    assert slides.list_slides(pid)[1]["needs_regeneration"] is True, "and a GET does not un-flag it"

    slides.undo_slide(pid, 1)
    slides.update_slide(pid, 2, voice_override="en-GB-RyanNeural")
    reloaded = ProjectManager(_inner(pid).parent)
    reloaded.load()
    assert _needing(reloaded) == [1, 2], "an undo and a voice override are edits too"


# -- the engine: atomic saves, unreadable files, every field -------------------------

def test_save_writes_a_temp_file_and_replaces_project_json_in_one_step(tmp_path, monkeypatch):
    pm = ProjectManager(tmp_path / "p")
    pm.create_project(tmp_path / "deck.pptx", ["a"], voice_id="v")
    before = pm.project_file.read_text()

    # A write that fails halfway leaves the previous file intact and no temp file behind.
    def broken_dump(data, f, **kwargs):
        f.write('{"partial": ')
        raise OSError("disk full")

    with monkeypatch.context() as m:
        m.setattr(engine.json, "dump", broken_dump)
        with pytest.raises(OSError):
            pm.update_slide_notes(0, "lost")
    assert pm.project_file.read_text() == before
    assert sorted(p.name for p in pm.project_dir.iterdir()) == ["audio", "images", "project.json"]

    # The write itself goes through os.replace of a same-directory temp file.
    replaced = []
    real_replace = engine.os.replace

    def spy(src, dst):
        replaced.append((Path(src), Path(dst), Path(src).read_text()))
        real_replace(src, dst)

    monkeypatch.setattr(engine.os, "replace", spy)
    pm.update_slide_notes(0, "kept")
    assert len(replaced) == 1
    src, dst, content = replaced[0]
    assert dst == pm.project_file and src.parent == pm.project_dir and src != dst
    assert json.loads(content)["slides"][0]["speaker_notes"] == "kept", "complete before it replaces the file"
    assert sorted(p.name for p in pm.project_dir.iterdir()) == ["audio", "images", "project.json"]


def test_replace_retries_while_a_reader_holds_the_file(tmp_path, monkeypatch):
    attempts = []
    real_replace = engine.os.replace

    def flaky(src, dst):
        attempts.append(1)
        if len(attempts) < 3:
            raise PermissionError("in use")
        real_replace(src, dst)

    monkeypatch.setattr(engine.os, "replace", flaky)
    monkeypatch.setattr(engine.time, "sleep", lambda s: None)
    pm = ProjectManager(tmp_path / "p")
    pm.create_project(tmp_path / "deck.pptx", ["a"], voice_id="v")
    assert len(attempts) == 3 and pm.project_file.is_file()


@pytest.mark.parametrize(
    "content",
    ["{not json", "[]", '{"slides": "nope"}', '{"pptx_path": "x", "slides": [1, 2]}', '{"slides": []}'],
)
def test_load_raises_for_an_unreadable_project_instead_of_returning_none(tmp_path, content):
    """Bad JSON, the wrong shape, a non-dict slide entry, a required field
    missing: all ProjectStateError - never None (which every caller took as
    "no project yet" and answered by recreating it from the deck)."""
    pm = ProjectManager(tmp_path / "p")
    pm.project_file.write_text(content)
    with pytest.raises(ProjectStateError) as exc:
        pm.load()
    assert "could not be read" in str(exc.value) and "Nothing was changed" in str(exc.value)
    assert str(tmp_path) not in str(exc.value), "no absolute path in a message the user sees"
    assert pm.project_file.read_text() == content
    assert ProjectManager(tmp_path / "absent").load() is None, "no file at all is still None"


def test_an_unreadable_deck_project_is_never_recreated(client):
    pid = _import_deck()
    slides.update_slide(pid, 0, speaker_notes="Precious edit.")
    _inner(pid).write_text("{not json")

    with pytest.raises(ProjectStateError):
        slides.list_slides(pid)
    assert _inner(pid).read_text() == "{not json", "left for the user to restore"

    r = client.get(f"/api/projects/{pid}/slides")
    assert r.status_code == 409 and "could not be read" in r.json()["detail"]
    assert client.patch(f"/api/projects/{pid}/slides/0", json={"speaker_notes": "x"}).status_code == 409
    assert client.post(f"/api/projects/{pid}/slides/render").status_code == 409
    assert client.get(f"/api/projects/{pid}/slides/0/image").status_code == 409
    assert client.get(f"/api/projects/{pid}/export/pptx").status_code == 409

    # The generate path (FileItem.load at the start of the job, then the
    # pipeline's own load) fails the job with that message instead of
    # starting the project over from the deck.
    r = client.post(f"/api/projects/{pid}/generate", json={})
    job = _wait_job(client, r.json()["job_id"])
    assert job["status"] == "error" and "could not be read" in job["error"]
    assert _inner(pid).read_text() == "{not json"

    item = types.SimpleNamespace(path=store.PROJECTS_DIR / pid / "deck.pptx", _projects_base=store.PROJECTS_DIR / pid,
                                 project_manager=None, reader=None, slide_count=3)
    with pytest.raises(ProjectStateError):
        processing.VideoProcessor()._get_or_create_project(item)
    assert _inner(pid).read_text() == "{not json"


def test_an_unreadable_pdf_project_is_never_recreated(client):
    pid = _import_pdf()
    slides.update_slide(pid, 0, speaker_notes="Narrate page one.")
    inner = _inner(pid, "pages")
    inner.write_text('{"pptx_path": "x", "slides": [1]}')

    with pytest.raises(ProjectStateError):
        slides.slides_payload(pid)
    item = file_item_module.FileItem(store.PROJECTS_DIR / pid / "pages.pdf", projects_base=store.PROJECTS_DIR / pid)
    with pytest.raises(ProjectStateError):
        item.load_pdf()
    assert inner.read_text() == '{"pptx_path": "x", "slides": [1]}'

    assert client.get(f"/api/projects/{pid}/slides").status_code == 409
    r = client.post(f"/api/projects/{pid}/generate", json={})
    job = _wait_job(client, r.json()["job_id"])
    assert job["status"] == "error" and "could not be read" in job["error"]
    assert inner.read_text() == '{"pptx_path": "x", "slides": [1]}'


def test_project_manager_round_trips_a_field_it_did_not_know_when_save_was_written(tmp_path, monkeypatch):
    """save() used to hand-list the keys, so a field added to the dataclasses
    was silently dropped. It now serialises every field: a wider state (what a
    later vertical adds) round-trips without touching save()/load()."""

    @dataclass
    class WiderState(engine.ProjectState):
        narrator_style: str = "calm"

    @dataclass
    class WiderSlide(engine.SlideRenderState):
        qa_score: int | None = None

    monkeypatch.setattr(engine, "ProjectState", WiderState)
    monkeypatch.setattr(engine, "SlideRenderState", WiderSlide)

    pm = engine.ProjectManager(tmp_path / "wider")
    pm.create_project(tmp_path / "deck.pptx", ["a", "b"], voice_id="v")
    pm.state.narrator_style = "brisk"
    pm.state.slides[1].qa_score = 8
    pm.update_slide_notes(0, "edited a")

    data = json.loads(pm.project_file.read_text())
    assert data["narrator_style"] == "brisk" and data["slides"][1]["qa_score"] == 8
    assert data["slides"][0]["notes_history"] == ["a"]

    state = engine.ProjectManager(tmp_path / "wider").load()
    assert state.narrator_style == "brisk"
    assert state.slides[1].qa_score == 8 and state.slides[0].speaker_notes == "edited a"


def test_project_manager_keeps_todays_on_disk_shape_and_reads_old_and_newer_files(tmp_path):
    pm = ProjectManager(tmp_path / "p")
    pm.create_project(tmp_path / "deck.pptx", ["a"], voice_id="v")
    data = json.loads(pm.project_file.read_text())
    assert set(data) == LEGACY_PROJECT_KEYS, "the 17 keys the hand-kept list wrote, no more, no less"
    assert set(data["slides"][0]) == LEGACY_SLIDE_KEYS
    assert LEGACY_PROJECT_KEYS == {f.name for f in fields(engine.ProjectState)}, "one key per field"
    assert LEGACY_SLIDE_KEYS == {f.name for f in fields(engine.SlideRenderState)}

    # A file written by a newer edition (unknown keys) still loads instead of
    # becoming unreadable (which generate would answer by recreating it).
    data["future_setting"] = 1
    data["slides"][0]["vision_tags"] = ["chart"]
    pm.project_file.write_text(json.dumps(data))
    state = ProjectManager(tmp_path / "p").load()
    assert state is not None and state.slides[0].speaker_notes == "a"
    assert not hasattr(state, "future_setting")


# -- images ------------------------------------------------------------------------

def test_ensure_images_exports_once_and_records_the_source(monkeypatch):
    pid = _import_deck()
    _fake_exporter(monkeypatch, backend="pillow")
    messages = []

    result = slides.ensure_images(pid, lambda fraction, message: messages.append((fraction, message)))

    assert result == {"cached": False, "images_source": "pillow", "slides": 3}
    payload = slides.slides_payload(pid)
    assert payload["slides_ready"] is True and payload["images_source"] == "pillow"
    assert payload["images_rendered_at"]
    assert all(s["has_image"] for s in payload["slides"])
    assert [s["has_animation"] for s in payload["slides"]] == [False, True, False]
    assert store.get_project(pid)["images_source"] == "pillow", "3b disables vision on the fallback"
    assert messages[-1] == (1.0, "Slide previews ready")
    assert any("Exporting slide 2/3" in message for _, message in messages)
    assert slides.manager(pid).all_images_ready(), "the engine's own check agrees (generate will not re-export)"

    runs = len(_FakeExporter.runs)
    assert slides.ensure_images(pid) == {"cached": True, "images_source": "pillow"}
    assert len(_FakeExporter.runs) == runs, "idempotent: nothing exported again"

    other = _import_deck("other.pptx")
    _fake_exporter(monkeypatch, backend="powerpoint")
    assert slides.ensure_images(other)["images_source"] == "powerpoint"
    assert store.get_project(other)["images_source"] == "powerpoint"


def test_ensure_images_survives_a_notes_edit_during_the_export(monkeypatch):
    """The export runs outside the project lock (minutes for a big deck); the
    image paths must land on the project as it is afterwards, not on the copy
    read before the export."""
    pid = _import_deck()
    _fake_exporter(monkeypatch)

    class _EditingExporter(_FakeExporter):
        def export_slides_as_images(self, progress_callback=None):
            slides.update_slide(pid, 1, speaker_notes="Typed while rendering.")
            return super().export_slides_as_images(progress_callback)

    import core.pptx_exporter as exporter_module

    monkeypatch.setattr(exporter_module, "PPTXExporter", _EditingExporter)
    slides.ensure_images(pid)
    listed = slides.list_slides(pid)
    assert listed[1]["speaker_notes"] == "Typed while rendering." and all(s["has_image"] for s in listed)


def test_concurrent_renders_queue_on_the_export_lock_and_say_so(monkeypatch):
    first, second = _import_deck("first.pptx"), _import_deck("second.pptx")
    entered, gate = threading.Event(), threading.Event()
    _fake_exporter(monkeypatch, entered=entered, gate=gate)
    logs = {first: [], second: []}
    results = {}

    def run(pid):
        results[pid] = slides.ensure_images(pid, lambda fraction, message: logs[pid].append(message))

    a = threading.Thread(target=run, args=(first,))
    b = threading.Thread(target=run, args=(second,))
    a.start()
    assert entered.wait(5), "a is inside its export, holding the lock"
    b.start()
    _wait_until(lambda: slides.WAITING_FOR_EXPORT in logs[second], what="b's waiting message")
    assert _FakeExporter.runs == [], "b has not exported while a holds the lock"
    gate.set()
    a.join(timeout=10)
    b.join(timeout=10)

    runs = sorted(_FakeExporter.runs, key=lambda r: r[2])
    assert [r[0] for r in runs] == ["first.pptx", "second.pptx"]
    assert all(r[1] for r in runs), "the lock was held during every export"
    assert runs[0][3] <= runs[1][2], "the exports never overlapped"
    assert logs[second][0] == slides.WAITING_FOR_EXPORT, "the queued job says why it waits"
    assert slides.WAITING_FOR_EXPORT not in logs[first]
    assert results[first]["cached"] is False and results[second]["cached"] is False
    assert not slides.SLIDE_EXPORT_LOCK.locked()


def test_ensure_images_checks_again_once_it_holds_the_export_lock(monkeypatch):
    """Two renders of one deck queued behind each other: the second finds the
    images the first just wrote and exports nothing."""
    pid = _import_deck()
    _fake_exporter(monkeypatch)
    messages = []
    result = {}

    slides.SLIDE_EXPORT_LOCK.acquire()  # "the first render" is exporting
    try:
        worker = threading.Thread(target=lambda: result.update(slides.ensure_images(pid, lambda f, m: messages.append(m))))
        worker.start()
        _wait_until(lambda: slides.WAITING_FOR_EXPORT in messages, what="the waiting message")
        # ... and finishes: the images land on disk while the second render waits.
        images = store.PROJECTS_DIR / pid / "deck_project" / "images"
        for i in range(3):
            (images / f"slide_{i + 1:03d}.png").write_bytes(PNG)
    finally:
        slides.SLIDE_EXPORT_LOCK.release()
    worker.join(timeout=10)
    assert result["cached"] is True
    assert _FakeExporter.runs == [], "nothing exported a second time"


def test_image_path_is_derived_from_the_index_and_kept_under_the_project(monkeypatch):
    pid = _import_deck()
    assert slides.image_path(pid, 0) is None, "not rendered yet"
    _fake_exporter(monkeypatch)
    slides.ensure_images(pid)

    expected = (store.PROJECTS_DIR / pid / "deck_project" / "images" / "slide_001.png").resolve()
    assert slides.image_path(pid, 0) == expected
    assert slides.image_path(pid, 2).name == "slide_003.png"
    assert slides.image_path(pid, 3) is None and slides.image_path(pid, -1) is None

    # A tampered project.json (an absolute path elsewhere) changes nothing: the
    # file is derived from the project directory and the index.
    secret = store.PROJECTS_DIR.parent / "secret.png"
    secret.write_bytes(PNG)
    pm = slides.manager(pid)
    pm.state.slides[0].image_path = str(secret)
    pm.save()
    assert slides.image_path(pid, 0) == expected

    # And a derivation that landed outside the project would be refused.
    monkeypatch.setattr(slides, "_image_file", lambda record, pm, index: secret)
    assert slides.image_path(pid, 0) is None


def test_moved_images_are_adopted_on_open(monkeypatch):
    """Files present but project.json pointing at another machine's paths: any
    open fixes the stored paths, so generate does not export again."""
    pid = _import_deck()
    _fake_exporter(monkeypatch)
    slides.ensure_images(pid)
    pm = slides.manager(pid)
    for slide in pm.state.slides:
        slide.image_path = "D:/elsewhere/" + Path(slide.image_path).name
    pm.save()

    assert slides.manager(pid).all_images_ready(), "adopted by the open itself"
    assert slides.ensure_images(pid)["cached"] is True
    assert len(_FakeExporter.runs) == 1


# -- PDFs ----------------------------------------------------------------------------

def test_pdf_projects_come_with_their_page_images_and_keep_their_edits():
    pid = _import_pdf()
    assert store.get_project(pid)["slide_count"] is None, "the import counts decks only"

    payload = slides.slides_payload(pid)
    assert payload["slides_ready"] is True, "a PDF's images are rendered as part of reading it"
    assert payload["images_source"] == "pdf"
    assert [s["index"] for s in payload["slides"]] == [0, 1]
    assert payload["slides"][0]["original_notes"] == "" and payload["slides"][0]["title"] is None
    assert store.get_project(pid)["slide_count"] == 2
    assert slides.image_path(pid, 1).name == "page_002.png"
    assert slides.ensure_images(pid) == {"cached": True, "images_source": "pdf"}

    slides.update_slide(pid, 0, speaker_notes="Narrate page one.")

    # The generate path re-reads the PDF through FileItem.load_pdf on every
    # run; it used to recreate the project, wiping the notes.
    item = file_item_module.FileItem(store.PROJECTS_DIR / pid / "pages.pdf", projects_base=store.PROJECTS_DIR / pid)
    assert item.load_pdf() is True
    assert item.project_manager.state.slides[0].speaker_notes == "Narrate page one."
    assert Path(item.project_manager.state.slides[1].image_path).name == "page_002.png"
    assert slides.reset_slide(pid, 0)["speaker_notes"] == "", "a PDF has no notes of its own"

    # Lost page images are rendered again, without PowerPoint.
    (store.PROJECTS_DIR / pid / "pages_project" / "images" / "page_001.png").unlink()
    assert slides.slides_payload(pid)["slides_ready"] is False
    messages = []
    assert slides.ensure_images(pid, lambda f, m: messages.append(m)) == {"cached": False, "images_source": "pdf", "slides": 2}
    assert slides.image_path(pid, 0) is not None and messages[-1] == "Slide previews ready"

    with pytest.raises(ValueError):
        slides.export_notes_pptx(pid)


# -- export ------------------------------------------------------------------------------

def test_export_notes_pptx_writes_the_edited_notes_into_a_copy():
    pid = _import_deck()
    slides.update_slide(pid, 1, speaker_notes="New notes for slide two.")
    slides.update_slide(pid, 2, speaker_notes="")

    out = slides.export_notes_pptx(pid)

    assert out == (store.PROJECTS_DIR / pid / "exports" / "deck-notes.pptx").resolve()
    exported = [s.notes_slide.notes_text_frame.text for s in Presentation(str(out)).slides]
    assert exported == ["Welcome to the deck.", "New notes for slide two.", ""]
    original = [s.notes_slide.notes_text_frame.text for s in Presentation(str(store.PROJECTS_DIR / pid / "deck.pptx")).slides]
    assert original == list(NOTES), "the source deck is untouched"


def test_export_refuses_a_name_that_leaves_the_project(client):
    pid = _import_deck()
    record = store.get_project(pid)
    record["name"] = "../../evil"
    store.save_project(record)
    with pytest.raises(ValueError):
        slides.export_notes_pptx(pid)
    assert client.get(f"/api/projects/{pid}/export/pptx").status_code == 400
    assert not (store.PROJECTS_DIR.parent / "evil-notes.pptx").exists()
    assert not list(store.PROJECTS_DIR.parent.glob("*.pptx"))


# -- the routes -------------------------------------------------------------------------

def test_routes_need_a_signed_in_user_and_a_deck_or_pdf(client):
    from api.app import app

    pid = _import_deck()
    video = store.import_upload("clip.mp4", b"v")["id"]
    anonymous = TestClient(app)
    assert anonymous.get(f"/api/projects/{pid}/slides").status_code == 401
    assert anonymous.patch(f"/api/projects/{pid}/slides/0", json={"speaker_notes": "x"}).status_code == 401
    assert anonymous.get(f"/api/projects/{pid}/slides/0/image").status_code == 401
    assert anonymous.get(f"/api/projects/{pid}/export/pptx").status_code == 401

    for path in ("slides", "slides/0/image", "export/pptx"):
        assert client.get(f"/api/projects/aabbccddeeff/{path}").status_code == 404, path
        assert client.get(f"/api/projects/{video}/{path}").status_code == 400, path
    assert client.post("/api/projects/aabbccddeeff/slides/render").status_code == 404
    assert client.post(f"/api/projects/{video}/slides/render").status_code == 400
    assert client.patch(f"/api/projects/{video}/slides/0", json={"speaker_notes": "x"}).status_code == 400
    assert client.post(f"/api/projects/{video}/slides/0/undo").status_code == 400
    assert not (store.PROJECTS_DIR / pid / "deck_project").exists(), "a refused request materialises nothing"


def test_get_slides_lists_the_deck(client):
    pid = _import_deck()
    r = client.get(f"/api/projects/{pid}/slides")
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["slides_ready"] is False and body["images_source"] is None
    assert [s["title"] for s in body["slides"]] == ["Slide 1 title", "Slide 2 title", "Slide 3 title"]
    assert body["slides"][2]["speaker_notes"] == "Third slide notes."
    assert client.get(f"/api/projects/{pid}").json()["slide_count"] == 3


def test_patch_slide_validates_and_returns_the_slide(client):
    pid = _import_deck()
    r = client.patch(f"/api/projects/{pid}/slides/0", json={"speaker_notes": "Hi", "pause_override": 1.5})
    assert r.status_code == 200, r.text
    assert r.json()["speaker_notes"] == "Hi" and r.json()["pause_override"] == 1.5
    assert r.json()["notes_history_depth"] == 1

    r = client.patch(f"/api/projects/{pid}/slides/7", json={"speaker_notes": "x"})
    assert r.status_code == 400 and "No slide with index 7" in r.json()["detail"]
    assert client.patch(f"/api/projects/{pid}/slides/0", json={"pause_override": 31}).status_code == 422
    assert client.patch(f"/api/projects/{pid}/slides/0", json={"pause_override": -0.5}).status_code == 422
    assert client.patch(f"/api/projects/{pid}/slides/0", json={"pause_override": True}).status_code == 422
    assert client.patch(f"/api/projects/{pid}/slides/0", json={"pause_override": 2}).status_code == 200, "an int is a number"
    assert client.patch(f"/api/projects/{pid}/slides/0", json={"speaker_notes": "x" * 20001}).status_code == 422
    assert client.patch(f"/api/projects/{pid}/slides/0", json={"colour": "red"}).status_code == 422

    r = client.patch(f"/api/projects/{pid}/slides/0", json={"voice_override": "af_heart"})
    assert r.status_code == 400 and "looks like a Kokoro voice" in r.json()["detail"]
    r = client.patch(f"/api/projects/{pid}/slides/0", json={"voice_override": "af_heart", "provider": "kokoro"})
    assert r.status_code == 200 and r.json()["voice_override"] == "af_heart"
    r = client.patch(f"/api/projects/{pid}/slides/0", json={"voice_override": None, "pause_override": None})
    assert r.status_code == 200 and r.json()["voice_override"] is None and r.json()["pause_override"] is None
    assert r.json()["speaker_notes"] == "Hi", "fields left out are left alone"


def test_undo_and_reset_routes(client):
    pid = _import_deck()
    assert client.post(f"/api/projects/{pid}/slides/0/undo").status_code == 409

    client.patch(f"/api/projects/{pid}/slides/0", json={"speaker_notes": "Changed"})
    r = client.post(f"/api/projects/{pid}/slides/0/undo")
    assert r.status_code == 200 and r.json()["speaker_notes"] == "Welcome to the deck."

    client.patch(f"/api/projects/{pid}/slides/0", json={"speaker_notes": "Changed again"})
    r = client.post(f"/api/projects/{pid}/slides/0/reset")
    assert r.status_code == 200
    assert r.json()["speaker_notes"] == "Welcome to the deck." and r.json()["notes_history_depth"] == 2
    assert client.post(f"/api/projects/{pid}/slides/9/undo").status_code == 400
    assert client.post(f"/api/projects/{pid}/slides/9/reset").status_code == 400


def test_bulk_route_saves_the_dirty_slides(client):
    pid = _import_deck()
    r = client.patch(f"/api/projects/{pid}/slides", json={"slides": [
        {"index": 0, "speaker_notes": "One"}, {"index": 2, "speaker_notes": "Three"},
    ]})
    assert r.status_code == 200, r.text
    assert [s["speaker_notes"] for s in r.json()["slides"]] == ["One", "", "Three"]

    r = client.patch(f"/api/projects/{pid}/slides", json={"slides": [{"index": 1, "speaker_notes": "Two"}, {"index": 5, "speaker_notes": "x"}]})
    assert r.status_code == 400
    assert [s["speaker_notes"] for s in slides.list_slides(pid)] == ["One", "", "Three"]
    assert client.patch(f"/api/projects/{pid}/slides", json={"slides": [{"index": -1, "speaker_notes": "x"}]}).status_code == 422
    assert client.patch(f"/api/projects/{pid}/slides", json={"slides": [{"index": True, "speaker_notes": "x"}]}).status_code == 422
    assert client.patch(f"/api/projects/{pid}/slides", json={"slides": [{"index": 0, "speaker_notes": "x", "title": "no"}]}).status_code == 422


def test_writes_are_refused_while_a_job_is_attached_to_the_project(client):
    pid = _import_deck()
    other = _import_deck("other.pptx")
    job = _BlockingJob(pid)
    try:
        assert jobs.active_for(pid)["id"] == job.id and jobs.get(job.id)["project_id"] == pid
        assert jobs.active_for(other) is None
        writes = (
            lambda: client.patch(f"/api/projects/{pid}/slides/0", json={"speaker_notes": "x"}),
            lambda: client.post(f"/api/projects/{pid}/slides/0/undo"),
            lambda: client.post(f"/api/projects/{pid}/slides/0/reset"),
            lambda: client.patch(f"/api/projects/{pid}/slides", json={"slides": [{"index": 0, "speaker_notes": "x"}]}),
            lambda: client.post(f"/api/projects/{pid}/slides/render"),
        )
        for write in writes:
            r = write()
            assert r.status_code == 409, r.text
            assert "A job is running for this project" in r.json()["detail"] and "generate" in r.json()["detail"]
        assert client.get(f"/api/projects/{pid}/slides").status_code == 200, "reads are fine"
        assert client.patch(f"/api/projects/{other}/slides/0", json={"speaker_notes": "x"}).status_code == 200, "another project is fine"
        assert slides.list_slides(pid)[0]["speaker_notes"] == "Welcome to the deck.", "nothing was written"
    finally:
        job.release.set()
    assert _wait_job(client, job.id)["status"] == "done"
    assert jobs.active_for(pid) is None
    assert client.patch(f"/api/projects/{pid}/slides/0", json={"speaker_notes": "x"}).status_code == 200


def test_render_job_then_the_image_route(client, monkeypatch):
    pid = _import_deck()
    _fake_exporter(monkeypatch)
    assert client.get(f"/api/projects/{pid}/slides/0/image").status_code == 404, "not rendered yet"

    r = client.post(f"/api/projects/{pid}/slides/render")
    assert r.status_code == 200 and "job_id" in r.json()
    job = _wait_job(client, r.json()["job_id"])
    assert job["status"] == "done" and job["kind"] == "render-slides" and job["project_id"] == pid, job
    assert job["result"] == {"cached": False, "images_source": "powerpoint", "slides": 3}

    r = client.get(f"/api/projects/{pid}/slides/0/image")
    assert r.status_code == 200
    assert r.headers["content-type"] == "image/png"
    assert r.headers["cache-control"] == "no-cache"
    assert r.content == PNG
    assert client.get(f"/api/projects/{pid}/slides/3/image").status_code == 404

    assert client.post(f"/api/projects/{pid}/slides/render").json() == {"cached": True}
    body = client.get(f"/api/projects/{pid}/slides").json()
    assert body["slides_ready"] is True and body["images_source"] == "powerpoint" and body["images_rendered_at"]


def test_a_second_render_request_returns_the_render_in_flight(client, monkeypatch):
    pid = _import_deck()
    entered, gate = threading.Event(), threading.Event()
    _fake_exporter(monkeypatch, entered=entered, gate=gate)

    first = client.post(f"/api/projects/{pid}/slides/render").json()["job_id"]
    assert entered.wait(5)
    assert client.post(f"/api/projects/{pid}/slides/render").json() == {"job_id": first}
    assert client.post(f"/api/projects/{pid}/slides/render").json() == {"job_id": first}
    gate.set()
    assert _wait_job(client, first)["status"] == "done"
    assert len(_FakeExporter.runs) == 1, "one export for three clicks"
    assert client.post(f"/api/projects/{pid}/slides/render").json() == {"cached": True}


def test_export_route_serves_the_deck_with_notes_as_an_attachment(client):
    pid = _import_deck()
    client.patch(f"/api/projects/{pid}/slides/2", json={"speaker_notes": "Exported."})

    r = client.get(f"/api/projects/{pid}/export/pptx")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.presentationml.presentation")
    assert r.headers["content-disposition"] == 'attachment; filename="deck-notes.pptx"'
    notes = [s.notes_slide.notes_text_frame.text for s in Presentation(io.BytesIO(r.content)).slides]
    assert notes == ["Welcome to the deck.", "", "Exported."]


# -- the request schemas ----------------------------------------------------------------

def test_slide_update_schema_bounds_and_forbids_unknown_keys():
    assert SlideUpdate().model_dump(exclude_unset=True) == {}
    assert SlideUpdate(pause_override=None).model_dump(exclude_unset=True) == {"pause_override": None}, "an explicit null clears"
    assert SlideUpdate(pause_override=30, speaker_notes="x" * 20000).pause_override == 30
    assert SlideUpdate(pause_override=2).pause_override == 2, "an int is a number"
    for kwargs in ({"pause_override": 30.5}, {"pause_override": -1}, {"pause_override": True}, {"pause_override": "2"},
                   {"speaker_notes": "x" * 20001}, {"alt_text": "x" * 2001}, {"colour": "red"}):
        with pytest.raises(ValidationError):
            SlideUpdate(**kwargs)
    for item in ({"index": -1, "speaker_notes": ""}, {"index": 0}, {"index": True, "speaker_notes": ""}, {"index": "0", "speaker_notes": ""}):
        with pytest.raises(ValidationError):
            SlidesBulkUpdate(slides=[item])


# -- the generate pipeline takes the same lock and reports its exporter -----------------

class _LockCheckingExporter:
    """Records whether SLIDE_EXPORT_LOCK was held while it exported."""

    held = []

    def __init__(self, pptx_path, output_dir):
        self.output_dir = Path(output_dir)
        self.backend = None

    def export_slides_as_images(self, progress_callback=None):
        _LockCheckingExporter.held.append(slides.SLIDE_EXPORT_LOCK.locked())
        image = self.output_dir / "slide_001.png"
        image.parent.mkdir(parents=True, exist_ok=True)
        image.write_bytes(PNG)
        self.backend = "powerpoint"
        return [ExportedSlide(index=0, image_path=image)]


def _pipeline_deck(tmp_path, monkeypatch):
    """A one-slide project with cached audio, so process_files reaches the
    export stage and then nothing else that needs a TTS or an encode."""
    _LockCheckingExporter.held = []
    monkeypatch.setattr(processing, "PPTXExporter", _LockCheckingExporter)
    monkeypatch.setattr(processing.VideoProcessor, "_create_tts_generator", lambda self: None)
    monkeypatch.setattr(processing.VideoProcessor, "_build_video", lambda self, pm, out, **kwargs: True)
    monkeypatch.setattr(processing.VideoProcessor, "_sidecar_outputs", lambda self, *args, **kwargs: {})
    pm = ProjectManager(tmp_path / "deck_project")
    pm.create_project(tmp_path / "deck.pptx", ["Notes"], voice_id="v")
    audio = tmp_path / "slide_001_audio.mp3"
    audio.write_bytes(b"mp3")
    pm.update_slide_audio(0, audio, 2.0, "v")
    monkeypatch.setattr(processing.VideoProcessor, "_get_or_create_project", lambda self, item: pm)
    item = types.SimpleNamespace(path=tmp_path / "deck.pptx", slide_count=1, project_manager=None, has_project=False)
    return pm, item


def test_process_files_exports_slides_under_the_export_lock(tmp_path, monkeypatch):
    pm, item = _pipeline_deck(tmp_path, monkeypatch)
    processor = processing.VideoProcessor(voice_id="v", subtitles="none")
    assert processor.images_backend is None

    assert processor.process_files([item], output_dir=tmp_path / "out") == 1
    assert _LockCheckingExporter.held == [True], "the export stage ran inside SLIDE_EXPORT_LOCK"
    assert Path(pm.state.slides[0].image_path).name == "slide_001.png"
    assert processor.images_backend == "powerpoint", "the exporter's backend, for the project record"
    assert not slides.SLIDE_EXPORT_LOCK.locked(), "released afterwards"

    again = processing.VideoProcessor(voice_id="v", subtitles="none")
    assert again.process_files([item], output_dir=tmp_path / "out") == 1
    assert _LockCheckingExporter.held == [True], "images present: no second export"
    assert again.images_backend is None, "nothing exported this run"


def test_process_files_reports_the_wait_while_another_export_holds_the_lock(tmp_path, monkeypatch):
    pm, item = _pipeline_deck(tmp_path, monkeypatch)
    messages = []
    processor = processing.VideoProcessor(voice_id="v", subtitles="none")
    result = []

    slides.SLIDE_EXPORT_LOCK.acquire()  # another job's export in progress
    try:
        worker = threading.Thread(
            target=lambda: result.append(processor.process_files([item], output_dir=tmp_path / "out",
                                                                 progress=lambda f, m: messages.append(m))),
        )
        worker.start()
        _wait_until(lambda: f"[1/1] deck.pptx: {slides.WAITING_FOR_EXPORT}" in messages, what="the waiting message")
        assert _LockCheckingExporter.held == [], "queued behind the other export"
    finally:
        slides.SLIDE_EXPORT_LOCK.release()
    worker.join(timeout=10)
    assert result == [1] and _LockCheckingExporter.held == [True]


def test_rebuild_and_prepare_take_the_lock_too(tmp_path, monkeypatch):
    pm, item = _pipeline_deck(tmp_path, monkeypatch)
    item.project_manager = pm
    pm.state.slides[0].image_path = None
    pm.save()
    processor = processing.VideoProcessor(voice_id="v", subtitles="none")
    assert processor.rebuild_files([item], output_dir=tmp_path / "out") == 1
    assert _LockCheckingExporter.held == [True] and processor.images_backend == "powerpoint"

    pm.state.slides[0].image_path = None
    pm.save()
    processor = processing.VideoProcessor(voice_id="v")
    assert processor.prepare_audio([item]) == 1
    assert _LockCheckingExporter.held == [True, True] and processor.images_backend == "powerpoint"


class _FakeRenderProcessor:
    """Stands in for VideoProcessor in the generate route: writes the MP4 and
    reports the exporter backend a test asks for."""

    backend = "pillow"

    def __init__(self, **kwargs):
        self.outputs = {}
        self.images_backend = None

    def process_files(self, files, output_dir, progress=None, preview_seconds=0):
        from utils.helpers import get_output_filename

        out = get_output_filename(files[0].path, output_dir)
        out.write_bytes(b"mp4")
        self.images_backend = type(self).backend
        return 1


class _FakeFileItem:
    def __init__(self, path, projects_base=None):
        self.path = path

    def load(self):
        return True

    def load_pdf(self):
        return True


def test_generate_records_where_the_images_came_from(client, monkeypatch):
    monkeypatch.setattr(processing, "VideoProcessor", _FakeRenderProcessor)
    monkeypatch.setattr(file_item_module, "FileItem", _FakeFileItem)

    pid = _import_deck()
    _FakeRenderProcessor.backend = "pillow"
    r = client.post(f"/api/projects/{pid}/generate", json={})
    job = _wait_job(client, r.json()["job_id"])
    assert job["status"] == "done" and job["project_id"] == pid
    saved = store.get_project(pid)
    assert saved["images_source"] == "pillow" and saved["images_rendered_at"]
    assert slides.slides_payload(pid)["images_source"] == "pillow"

    # No export this run (images cached): the record keeps what it knew.
    _FakeRenderProcessor.backend = None
    stamp = saved["images_rendered_at"]
    assert _wait_job(client, client.post(f"/api/projects/{pid}/generate", json={}).json()["job_id"])["status"] == "done"
    assert store.get_project(pid)["images_source"] == "pillow"
    assert store.get_project(pid)["images_rendered_at"] == stamp

    pdf = store.import_upload("pages.pdf", b"%PDF-fake")["id"]
    assert _wait_job(client, client.post(f"/api/projects/{pdf}/generate", json={}).json()["job_id"])["status"] == "done"
    assert store.get_project(pdf)["images_source"] == "pdf", "a PDF's pages are its own renders"

    with pytest.raises(ValueError):
        slides.record_images_source(pid, "libreoffice")
