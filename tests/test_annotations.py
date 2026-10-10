"""Annotations on a still (T3b): the record, the four routes and the one oracle.

- a new still says it is not annotated and shows its own size; one record can
  be read; a record from before annotations existed gets the defaults;
- Save keeps the still untouched and the export beside it, and from then on
  the image, the download and Add to deck all show the annotated picture
  (``image_path``, the one oracle), ``?original=1`` the still; Revert deletes
  the export and the document, and everything shows the still again;
- the deck slide letterboxes the CROP's size, as it letterboxes a still;
- ownership is the still's: another editor is refused, an administrator is not;
  the editor's routes work from another machine (the browser edition);
- a document that does not validate against the still, a picture that is not
  a PNG of the right size, and a body over the cap are refused (400 / 413)
  with a plain line, and nothing is written;
- a record that claims ``annotated`` while the export is gone is served as the
  original, and the log says so - never a 500;
- the audit rows, and the limits rows that name the constants.

Every still here is a Pillow-made PNG through the fake grab: nothing reads
the screen.
"""

from __future__ import annotations

import io
import json
import logging
import struct
import zlib
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation

import api.routers.capture as capture_router
from api import store as auth_store
from services import capture, captures, recordings
from services import projects as store
from utils.config import config

PASSWORD = "Owner-pass-12345"
LOOPBACK = ("127.0.0.1", 50000)
RED, BLUE, GREEN, WHITE = (200, 30, 30), (0, 0, 255), (0, 200, 0), (255, 255, 255)
LIMITS_MD = Path(__file__).resolve().parent.parent / "docs" / "reference" / "limits.md"
PNG_MAGIC_BYTES = captures.PNG_MAGIC


@pytest.fixture
def rig(tmp_path, monkeypatch):
    monkeypatch.setattr(store, "PROJECTS_DIR", tmp_path / "projects")
    monkeypatch.setattr(captures, "CAPTURES_DIR", tmp_path / "captures")
    monkeypatch.setattr(captures, "TEMP_DIR", tmp_path / "temp")
    monkeypatch.setattr(capture_router, "SCREEN_DIR", tmp_path / "temp" / "capture")
    monkeypatch.setattr(recordings, "RECORDINGS_DIR", tmp_path / "temp" / "recordings")
    monkeypatch.setattr(config, "_config", {})
    monkeypatch.setattr(config, "save", lambda: None)
    monkeypatch.setattr(auth_store, "DATA_DIR", tmp_path / "auth")
    monkeypatch.setattr(auth_store, "DB_PATH", tmp_path / "auth" / "media_studio.db")
    from api.app import app

    with TestClient(app, client=LOOPBACK) as admin:
        seeded = auth_store.authenticate("admin", "admin")
        auth_store.change_password(seeded["id"], "admin", must_change=False)
        assert admin.post("/api/auth/login", json={"username": "admin", "password": "admin"}).status_code == 200
        auth_store.create_user("other", PASSWORD, "Otto Other", role="editor", must_change_password=False)
        other = TestClient(app, client=LOOPBACK)
        assert other.post("/api/auth/login", json={"username": "other", "password": PASSWORD}).status_code == 200
        yield {"admin": admin, "other": other, "tmp": tmp_path, "app": app}


# -- helpers --------------------------------------------------------------------------

def _fake_grab(colour=RED):
    """A stand-in for gdigrab: a PNG of the asked size and colour."""

    def grab(x, y, width, height, *, cursor, out_dir):
        capture.validate_region(x, y, width, height)
        Path(out_dir).mkdir(parents=True, exist_ok=True)
        out = Path(out_dir) / f"still-fake-{width}x{height}.png"
        Image.new("RGB", (width, height), colour).save(out)
        return out

    return grab


def _png(size: tuple[int, int], colour=BLUE) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", size, colour).save(buf, format="PNG")
    return buf.getvalue()


def _still(rig, monkeypatch, width: int, height: int, colour=RED, name: str = "Still", who: str = "admin") -> str:
    monkeypatch.setattr(captures.capture, "take_still", _fake_grab(colour))
    r = rig[who].post("/api/captures", json={"x": 0, "y": 0, "width": width, "height": height, "name": name})
    assert r.status_code == 200, r.text
    return r.json()["id"]


def _deck_bytes(slide_count: int = 1) -> bytes:
    prs = Presentation()
    for i in range(slide_count):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"Slide {i + 1}"
        slide.placeholders[1].text = f"Bullet {i + 1}"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _arrow(**over) -> dict:
    return {"id": "a1", "type": "arrow", "colour": "#e53935", "stroke": "medium", "x1": 2.0, "y1": 2.0, "x2": 20.0, "y2": 12.0, **over}


def _line(**over) -> dict:
    return {**_arrow(id="l1", type="line"), **over}


def _box(type_: str, id_: str, **over) -> dict:
    return {"id": id_, "type": type_, "colour": "#1e88e5", "stroke": "thin", "x": 1.0, "y": 1.0, "width": 10.0, "height": 8.0, **over}


def _text(**over) -> dict:
    return {"id": "t1", "type": "text", "colour": "#000000", "stroke": "medium", "x": 3.0, "y": 3.0, "width": 30.0, "height": 12.0,
            "text": "Hello", "text_size": "small", **over}


def _callout(**over) -> dict:
    return {"id": "c1", "type": "callout", "colour": "#fb8c00", "stroke": "thick", "x": 3.0, "y": 3.0, "width": 30.0, "height": 12.0,
            "tail_x": 1.0, "tail_y": 25.0, "text": "Look here", "text_size": "large", **over}


def _step(id_: str = "s1", sequence: int = 0, **over) -> dict:
    return {"id": id_, "type": "step", "colour": "#e53935", "stroke": "medium", "x": 10.0, "y": 10.0, "sequence": sequence, **over}


def _magnify(**over) -> dict:
    return {"id": "m1", "type": "magnify", "colour": "#43a047", "stroke": "medium", "x": 1.0, "y": 1.0, "width": 10.0, "height": 8.0,
            "lens_x": 15.0, "lens_y": 1.0, "factor": 2, **over}


def _doc(shapes=(), crop=None) -> dict:
    return {"version": 1, "shapes": list(shapes), "crop": crop}


def _put(client: TestClient, cid: str, document, png: bytes):
    """``document`` a dict (sent as JSON) or a string (sent as is); ``("file", dict)`` sends it as a FILE part."""
    files = {"image": ("annotated.png", png, "image/png")}
    if isinstance(document, tuple):
        files["document"] = ("document.json", json.dumps(document[1]).encode("utf-8"), "application/json")
        return client.put(f"/api/captures/{cid}/annotations", files=files)
    text = document if isinstance(document, str) else json.dumps(document)
    return client.put(f"/api/captures/{cid}/annotations", data={"document": text}, files=files)


def _chunk(kind: bytes, payload: bytes) -> bytes:
    return struct.pack(">I", len(payload)) + kind + payload + struct.pack(">I", zlib.crc32(kind + payload) & 0xFFFFFFFF)


def _crafted_png(width: int, height: int, idat: bytes = b"\0" * 16) -> bytes:
    """A structurally valid PNG - correct chunk lengths and CRCs - whose IHDR claims ``width`` x
    ``height`` and whose IDAT holds whatever ``idat`` is: a bomb-sized header, or data that cannot be
    decoded. About 100 bytes."""
    ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    return PNG_MAGIC_BYTES + _chunk(b"IHDR", ihdr) + _chunk(b"IDAT", idat) + _chunk(b"IEND", b"")


def _pixel(client: TestClient, cid: str, xy=(0, 0), original: bool = False):
    r = client.get(f"/api/captures/{cid}/image" + ("?original=1" if original else ""))
    assert r.status_code == 200, r.text
    return Image.open(io.BytesIO(r.content)).convert("RGB").getpixel(xy)


def _signed_in_from(app, client: tuple[str, int]) -> TestClient:
    c = TestClient(app, client=client)
    assert c.post("/api/auth/login", json={"username": "admin", "password": "admin"}).status_code == 200
    return c


# -- the record -------------------------------------------------------------------------

def test_a_new_still_is_not_annotated_and_one_record_can_be_read(rig, monkeypatch):
    c = rig["admin"]
    cid = _still(rig, monkeypatch, 40, 30)
    r = c.get(f"/api/captures/{cid}")
    assert r.status_code == 200, r.text
    rec = r.json()
    assert rec["id"] == cid and rec["annotated"] is False and rec["annotations"] is None
    assert (rec["output_width"], rec["output_height"]) == (40, 30)
    assert "annotated_at" not in rec
    listed = c.get("/api/captures").json()["captures"]
    assert [x["id"] for x in listed] == [cid]
    assert listed[0]["annotated"] is False and (listed[0]["output_width"], listed[0]["output_height"]) == (40, 30)
    assert "annotations" not in listed[0], "the list carries no documents"
    assert c.get("/api/captures/000000000000").status_code == 404
    assert c.get("/api/captures/not-an-id").status_code == 404


def test_a_record_from_before_annotations_existed_gets_the_defaults(rig, monkeypatch):
    """A capture.json written by 0.14.x has none of the new fields: the list and the record fill them
    in, and the image is the still."""
    c = rig["admin"]
    cid = _still(rig, monkeypatch, 24, 16)
    path = rig["tmp"] / "captures" / cid / "capture.json"
    rec = json.loads(path.read_text(encoding="utf-8"))
    for key in ("annotated", "output_width", "output_height"):
        rec.pop(key)
    path.write_text(json.dumps(rec), encoding="utf-8")
    one = c.get(f"/api/captures/{cid}").json()
    assert one["annotated"] is False and (one["output_width"], one["output_height"]) == (24, 16) and one["annotations"] is None
    assert c.get("/api/captures").json()["captures"][0]["annotated"] is False
    assert _pixel(c, cid) == RED


# -- save, the oracle, revert ---------------------------------------------------------------

def test_save_keeps_the_original_and_every_route_shows_the_annotated_picture_until_revert(rig, monkeypatch):
    c = rig["admin"]
    cid = _still(rig, monkeypatch, 40, 30, RED, name="Notepad")
    deck = c.post("/api/projects/import", files={"file": ("deck.pptx", _deck_bytes(1), "application/octet-stream")}).json()
    doc = _doc([_arrow(), _step()])

    r = _put(c, cid, doc, _png((40, 30), BLUE))
    assert r.status_code == 200, r.text
    rec = r.json()
    assert rec["annotated"] is True and rec["annotated_at"] and rec["annotations"] == doc
    assert (rec["output_width"], rec["output_height"]) == (40, 30)
    assert c.get(f"/api/captures/{cid}").json()["annotations"] == doc
    assert c.get("/api/captures").json()["captures"][0]["annotated"] is True

    # On disk: the still untouched, the export beside it, the record with the document.
    folder = rig["tmp"] / "captures" / cid
    assert Image.open(folder / "still.png").convert("RGB").getpixel((0, 0)) == RED
    assert Image.open(folder / "annotated.png").convert("RGB").getpixel((0, 0)) == BLUE
    stored = json.loads((folder / "capture.json").read_text(encoding="utf-8"))
    assert stored["annotated"] is True and stored["annotations"] == doc

    # The one oracle: image, download and add-to-deck show the annotated picture; the original on request.
    assert _pixel(c, cid) == BLUE
    assert _pixel(c, cid, original=True) == RED
    dl = c.get(f"/api/captures/{cid}/download")
    assert dl.status_code == 200 and 'filename="Notepad.png"' in dl.headers["content-disposition"]
    assert Image.open(io.BytesIO(dl.content)).convert("RGB").getpixel((0, 0)) == BLUE
    added = c.post(f"/api/captures/{cid}/add-to-deck", json={"project_id": deck["id"]})
    assert added.status_code == 200, added.text
    slide = Image.open(io.BytesIO(c.get(f"/api/projects/{deck['id']}/slides/{added.json()['index']}/image").content)).convert("RGB")
    assert slide.getpixel((960, 540)) == BLUE

    # Revert: the export and the document go; the still is what everything shows.
    assert c.delete(f"/api/captures/{cid}/annotations").status_code == 204
    assert not (folder / "annotated.png").exists() and (folder / "still.png").is_file()
    rec = c.get(f"/api/captures/{cid}").json()
    assert rec["annotated"] is False and rec["annotations"] is None and "annotated_at" not in rec
    assert (rec["output_width"], rec["output_height"]) == (40, 30)
    assert _pixel(c, cid) == RED and _pixel(c, cid, original=True) == RED
    assert Image.open(io.BytesIO(c.get(f"/api/captures/{cid}/download").content)).convert("RGB").getpixel((0, 0)) == RED
    added = c.post(f"/api/captures/{cid}/add-to-deck", json={"project_id": deck["id"]}).json()
    slide = Image.open(io.BytesIO(c.get(f"/api/projects/{deck['id']}/slides/{added['index']}/image").content)).convert("RGB")
    assert slide.getpixel((960, 540)) == RED
    # Reverting what is not annotated is nothing to do, not an error.
    assert c.delete(f"/api/captures/{cid}/annotations").status_code == 204


def test_the_deck_slide_letterboxes_the_crops_size(rig, monkeypatch):
    """A crop of 100 x 400 out of a 400 x 400 still: the export is the crop's size, and Add to deck
    letterboxes THAT - a tall picture, so bars at the sides and the picture the full height."""
    c = rig["admin"]
    cid = _still(rig, monkeypatch, 400, 400, RED)
    deck = c.post("/api/projects/import", files={"file": ("deck.pptx", _deck_bytes(1), "application/octet-stream")}).json()
    r = _put(c, cid, _doc([_arrow()], crop={"x": 50, "y": 0, "width": 100, "height": 400}), _png((100, 400), GREEN))
    assert r.status_code == 200, r.text
    assert (r.json()["output_width"], r.json()["output_height"]) == (100, 400)
    assert c.get("/api/captures").json()["captures"][0]["output_width"] == 100
    assert Image.open(io.BytesIO(c.get(f"/api/captures/{cid}/image").content)).size == (100, 400)
    added = c.post(f"/api/captures/{cid}/add-to-deck", json={"project_id": deck["id"]}).json()
    slide = Image.open(io.BytesIO(c.get(f"/api/projects/{deck['id']}/slides/{added['index']}/image").content)).convert("RGB")
    assert slide.size == (1920, 1080)
    assert slide.getpixel((960, 540)) == GREEN and slide.getpixel((960, 5)) == GREEN
    assert slide.getpixel((10, 540)) == WHITE and slide.getpixel((1910, 540)) == WHITE


def test_every_shape_type_round_trips_through_the_record(rig, monkeypatch):
    c = rig["admin"]
    cid = _still(rig, monkeypatch, 60, 40)
    shapes = [
        _arrow(), _line(), _box("rect", "r1"), _box("ellipse", "e1"), _box("blur", "b1"), _box("highlight", "h1", colour="#fdd835"),
        _text(), _callout(), _step("s1", 0), _step("s2", 0), _step("s3", 1), _magnify(factor=3),
    ]
    doc = _doc(shapes, crop={"x": 4, "y": 2, "width": 50, "height": 36})
    r = _put(c, cid, doc, _png((50, 36)))
    assert r.status_code == 200, r.text
    assert r.json()["annotations"] == doc
    assert c.get(f"/api/captures/{cid}").json()["annotations"] == doc
    # A shape may hang off the edge (the renderer clips it): -size .. 2 x size is legal.
    off = _doc([_arrow(x1=-60.0, y1=-40.0, x2=120.0, y2=80.0)])
    assert _put(c, cid, off, _png((60, 40))).status_code == 200


def test_delete_still_removes_the_export_too(rig, monkeypatch):
    c = rig["admin"]
    cid = _still(rig, monkeypatch, 20, 20)
    assert _put(c, cid, _doc([_arrow()]), _png((20, 20))).status_code == 200
    folder = rig["tmp"] / "captures" / cid
    assert (folder / "annotated.png").is_file()
    assert c.delete(f"/api/captures/{cid}").status_code == 204
    assert not folder.exists()
    assert c.get(f"/api/captures/{cid}").status_code == 404
    assert _put(c, cid, _doc([_arrow()]), _png((20, 20))).status_code == 404
    assert c.delete(f"/api/captures/{cid}/annotations").status_code == 404


# -- ownership ---------------------------------------------------------------------------------

def test_the_annotations_belong_to_the_stills_owner(rig, monkeypatch):
    admin, other = rig["admin"], rig["other"]
    cid = _still(rig, monkeypatch, 20, 20)
    assert other.get(f"/api/captures/{cid}").status_code == 403
    assert _put(other, cid, _doc([_arrow()]), _png((20, 20))).status_code == 403
    assert other.delete(f"/api/captures/{cid}/annotations").status_code == 403
    assert other.get(f"/api/captures/{cid}/image?original=1").status_code == 403
    assert not (rig["tmp"] / "captures" / cid / "annotated.png").exists()
    # The other editor's own still: theirs to annotate, and an administrator's too.
    theirs = _still(rig, monkeypatch, 20, 20, who="other")
    assert _put(other, theirs, _doc([_arrow()]), _png((20, 20))).status_code == 200
    assert admin.get(f"/api/captures/{theirs}").json()["annotated"] is True
    assert _put(admin, theirs, _doc([_step()]), _png((20, 20))).status_code == 200
    assert admin.delete(f"/api/captures/{theirs}/annotations").status_code == 204
    assert other.get(f"/api/captures/{theirs}").json()["annotated"] is False


def test_the_editor_works_from_another_machine(rig, monkeypatch):
    """Annotating touches no screen: unlike taking a still, it is not this computer's only (the browser
    edition annotates too)."""
    cid = _still(rig, monkeypatch, 20, 20)
    remote = _signed_in_from(rig["app"], ("10.20.30.40", 50123))
    assert remote.get(f"/api/captures/{cid}").status_code == 200
    assert _put(remote, cid, _doc([_arrow()]), _png((20, 20))).status_code == 200
    assert remote.get(f"/api/captures/{cid}/image?original=1").status_code == 200
    assert remote.delete(f"/api/captures/{cid}/annotations").status_code == 204


# -- validation ---------------------------------------------------------------------------------

def _jpeg(size) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", size, BLUE).save(buf, format="JPEG")
    return buf.getvalue()


BAD = [
    ("bad json", "{not json", _png((40, 30)), "not valid JSON"),
    ("not an object", json.dumps([1, 2]), _png((40, 30)), "not valid"),
    ("another version", _doc([_arrow()]) | {"version": 2}, _png((40, 30)), "version"),
    ("too many shapes", _doc([_step(f"s{i}") for i in range(captures.MAX_SHAPES + 1)]), _png((40, 30)), "shapes"),
    ("text too long", _doc([_text(text="x" * (captures.MAX_TEXT_CHARS + 1))]), _png((40, 30)), "text"),
    ("a NaN", json.dumps(_doc([_arrow(x1=float("nan"))])), _png((40, 30)), "x1"),
    ("an infinity", json.dumps(_doc([_arrow(y2=float("inf"))])), _png((40, 30)), "y2"),
    ("out of range", _doc([_arrow(x2=3 * 40.0)]), _png((40, 30)), "outside the still"),
    ("out of range below", _doc([_step(y=-31.0)]), _png((40, 30)), "outside the still"),
    ("a zero-size box", _doc([_box("rect", "r1", width=0.0)]), _png((40, 30)), "width"),
    ("an unknown tool", _doc([_arrow(type="pen")]), _png((40, 30)), "type"),
    ("an unknown field", _doc([_arrow(spin=1)]), _png((40, 30)), "spin"),
    ("a bad colour", _doc([_arrow(colour="red")]), _png((40, 30)), "colour"),
    ("a bad lens factor", _doc([_magnify(factor=4)]), _png((40, 30)), "factor"),
    ("a crop too small", _doc([], crop={"x": 0, "y": 0, "width": 10, "height": 10}), _png((10, 10)), "crop"),
    ("a crop outside", _doc([], crop={"x": 30, "y": 0, "width": 20, "height": 20}), _png((20, 20)), "past the still's edge"),
    ("a crop with a fraction", _doc([], crop={"x": 0.5, "y": 0, "width": 20, "height": 20}), _png((20, 20)), "crop"),
    ("a png of the wrong size", _doc([_arrow()]), _png((41, 30)), "41x30"),
    ("a png of the still's size under a crop", _doc([], crop={"x": 0, "y": 0, "width": 20, "height": 20}), _png((40, 30)), "20x20"),
    ("a jpeg", _doc([_arrow()]), _jpeg((40, 30)), "PNG"),
    ("png magic over garbage", _doc([_arrow()]), captures.PNG_MAGIC + b"\0" * 64, "PNG"),
    ("no picture at all", _doc([_arrow()]), b"", "PNG"),
    # A header claiming a bomb-sized picture: refused by the size comparison on the IHDR, before Pillow's
    # decompression-bomb check (a bare Exception) could ever see it. 100 bytes, correct CRCs.
    ("a header claiming 20000x10000", _doc([_arrow()]), _crafted_png(20000, 10000), "20000x10000"),
    # A well-formed PNG (chunks and CRCs right, the header the still's size) whose pixel data cannot be
    # decoded: ``verify`` passes it, the decode at save refuses it, so no consumer ever sees it.
    ("a png whose data cannot be decoded", _doc([_arrow()]), _crafted_png(40, 30, b"not-zlib-at-all!"), "not a valid PNG"),
    # The document sent as a FILE part rather than a string field.
    ("the document as a file part", ("file", _doc([_arrow()])), _png((40, 30)), "document"),
]


@pytest.mark.parametrize("label, document, png, said", BAD, ids=[b[0] for b in BAD])
def test_a_bad_document_or_picture_is_refused_with_400_and_nothing_is_written(rig, monkeypatch, label, document, png, said):
    c = rig["admin"]
    cid = _still(rig, monkeypatch, 40, 30)
    r = _put(c, cid, document, png)
    assert r.status_code == 400, f"{label}: {r.status_code} {r.text}"
    detail = r.json()["detail"]
    assert isinstance(detail, str) and said.lower() in detail.lower(), f"{label}: {detail!r}"
    folder = rig["tmp"] / "captures" / cid
    assert not (folder / "annotated.png").exists(), label
    assert json.loads((folder / "capture.json").read_text(encoding="utf-8"))["annotated"] is False, label
    assert c.get(f"/api/captures/{cid}").json()["annotated"] is False
    assert _pixel(c, cid) == RED


def test_a_form_without_its_two_parts_is_refused(rig, monkeypatch):
    c = rig["admin"]
    cid = _still(rig, monkeypatch, 40, 30)
    r = c.put(f"/api/captures/{cid}/annotations", files={"image": ("a.png", _png((40, 30)), "image/png")})
    assert r.status_code == 400 and "document" in r.json()["detail"]
    r = c.put(f"/api/captures/{cid}/annotations", data={"document": json.dumps(_doc([_arrow()]))})
    assert r.status_code == 400 and "image" in r.json()["detail"]
    r = c.put(f"/api/captures/{cid}/annotations", content=b"not a form", headers={"Content-Type": "text/plain"})
    assert r.status_code == 400
    assert not (rig["tmp"] / "captures" / cid / "annotated.png").exists()


def test_a_body_over_the_cap_is_refused_with_413_and_nothing_is_written(rig, monkeypatch):
    """Two refusals at the cap: the whole body by its length before it is read (the document's own
    allowance is lowered to nothing here, so the picture alone is the cap), and the picture part once
    the form is parsed (a cap the body as a whole fits under)."""
    c = rig["admin"]
    cid = _still(rig, monkeypatch, 100, 100)
    noise = Image.effect_noise((100, 100), 64).convert("RGB")
    buf = io.BytesIO()
    noise.save(buf, format="PNG")
    big = buf.getvalue()
    assert len(big) > 2000
    monkeypatch.setattr(captures, "MAX_ANNOTATED_BYTES", 2000)
    # The picture part: the body fits (the document allowance is a megabyte), the picture does not.
    r = _put(c, cid, _doc([_arrow()]), big)
    assert r.status_code == 413, r.text
    assert "may not exceed" in r.json()["detail"]
    # The body: with no document allowance the declared length alone refuses it.
    monkeypatch.setattr(captures, "MAX_DOCUMENT_BYTES", 0)
    r = _put(c, cid, _doc([_arrow()]), big)
    assert r.status_code == 413, r.text
    folder = rig["tmp"] / "captures" / cid
    assert not (folder / "annotated.png").exists()
    assert c.get(f"/api/captures/{cid}").json()["annotated"] is False
    # Under the cap it still saves.
    monkeypatch.setattr(captures, "MAX_ANNOTATED_BYTES", 64 * 1024 * 1024)
    monkeypatch.setattr(captures, "MAX_DOCUMENT_BYTES", 1024 * 1024)
    assert _put(c, cid, _doc([_arrow()]), big).status_code == 200


def test_the_png_check_itself():
    captures.check_png(_png((8, 6)), (8, 6))
    with pytest.raises(ValueError, match="8x6"):
        captures.check_png(_png((8, 6)), (6, 8))
    with pytest.raises(ValueError, match="PNG"):
        captures.check_png(_jpeg((8, 6)), (8, 6))
    with pytest.raises(captures.TooLarge):
        captures.check_png(b"\0" * (captures.MAX_ANNOTATED_BYTES + 1), (8, 6))
    # The bomb-sized header is a plain ValueError naming the header's size, from the IHDR read, never
    # Pillow's DecompressionBombError; the garbage data is a plain ValueError from the decode.
    with pytest.raises(ValueError, match="20000x10000"):
        captures.check_png(_crafted_png(20000, 10000), (8, 6))
    with pytest.raises(ValueError, match="not a valid PNG"):
        captures.check_png(_crafted_png(8, 6, b"not-zlib-at-all!"), (8, 6))
    with pytest.raises(ValueError, match="not a valid PNG"):
        captures.check_png(PNG_MAGIC_BYTES + b"\0\0\0\rIHDR", (8, 6))  # a header cut short


def test_a_still_past_the_pixel_bound_cannot_be_annotated(rig, monkeypatch):
    """The server holds the editor's 100-million-pixel bound on the picture a save would carry (the
    crop's size, else the still's): past it, 400 before any PNG is looked at, and nothing written."""
    c = rig["admin"]
    cid = _still(rig, monkeypatch, 20, 20)
    path = rig["tmp"] / "captures" / cid / "capture.json"
    rec = json.loads(path.read_text(encoding="utf-8"))
    rec["width"], rec["height"] = 20000, 6000
    path.write_text(json.dumps(rec), encoding="utf-8")
    r = _put(c, cid, _doc([_arrow()]), _crafted_png(20000, 6000))
    assert r.status_code == 400, r.text
    assert "million" in r.json()["detail"] and "20000x6000" in r.json()["detail"]
    assert not (rig["tmp"] / "captures" / cid / "annotated.png").exists()
    assert c.get(f"/api/captures/{cid}").json()["annotated"] is False
    # The bound is on the PICTURE: a crop inside it is allowed.
    r = _put(c, cid, _doc([_arrow()], crop={"x": 0, "y": 0, "width": 100, "height": 100}), _png((100, 100)))
    assert r.status_code == 200, r.text


def test_add_to_deck_answers_400_when_the_annotated_picture_cannot_be_read(rig, monkeypatch):
    """The consumer's own robustness: an export damaged on disk after the save (the save itself refuses
    one) is a plain 400 from Add to deck, never a 500, and the deck is untouched."""
    c = rig["admin"]
    cid = _still(rig, monkeypatch, 40, 30)
    deck = c.post("/api/projects/import", files={"file": ("deck.pptx", _deck_bytes(1), "application/octet-stream")}).json()
    assert _put(c, cid, _doc([_arrow()]), _png((40, 30))).status_code == 200
    (rig["tmp"] / "captures" / cid / "annotated.png").write_bytes(_crafted_png(40, 30, b"not-zlib-at-all!"))
    r = c.post(f"/api/captures/{cid}/add-to-deck", json={"project_id": deck["id"]})
    assert r.status_code == 400, r.text
    assert r.json()["detail"] == "The annotated picture cannot be read: revert the annotations and save them again."
    assert c.get(f"/api/projects/{deck['id']}").json()["slide_count"] == 1
    # Revert cures it: the original is untouched.
    assert c.delete(f"/api/captures/{cid}/annotations").status_code == 204
    assert c.post(f"/api/captures/{cid}/add-to-deck", json={"project_id": deck["id"]}).status_code == 200


# -- the missing export ------------------------------------------------------------------------

def test_a_record_that_claims_annotated_without_the_export_serves_the_original_and_says_so(rig, monkeypatch, caplog):
    c = rig["admin"]
    cid = _still(rig, monkeypatch, 40, 30, RED)
    deck = c.post("/api/projects/import", files={"file": ("deck.pptx", _deck_bytes(1), "application/octet-stream")}).json()
    assert _put(c, cid, _doc([_arrow()]), _png((40, 30), BLUE)).status_code == 200
    (rig["tmp"] / "captures" / cid / "annotated.png").unlink()
    with caplog.at_level(logging.WARNING):
        assert _pixel(c, cid) == RED
    assert any("annotated.png is missing" in r.getMessage() and cid in r.getMessage() for r in caplog.records)
    assert c.get(f"/api/captures/{cid}/download").status_code == 200
    assert c.post(f"/api/captures/{cid}/add-to-deck", json={"project_id": deck["id"]}).status_code == 200
    assert c.get(f"/api/captures/{cid}").json()["annotated"] is True, "the record is not rewritten behind the user's back"
    assert captures.image_path(cid, original=True).name == "still.png"


# -- the audit rows ------------------------------------------------------------------------------

def test_the_audit_rows_say_how_many_shapes_and_whether_cropped(rig, monkeypatch):
    c = rig["admin"]
    cid = _still(rig, monkeypatch, 400, 400, name="Dialog")
    assert _put(c, cid, _doc([_arrow(), _step()]), _png((400, 400))).status_code == 200
    assert _put(c, cid, _doc([_callout(text="secret words")], crop={"x": 0, "y": 0, "width": 100, "height": 400}), _png((100, 400))).status_code == 200
    assert c.delete(f"/api/captures/{cid}/annotations").status_code == 204
    rows = auth_store.list_audit(action="capture.annotate")
    assert [r["detail"] for r in rows] == ["1 shape, cropped 100x400", "2 shapes, not cropped"]
    assert all(r["entity"] == "capture" and r["entity_id"] == cid for r in rows)
    assert "secret words" not in json.dumps(rows), "a callout's text never reaches the log"
    reverted = auth_store.list_audit(action="capture.revert")
    assert [r["detail"] for r in reverted] == ["Dialog"]
    # A refused save and a revert of nothing record nothing.
    assert _put(c, cid, _doc([_arrow(x2=9999.0)]), _png((400, 400))).status_code == 400
    assert c.delete(f"/api/captures/{cid}/annotations").status_code == 204
    assert len(auth_store.list_audit(action="capture.annotate")) == 2
    assert len(auth_store.list_audit(action="capture.revert")) == 1


# -- the limits rows ----------------------------------------------------------------------------

def test_the_limits_rows_name_the_annotation_constants():
    """The constants behind the rows (``tests/test_docs.py`` checks every row's line references)."""
    text = LIMITS_MD.read_text(encoding="utf-8")
    rows = [line for line in text.splitlines() if line.startswith("| Annotat")]
    assert rows, "no Annotation rows in docs/reference/limits.md"
    joined = "\n".join(rows)
    for number in (captures.MAX_SHAPES, captures.MAX_TEXT_CHARS, captures.MAX_ANNOTATED_BYTES // 1024 ** 2, captures.MIN_CROP_SIDE):
        assert str(number) in joined, number
    assert "100 million" in joined or "100,000,000" in joined
    for name in ("MAX_SHAPES", "MAX_TEXT_CHARS", "MAX_ANNOTATED_BYTES", "MAX_ANNOTATE_PIXELS"):
        assert name in text, name
