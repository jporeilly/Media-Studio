"""The AI assistant over the slide editor (porting vertical 3b):
``services/ai_slides.py``, its routes, the one-job-per-project gate the AI
jobs share with generate / re-voice / render, the job cancel flag, and the
small extensions of ``services/slides.py`` it stands on.

Ollama is never called: ``core.ollama_client.generate`` / ``check_connection``
/ ``list_models`` are monkeypatched with canned, recorded answers (every
per-slide prompt - notes, enhance, tone, translation, pacing, the QA review -
goes through ``generate``), and the engine's remaining network users (the Q&A
generator, the analyzer's model prose) are replaced at their module
attributes. A real small deck is built with python-pptx (three slides, notes
on two), so the prompts carry the deck's real title and text; the store, the
config and the auth DB are isolated to a tmp dir, as in tests/test_slides.py.
"""

import io
import json
import threading
import time
import urllib.error
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pydantic import ValidationError

from api.schemas import (
    AiEnhanceOneRequest, AiEnhanceRequest, AiNotesRequest, AiQaDocRequest, AiQaFixRequest, AiToneRequest,
    AiTranslateRequest, RenderRequest,
)
from core import ollama_client, qa_generator, slide_analyzer
from core.pptx_exporter import ExportedSlide
from services import ai_slides, jobs, slides
from services import file_item as file_item_module
from services import processing
from services import projects as store
from services import voices as voice_lists
from utils.config import config

NOTES = ("Welcome to the deck.", "", "Third slide notes.")
MODEL = "gemma3:12b"
URL = "http://localhost:11434"


@pytest.fixture(autouse=True)
def isolated(tmp_path, monkeypatch):
    """A throwaway project store, an in-memory config with a vision-capable
    model configured, a throwaway auth DB - and an Ollama that is reachable,
    lists the configured model, and refuses any generate a test did not
    script (so no test can fall through to the network)."""
    monkeypatch.setattr(store, "PROJECTS_DIR", tmp_path / "projects")
    monkeypatch.setattr(config, "_config", {"ollama_model": MODEL})
    monkeypatch.setattr(config, "save", lambda: None)
    monkeypatch.setattr(slides, "_deck_cache", {})
    monkeypatch.setattr(slides, "_locks", {})
    monkeypatch.setattr(processing, "TEMP_DIR", tmp_path / "temp")

    from api import store as auth_store

    monkeypatch.setattr(auth_store, "DATA_DIR", tmp_path / "auth")
    monkeypatch.setattr(auth_store, "DB_PATH", tmp_path / "auth" / "media_studio.db")

    monkeypatch.setattr(ollama_client, "check_connection", lambda base_url=None, timeout=None: True)
    monkeypatch.setattr(
        ollama_client, "list_models", lambda base_url=None, timeout=None: [ollama_client.OllamaModel(MODEL, 1)],
    )

    def refuse(*args, **kwargs):
        raise AssertionError("ollama_client.generate was called without a scripted answer")

    monkeypatch.setattr(ollama_client, "generate", refuse)


@pytest.fixture
def client():
    """A logged-in TestClient (the seeded admin, past its first password change)."""
    from api import store as auth_store
    from api.app import app

    with TestClient(app) as c:
        seeded = auth_store.authenticate("admin", "admin")
        auth_store.change_password(seeded["id"], "admin", must_change=False)
        r = c.post("/api/auth/login", json={"username": "admin", "password": "admin"})
        assert r.status_code == 200
        yield c


def _deck_bytes(notes=NOTES) -> bytes:
    prs = Presentation()
    for i, note in enumerate(notes):
        slide = prs.slides.add_slide(prs.slide_layouts[1])  # title + content
        slide.shapes.title.text = f"Slide {i + 1} title"
        slide.placeholders[1].text = f"Bullet {i + 1}"
        if note:
            slide.notes_slide.notes_text_frame.text = note
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _import_deck(name="deck.pptx", notes=NOTES, owner: dict | None = None) -> str:
    """A deck in the store. With no ``owner`` the record carries none, which
    makes it admin-owned (api/deps.py) - right for the admin `client` fixture;
    a test acting as an editor must name that editor as the owner."""
    return store.import_upload(
        name, _deck_bytes(notes),
        owner_id=owner["id"] if owner else None,
        owner_name=owner["display_name"] if owner else None,
    )["id"]


def _pdf_bytes(pages=2) -> bytes:
    import fitz

    doc = fitz.open()
    for i in range(pages):
        doc.new_page().insert_text((72, 72), f"Page {i + 1}\nSome page text")
    data = doc.tobytes()
    doc.close()
    return data


def _png_bytes() -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (4, 4), (200, 30, 30)).save(buf, "PNG")
    return buf.getvalue()


PNG = _png_bytes()


def _render_files(pid: str) -> None:
    """One PNG per slide, as an export before ``images_source`` existed left them."""
    count = len(slides.list_slides(pid))
    images = store.PROJECTS_DIR / pid / "deck_project" / "images"
    for i in range(count):
        (images / f"slide_{i + 1:03d}.png").write_bytes(PNG)


def _render(pid: str, source: str = "powerpoint") -> None:
    """Pretend the previews were rendered: one PNG per slide, and the source recorded."""
    _render_files(pid)
    slides.record_images_source(pid, source)


def _notes(pid: str) -> list[str]:
    return [s["speaker_notes"] for s in slides.list_slides(pid)]


def _flags(pid: str) -> list[bool]:
    return [s["ai_enhanced"] for s in slides.list_slides(pid)]


class FakeOllama:
    """Stands in for core.ollama_client.generate: answers from ``answer`` (a
    text, or a callable (prompt, call number) -> text; an Exception instance
    is raised) and records every call with its keyword arguments. The jobs'
    warm-up call is answered "OK" and counted apart (``warmups``), so
    ``calls`` and the call numbers are the real prompts only."""

    def __init__(self, answer):
        self.answer = answer
        self.calls = []
        self.warmups = 0

    def __call__(self, prompt, model, system=None, base_url=None, timeout=120.0, images=None):
        if prompt == ai_slides.WARM_UP_PROMPT:
            self.warmups += 1
            return "OK"
        self.calls.append({
            "prompt": prompt, "model": model, "system": system, "base_url": base_url,
            "timeout": timeout, "images": list(images or []),
        })
        out = self.answer(prompt, len(self.calls)) if callable(self.answer) else self.answer
        if isinstance(out, Exception):
            raise out
        return out


class GatedOllama(FakeOllama):
    """A generate that reports its first real call (``entered``) and then
    waits for the test to open ``gate`` - so a test can act while the job is
    inside a model call, ordering threads on events rather than sleeps."""

    def __init__(self, answer, entered: threading.Event, gate: threading.Event):
        super().__init__(answer)
        self.entered, self.gate = entered, gate

    def __call__(self, prompt, *args, **kwargs):
        if prompt != ai_slides.WARM_UP_PROMPT:
            self.entered.set()
            assert self.gate.wait(10), "the test never opened the gate"
        return super().__call__(prompt, *args, **kwargs)


def _script(monkeypatch, answer) -> FakeOllama:
    fake = FakeOllama(answer)
    monkeypatch.setattr(ollama_client, "generate", fake)
    return fake


def _gated(monkeypatch, answer="Written.") -> tuple[GatedOllama, threading.Event, threading.Event]:
    entered, gate = threading.Event(), threading.Event()
    fake = GatedOllama(answer, entered, gate)
    monkeypatch.setattr(ollama_client, "generate", fake)
    return fake, entered, gate


def _http_error(code=500, body=b'{"error":"boom"}') -> urllib.error.HTTPError:
    return urllib.error.HTTPError(f"{URL}/api/generate", code, "Server Error", None, io.BytesIO(body))


def _refused() -> urllib.error.URLError:
    return urllib.error.URLError(ConnectionRefusedError(111, "refused"))


def _wait_job(client, job_id, timeout=10.0) -> dict:
    deadline = time.time() + timeout
    while time.time() < deadline:
        job = client.get(f"/api/jobs/{job_id}").json()
        if job["status"] in ("done", "error"):
            return job
        time.sleep(0.02)
    raise AssertionError("job did not finish in time")


def _run(client, pid, path, body=None) -> dict:
    r = client.post(f"/api/projects/{pid}/ai/{path}", json=body if body is not None else {})
    assert r.status_code == 200, r.text
    return _wait_job(client, r.json()["job_id"])


class _BlockingJob:
    """A job attached to a project that runs until the test releases it."""

    def __init__(self, pid, kind="generate"):
        self.started, self.release = threading.Event(), threading.Event()

        def work(progress):
            self.started.set()
            self.release.wait(10)
            return "ok"

        self.id = jobs.submit(kind, work, project_id=pid)
        assert self.started.wait(5)


class _FakeExporter:
    """Stands in for core.pptx_exporter.PPTXExporter: one PNG per slide; blocks
    on ``gate`` (after setting ``entered``) when a test gives it one."""

    entered: threading.Event | None = None
    gate: threading.Event | None = None

    def __init__(self, pptx_path, output_dir):
        self.pptx_path, self.output_dir = Path(pptx_path), Path(output_dir)
        self.backend = None

    def export_slides_as_images(self, progress_callback=None):
        cls = type(self)
        if cls.entered is not None:
            cls.entered.set()
        if cls.gate is not None:
            assert cls.gate.wait(10)
        out = []
        for i in range(len(Presentation(str(self.pptx_path)).slides)):
            image = self.output_dir / f"slide_{i + 1:03d}.png"
            image.write_bytes(PNG)
            out.append(ExportedSlide(index=i, image_path=image, video_path=None, has_animation=False))
        self.backend = "powerpoint"
        return out


def _fake_exporter(monkeypatch, entered=None, gate=None):
    import core.pptx_exporter as exporter_module

    _FakeExporter.entered, _FakeExporter.gate = entered, gate
    monkeypatch.setattr(exporter_module, "PPTXExporter", _FakeExporter)


class _FakeFileItem:
    def __init__(self, path, projects_base=None):
        self.path = path

    def load(self):
        return True

    def load_pdf(self):
        return True


class _FakeRenderProcessor:
    """A generate that writes a stub MP4 and one subtitle sidecar."""

    def __init__(self, **kwargs):
        self.outputs = {}
        self.images_backend = None

    def process_files(self, files, output_dir, progress=None, preview_seconds=0):
        from utils.helpers import get_output_filename

        out = get_output_filename(files[0].path, output_dir)
        out.write_bytes(b"MP4")
        out.with_suffix(".srt").write_bytes(b"SRT")
        self.outputs = {"srt": out.with_suffix(".srt").name}
        return 1

    def _revoice_video(self, pm, source, out, progress=None):
        out.write_bytes(b"MP4")
        return True


# -- status and the vision gate ------------------------------------------------

def test_status_offline_and_online(client, monkeypatch):
    pid = _import_deck()
    monkeypatch.setattr(ollama_client, "check_connection", lambda base_url=None, timeout=None: False)
    body = client.get(f"/api/projects/{pid}/ai/status").json()
    assert body["available"] is False and body["models"] == [] and body["model_installed"] is None
    assert body["configured_model"] == MODEL and body["base_url"] == URL
    assert body["vision_capable"] is True and body["vision_usable"] is False
    assert "not been rendered yet" in body["vision_reason"]

    monkeypatch.setattr(ollama_client, "check_connection", lambda base_url=None, timeout=None: True)
    body = client.get(f"/api/projects/{pid}/ai/status").json()
    assert body["available"] is True and body["models"] == [MODEL] and body["model_installed"] is True
    assert [t["name"] for t in body["tones"]] == ["Technical", "Executive", "Student", "Sales", "Casual", "Formal", "Custom"]

    config._config["ollama_model"] = "mistral"
    assert client.get(f"/api/projects/{pid}/ai/status").json()["model_installed"] is False


def test_vision_is_usable_only_for_a_vision_model_over_real_renders(client):
    pid = _import_deck()
    assert ai_slides.vision_for_project(pid) == (False, "the slide previews have not been rendered yet — render them first")

    _render(pid, "pillow")
    usable, reason = ai_slides.vision_for_project(pid)
    assert usable is False and "title-only" in reason, "the Pillow fallback is never shown to a model"

    _render(pid, "powerpoint")
    assert ai_slides.vision_for_project(pid) == (True, None)
    body = client.get(f"/api/projects/{pid}/ai/status").json()
    assert body["vision_usable"] is True and body["vision_reason"] is None

    config._config["ollama_model"] = "llama3.1:8b"
    usable, reason = ai_slides.vision_for_project(pid)
    assert usable is False and reason.startswith("llama3.1:8b is not a vision model")
    for name in ("gemma3:12b-it-qat", "llava:7b", "moondream:1.8b", "llama3.2-vision:11b", "qwen2.5vl:7b", "minicpm-v", "bakllava"):
        assert ai_slides.vision_capable(name), name
    assert not ai_slides.vision_capable("mistral") and not ai_slides.vision_capable("")

    (store.PROJECTS_DIR / pid / "deck_project" / "images" / "slide_002.png").unlink()
    config._config["ollama_model"] = MODEL
    assert ai_slides.vision_for_project(pid)[0] is False, "every slide must have its image"


def test_previews_rendered_before_their_source_was_recorded_ask_for_a_render_again(client, monkeypatch):
    """The sample decks: 40 images on disk, no ``images_source`` - vision was
    unreachable (not usable, and the Render button hidden since every image
    exists). Now the reason says so and ``force`` re-exports them."""
    pid = _import_deck()
    _render_files(pid)
    payload = client.get(f"/api/projects/{pid}/slides").json()
    assert payload["slides_ready"] is True and payload["images_source"] is None

    body = client.get(f"/api/projects/{pid}/ai/status").json()
    assert body["vision_usable"] is False
    assert body["vision_reason"] == "the slide previews were rendered before their source was recorded — render them again to enable vision"

    assert client.post(f"/api/projects/{pid}/slides/render").json() == {"cached": True}, "without force, the files are enough"
    assert client.post(f"/api/projects/{pid}/slides/render", json={"forced": True}).status_code == 422

    _fake_exporter(monkeypatch)
    r = client.post(f"/api/projects/{pid}/slides/render", json={"force": True})
    assert r.status_code == 200 and "job_id" in r.json(), r.text
    job = _wait_job(client, r.json()["job_id"])
    assert job["status"] == "done" and job["result"] == {"cached": False, "images_source": "powerpoint", "slides": 3}
    assert client.get(f"/api/projects/{pid}/slides").json()["images_source"] == "powerpoint"
    assert client.get(f"/api/projects/{pid}/ai/status").json()["vision_usable"] is True
    assert RenderRequest().force is False


def test_build_prompt_mirrors_slidestudio():
    notes = ai_slides.build_prompt("Agenda", "One\nTwo", "", image=False, mode="notes")
    assert notes == "Slide title: Agenda\n\nSlide content:\nOne\nTwo\n\n" + ai_slides.NOTES_INSTRUCTION
    with_image = ai_slides.build_prompt("Agenda", None, "", image=True, mode="notes")
    assert with_image.startswith(ai_slides.IMAGE_PREAMBLE["notes"] + "\n\nSlide title: Agenda")

    enhance = ai_slides.build_prompt("Agenda", "One", "Old notes.", image=True, mode="enhance")
    assert enhance == (
        ai_slides.IMAGE_PREAMBLE["enhance"] + "\n\nSlide title: Agenda\n\nSlide content:\nOne\n\n"
        "Existing speaker notes:\nOld notes.\n\n" + ai_slides.ENHANCE_INSTRUCTION
    )
    assert ai_slides.build_prompt(None, None, "", image=False, mode="enhance") == ai_slides.GENERATE_INSTRUCTION
    with pytest.raises(ValueError):
        ai_slides.build_prompt("x", None, "", image=False, mode="other")


def test_engine_prompts_are_rebuilt_verbatim():
    from core.tone_adapter import TONE_PRESETS

    assert ai_slides.tone_prompt("Hi.", "Executive") == TONE_PRESETS["Executive"]["prompt"] + "\n\nOriginal text:\nHi.\n\nRewritten text:"
    assert ai_slides.tone_prompt("Hi.", "Custom", "Make it rhyme.") == "Make it rhyme.\n\nOriginal text:\nHi.\n\nRewritten text:"
    plain = ai_slides.translate_prompt("Hello there.", "French")
    assert plain == (
        "Translate the text below into French. Output only the translation — no preamble, quotes, or explanation."
        "\n\nText:\nHello there.\n\nTranslation:"
    )
    marked = ai_slides.translate_prompt("Hello [pause:1s] there.", "French")
    assert "Keep any bracketed markup such as [pause:1s], [break] or [emphasis] exactly as written" in marked
    assert ai_slides.pacing_prompt("Now, go.") == ai_slides.PACING_INSTRUCTION + "\n\nNow, go."


def test_call_maps_the_failures_to_readable_messages(monkeypatch):
    _script(monkeypatch, _http_error(404, b'{"error":"model \'gemma3:12b\' not found"}'))
    with pytest.raises(ai_slides.AiError, match="404 for model 'gemma3:12b': model 'gemma3:12b' not found"):
        ai_slides._call("hi")
    _script(monkeypatch, TimeoutError())
    with pytest.raises(ai_slides.AiError, match="did not answer within 120 s"):
        ai_slides._call("hi")
    _script(monkeypatch, json.JSONDecodeError("Expecting value", "<html>", 0))
    with pytest.raises(ai_slides.AiError, match="not JSON"):
        ai_slides._call("hi")

    # A dropped connection is "unavailable" only when a probe finds the server gone.
    for dropped in (_refused(), ConnectionResetError(104, "reset"), OSError("socket closed")):
        _script(monkeypatch, dropped)
        with pytest.raises(ai_slides.AiError, match="the server still answers") as info:
            ai_slides._call("hi")
        assert not isinstance(info.value, ai_slides.AiUnavailable)
    monkeypatch.setattr(ollama_client, "check_connection", lambda base_url=None, timeout=None: False)
    for dropped in (_refused(), ConnectionResetError(104, "reset"), OSError("socket closed")):
        _script(monkeypatch, dropped)
        with pytest.raises(ai_slides.AiUnavailable, match=f"not reachable at {URL}"):
            ai_slides._call("hi")


def test_call_sends_the_studio_system_prompt_or_none(monkeypatch):
    fake = _script(monkeypatch, "ok")
    ai_slides._call("hi")
    ai_slides._call("hi", system=ai_slides.NO_SYSTEM)
    config._config["ollama_system_prompt"] = "Be brief."
    ai_slides._call("hi")
    assert [c["system"] for c in fake.calls] == [ollama_client.DEFAULT_SYSTEM_PROMPT, "", "Be brief."]


# -- generate notes / enhance ----------------------------------------------------

def test_generate_notes_job_fills_the_empty_slides(client, monkeypatch):
    pid = _import_deck()
    fake = _script(monkeypatch, lambda prompt, n: f"Generated notes {n}.")

    job = _run(client, pid, "notes", {})
    assert job["status"] == "done" and job["kind"] == "ai-notes" and job["project_id"] == pid
    assert job["result"] == {
        "total": 1, "done": 1, "failed": 0, "skipped": 0, "unchanged": 0, "errors": [], "cancelled": False,
        "vision": False, "vision_reason": None,
    }
    assert _notes(pid) == ["Welcome to the deck.", "Generated notes 1.", "Third slide notes."]
    assert _flags(pid) == [False, True, False]
    assert slides.list_slides(pid)[1]["notes_history_depth"] == 1, "an AI write is an edit like any other"

    assert fake.warmups == 1 and len(fake.calls) == 1, "the model is warmed once, then one call per slide"
    call = fake.calls[0]
    assert call["prompt"] == "Slide title: Slide 2 title\n\nSlide content:\nBullet 2\n\n" + ai_slides.NOTES_INSTRUCTION
    assert call["model"] == MODEL and call["base_url"] == URL and call["timeout"] == 120.0
    assert call["system"] == ollama_client.DEFAULT_SYSTEM_PROMPT and call["images"] == []
    assert job["message"] == "Generate notes: 1 of 1 slides"
    assert job["user_id"], "the job remembers who started it"

    r = client.post(f"/api/projects/{pid}/ai/notes", json={})
    assert r.status_code == 400 and r.json()["detail"] == "Every slide already has notes."


def test_notes_job_attaches_images_only_when_vision_is_usable(client, monkeypatch):
    pid = _import_deck()
    fake = _script(monkeypatch, "Seen.")

    _render(pid, "pillow")
    job = _run(client, pid, "notes", {"use_vision": True, "scope": "all", "slide_indexes": [1]})
    assert job["result"]["vision"] is False and "title-only" in job["result"]["vision_reason"]
    assert fake.calls[-1]["images"] == [] and "showing you the slide image" not in fake.calls[-1]["prompt"]

    _render(pid, "powerpoint")
    job = _run(client, pid, "notes", {"use_vision": True, "scope": "all", "slide_indexes": [1, 1, 0]})
    assert job["result"]["vision"] is True and job["result"]["done"] == 2, "duplicate indexes are sent once"
    assert [Path(c["images"][0]).name for c in fake.calls[-2:]] == ["slide_002.png", "slide_001.png"]
    assert fake.calls[-1]["prompt"].startswith(ai_slides.IMAGE_PREAMBLE["notes"])

    r = client.post(f"/api/projects/{pid}/ai/notes", json={"scope": "all", "slide_indexes": [7]})
    assert r.status_code == 400 and "No slide with index 7" in r.json()["detail"]


def test_enhance_all_counts_failures_and_keeps_going(client, monkeypatch):
    pid = _import_deck()

    def answer(prompt, n):
        return _http_error(500) if n == 2 else f"Better {n}."

    fake = _script(monkeypatch, answer)
    job = _run(client, pid, "enhance", {})
    assert job["status"] == "done" and job["kind"] == "ai-enhance"
    result = job["result"]
    assert (result["total"], result["done"], result["failed"], result["skipped"]) == (3, 2, 1, 0)
    assert result["errors"] == ["Slide 2: Ollama answered 500 for model 'gemma3:12b': boom"]
    assert _notes(pid) == ["Better 1.", "", "Better 3."] and _flags(pid) == [True, False, True]
    assert job["message"] == "Enhance: 2 of 3 slides, 1 failed"

    prompts = [c["prompt"] for c in fake.calls]
    assert "Existing speaker notes:\nWelcome to the deck.\n\n" + ai_slides.ENHANCE_INSTRUCTION in prompts[0]
    assert prompts[1].endswith(ai_slides.GENERATE_INSTRUCTION), "an empty slide gets notes from its content"
    assert prompts[1].startswith("Slide title: Slide 2 title")
    assert AiEnhanceRequest().scope == "all" and AiNotesRequest().scope == "empty"


def test_enhance_stops_when_ollama_goes_away_and_keeps_what_it_wrote(client, monkeypatch):
    pid = _import_deck()
    up = {"value": True}
    monkeypatch.setattr(ollama_client, "check_connection", lambda base_url=None, timeout=None: up["value"])

    def answer(prompt, n):
        if n == 2:
            up["value"] = False  # the server dies mid-job: the probe after the dropped call finds it gone
            return _refused()
        return "Kept."

    _script(monkeypatch, answer)
    job = _run(client, pid, "enhance", {"scope": "all"})
    assert job["status"] == "error"
    assert "Ollama is not reachable" in job["error"] and "Stopped after 1 of 3 slides" in job["error"]
    assert _notes(pid)[0] == "Kept." and _notes(pid)[2] == "Third slide notes."


def test_a_dropped_call_while_the_server_answers_is_one_failed_slide(client, monkeypatch):
    pid = _import_deck()
    _script(monkeypatch, lambda prompt, n: _refused() if n == 1 else "Fine.")
    job = _run(client, pid, "enhance", {"scope": "all"})
    assert job["status"] == "done" and job["result"]["failed"] == 1 and job["result"]["done"] == 2
    assert "the server still answers" in job["result"]["errors"][0]


def test_a_model_that_will_not_load_fails_the_job_before_the_first_slide(client, monkeypatch):
    pid = _import_deck()

    def generate(prompt, model, system=None, base_url=None, timeout=120.0, images=None):
        raise _http_error(404, b'{"error":"model not found"}')

    monkeypatch.setattr(ollama_client, "generate", generate)
    job = _run(client, pid, "enhance", {})
    assert job["status"] == "error" and job["error"].startswith("The model could not be loaded: Ollama answered 404")
    assert _notes(pid) == list(NOTES)


def test_offline_refuses_to_start_a_job_and_no_job_exists(client, monkeypatch):
    pid = _import_deck()
    monkeypatch.setattr(ollama_client, "check_connection", lambda base_url=None, timeout=None: False)
    r = client.post(f"/api/projects/{pid}/ai/enhance", json={})
    assert r.status_code == 503 and "Ollama is not reachable" in r.json()["detail"]
    assert jobs.active_for(pid) is None


def test_cancel_stops_a_job_between_slides(client, monkeypatch):
    pid = _import_deck()
    fake, entered, gate = _gated(monkeypatch)

    job_id = client.post(f"/api/projects/{pid}/ai/enhance", json={}).json()["job_id"]
    assert entered.wait(5)
    r = client.post(f"/api/jobs/{job_id}/cancel")
    assert r.status_code == 200 and r.json()["cancel_requested"] is True and r.json()["status"] == "running"
    gate.set()

    job = _wait_job(client, job_id)
    assert job["status"] == "done" and job["message"] == "Cancelled"
    assert job["result"]["cancelled"] is True and job["result"]["done"] == 1 and job["result"]["total"] == 3
    assert len(fake.calls) == 1, "no call after the flag was seen"
    assert _notes(pid)[0] == "Written." and _notes(pid)[2] == "Third slide notes."
    assert jobs.active_for(pid) is None
    assert client.post(f"/api/jobs/{job_id}/cancel").json()["status"] == "done", "a finished job is returned as it is"
    assert client.post("/api/jobs/nope/cancel").status_code == 404


def test_only_the_starter_or_an_admin_may_cancel(client, monkeypatch):
    from api import store as auth_store
    from api.app import app

    editor_account = auth_store.create_user("editor", "Editor-pass-1", "Editor", role="editor", must_change_password=False)
    with TestClient(app) as editor:
        assert editor.post("/api/auth/login", json={"username": "editor", "password": "Editor-pass-1"}).status_code == 200

        # The admin's job: the editor may not cancel it, the admin may.
        pid = _import_deck()
        fake, entered, gate = _gated(monkeypatch)
        job_id = client.post(f"/api/projects/{pid}/ai/enhance", json={}).json()["job_id"]
        assert entered.wait(5)
        r = editor.post(f"/api/jobs/{job_id}/cancel")
        assert r.status_code == 403 and "Only the user who started this job" in r.json()["detail"]
        assert jobs.get(job_id)["cancel_requested"] is False
        # Nor may they WATCH it: a job carries its project id, its progress
        # messages and the names of the files it produced, so reading someone
        # else's was the last cross-tenant leak once projects gained owners.
        assert editor.get(f"/api/jobs/{job_id}").status_code == 403
        assert client.get(f"/api/jobs/{job_id}").status_code == 200, "the starter still reads their own"
        assert client.post(f"/api/jobs/{job_id}/cancel").status_code == 200
        gate.set()
        assert _wait_job(client, job_id)["result"]["cancelled"] is True

        # The editor's own job on the editor's own project: the editor may
        # cancel it, and so may the admin.
        other = _import_deck("other.pptx", owner=editor_account)
        fake, entered, gate = _gated(monkeypatch)
        job_id = editor.post(f"/api/projects/{other}/ai/enhance", json={}).json()["job_id"]
        assert entered.wait(5)
        assert client.get(f"/api/jobs/{job_id}").status_code == 200, "an admin reads anyone's"
        assert client.post(f"/api/jobs/{job_id}/cancel").status_code == 200, "an admin cancels anyone's"
        gate.set()
        assert _wait_job(client, job_id)["result"]["cancelled"] is True
        fake, entered, gate = _gated(monkeypatch)
        job_id = editor.post(f"/api/projects/{other}/ai/enhance", json={}).json()["job_id"]
        assert entered.wait(5)
        assert editor.post(f"/api/jobs/{job_id}/cancel").status_code == 200, "the starter cancels their own"
        gate.set()
        _wait_job(client, job_id)


def test_one_job_per_project_across_generate_render_and_ai(client, monkeypatch):
    """A render started over a running AI job would save its stale copy of the
    notes over everything the AI loop wrote: every job start on a busy
    project is the same 409, whichever kind holds it."""
    monkeypatch.setattr(processing, "VideoProcessor", _FakeRenderProcessor)
    monkeypatch.setattr(file_item_module, "FileItem", _FakeFileItem)
    pid = _import_deck()
    starts = (
        lambda: client.post(f"/api/projects/{pid}/generate", json={}),
        lambda: client.post(f"/api/projects/{pid}/slides/render"),
        lambda: client.post(f"/api/projects/{pid}/slides/render", json={"force": True}),
        lambda: client.post(f"/api/projects/{pid}/ai/notes", json={}),
        lambda: client.post(f"/api/projects/{pid}/ai/enhance", json={}),
        lambda: client.post(f"/api/projects/{pid}/ai/qa"),
        lambda: client.post(f"/api/projects/{pid}/ai/tone", json={"tone": "Casual"}),
        lambda: client.post(f"/api/projects/{pid}/ai/translate", json={"language": "French"}),
        lambda: client.post(f"/api/projects/{pid}/ai/pacing", json={"use_ai": True}),
        lambda: client.post(f"/api/projects/{pid}/ai/pacing", json={}),
        lambda: client.post(f"/api/projects/{pid}/ai/qa-doc", json={"num_questions": 3}),
        lambda: client.post(f"/api/projects/{pid}/slides/0/ai/qa-fix", json={"criterion": "grammar", "issue": "typo"}),
        lambda: client.patch(f"/api/projects/{pid}/slides/0", json={"speaker_notes": "x"}),
    )

    # An AI job holds the project.
    fake, entered, gate = _gated(monkeypatch)
    ai_job = client.post(f"/api/projects/{pid}/ai/notes", json={}).json()["job_id"]
    assert entered.wait(5)
    try:
        for start in starts:
            r = start()
            assert r.status_code == 409 and "ai-notes" in r.json()["detail"], r.text
        assert client.get(f"/api/projects/{pid}/ai/status").status_code == 200, "reads are fine"
    finally:
        gate.set()
    assert _wait_job(client, ai_job)["status"] == "done"
    assert jobs.active_for(pid) is None

    # A generate job holds the project.
    blocking = _BlockingJob(pid, "generate")
    try:
        for start in starts:
            r = start()
            assert r.status_code == 409 and "generate" in r.json()["detail"], r.text
    finally:
        blocking.release.set()
    assert _wait_job(client, blocking.id)["status"] == "done"

    # A render holds it (an AI start is refused, a second render joins it).
    entered, gate = threading.Event(), threading.Event()
    _fake_exporter(monkeypatch, entered=entered, gate=gate)
    render = client.post(f"/api/projects/{pid}/slides/render").json()["job_id"]
    assert entered.wait(5)
    try:
        assert client.post(f"/api/projects/{pid}/ai/notes", json={}).status_code == 409
        assert client.post(f"/api/projects/{pid}/generate", json={}).status_code == 409
        assert client.post(f"/api/projects/{pid}/slides/render").json() == {"job_id": render}
    finally:
        gate.set()
    assert _wait_job(client, render)["status"] == "done"

    # Idle again: every start is a 200.
    _script(monkeypatch, "Fine.")
    assert _wait_job(client, client.post(f"/api/projects/{pid}/ai/notes", json={"scope": "all"}).json()["job_id"])["status"] == "done"
    assert _wait_job(client, client.post(f"/api/projects/{pid}/generate", json={}).json()["job_id"])["status"] == "done"
    assert client.post(f"/api/projects/{pid}/slides/render").json() == {"cached": True}
    assert client.patch(f"/api/projects/{pid}/slides/0", json={"speaker_notes": "x"}).status_code == 200


def test_one_job_per_project_for_re_voice_and_transcribe(client, monkeypatch):
    monkeypatch.setattr(processing, "VideoProcessor", _FakeRenderProcessor)
    video = store.import_upload("clip.mp4", b"mp4")["id"]
    store.set_transcript(video, [{"start": 0.0, "end": 1.0, "text": "Hello."}])
    blocking = _BlockingJob(video, "transcribe")
    try:
        r = client.post(f"/api/projects/{video}/revoice", json={"voice_id": "en-US-AriaNeural"})
        assert r.status_code == 409 and "transcribe" in r.json()["detail"]
        assert client.post(f"/api/projects/{video}/transcribe", json={}).status_code == 409
    finally:
        blocking.release.set()
    assert _wait_job(client, blocking.id)["status"] == "done"
    r = client.post(f"/api/projects/{video}/revoice", json={"voice_id": "en-US-AriaNeural"})
    assert r.status_code == 200, r.text
    assert _wait_job(client, r.json()["job_id"])["status"] == "done"


def test_start_is_atomic_under_one_lock():
    pid = _import_deck()
    blocking = _BlockingJob(pid, "generate")
    try:
        with pytest.raises(jobs.ProjectBusy) as info:
            jobs.start("ai-notes", lambda progress: None, project_id=pid)
        assert info.value.job["id"] == blocking.id and "generate" in str(info.value)
        with pytest.raises(jobs.ProjectBusy):
            jobs.require_idle(pid)
        assert jobs.start("generate", lambda progress: None, project_id=pid, reuse=True) == blocking.id, "reuse joins the same kind"
        with pytest.raises(jobs.ProjectBusy):
            jobs.start("render-slides", lambda progress: None, project_id=pid, reuse=True)
    finally:
        blocking.release.set()
    while jobs.active_for(pid):
        time.sleep(0.01)
    jobs.require_idle(pid)


# -- the QA review -------------------------------------------------------------------

def test_parse_qa_json_direct_fenced_and_braces():
    payload = {"score": 8, "slides": [{"slide": 1, "grammar": "None"}]}
    assert ai_slides.parse_qa_json(json.dumps(payload)) == payload
    assert ai_slides.parse_qa_json("Sure!\n```json\n" + json.dumps(payload) + "\n```\nDone.") == payload
    assert ai_slides.parse_qa_json("Here you go: " + json.dumps(payload) + " hope it helps") == payload
    assert ai_slides.parse_qa_json("no json here") == {}
    assert ai_slides.parse_qa_json("{broken") == {}


def test_qa_review_maps_by_slide_number_and_persists(client, monkeypatch):
    pid = _import_deck()
    # The model answers OUT OF ORDER: slide 3 first. The numbers, not the positions, decide.
    answer = {
        "score": "8.0",
        "slides": [
            {"slide": 3, "grammar": "None", "tone": "Too casual", "flow": "ok", "transitions": "No link to slide 2"},
            {"slide": "1", "grammar": "'deck' should be 'Deck'", "tone": "none", "flow": None, "transitions": " N/A "},
        ],
    }
    fake = _script(monkeypatch, "```json\n" + json.dumps(answer) + "\n```")

    job = _run(client, pid, "qa")
    assert job["status"] == "done" and job["kind"] == "ai-qa"
    assert job["result"] == {"score": 8, "issues": 3, "slides": 2, "passes": 1, "cancelled": False, "saved": True}
    assert job["message"] == "QA score 8/10 — 3 issue(s)"

    review = client.get(f"/api/projects/{pid}/slides").json()["qa_review"]
    assert review["score"] == 8 and review["model"] == MODEL and review["issues"] == 3 and review["run_at"]
    assert review["slides"] == [
        {"index": 0, "grammar": "'deck' should be 'Deck'", "tone": "None", "flow": "None", "transitions": "None", "fixed": []},
        {"index": 2, "grammar": "None", "tone": "Too casual", "flow": "None", "transitions": "No link to slide 2", "fixed": []},
    ], "the empty slide 2 is not reviewed; every no-issue spelling becomes 'None'"
    assert _notes(pid) == list(NOTES), "a review changes no notes"

    assert fake.warmups == 1 and len(fake.calls) == 1
    call = fake.calls[0]
    assert call["timeout"] == 180.0 and call["system"] == ai_slides.QA_SYSTEM
    assert call["prompt"].startswith("Review the following speaker notes across 2 slides.")
    assert '{"slide": 1, "grammar"' in call["prompt"] and '{"slide": 3, "grammar"' in call["prompt"]
    assert "--- Slide 1 ---\nWelcome to the deck." in call["prompt"] and "--- Slide 3 ---\nThird slide notes." in call["prompt"]


def test_qa_review_maps_positionally_only_without_slide_numbers(client, monkeypatch):
    pid = _import_deck()
    _script(monkeypatch, json.dumps({"score": 9, "slides": [{"grammar": "first"}, {"grammar": "second"}]}))
    assert _run(client, pid, "qa")["status"] == "done"
    entries = store.get_project(pid)["qa_review"]["slides"]
    assert [(e["index"], e["grammar"]) for e in entries] == [(0, "first"), (2, "second")]

    # One number missing, one present: not every entry carries one, so positional again.
    _script(monkeypatch, json.dumps({"slides": [{"slide": 3, "grammar": "a"}, {"grammar": "b"}]}))
    assert _run(client, pid, "qa")["status"] == "done"
    entries = store.get_project(pid)["qa_review"]["slides"]
    assert [(e["index"], e["grammar"]) for e in entries] == [(0, "a"), (2, "b")]
    assert ai_slides._slide_number(3.0) == 3 and ai_slides._slide_number(" 7 ") == 7
    assert ai_slides._slide_number(True) is None and ai_slides._slide_number(2.5) is None and ai_slides._slide_number("x") is None


def test_qa_review_refuses_mismatched_answers_and_keeps_the_previous_review(client, monkeypatch):
    pid = _import_deck()
    previous = {"score": 5, "slides": [], "issues": 0, "run_at": "then", "model": "old"}
    record = store.get_project(pid)
    record["qa_review"] = previous
    store.save_project(record)

    # Same length, but slide 1 dropped and slide 3 duplicated: refused rather than mis-mapped.
    _script(monkeypatch, json.dumps({"score": 9, "slides": [{"slide": 3, "grammar": "x"}, {"slide": 3, "grammar": "y"}]}))
    job = _run(client, pid, "qa")
    assert job["status"] == "error"
    assert "slide numbers (3, 3) are not the slides sent (1, 3)" in job["error"] and "nothing was saved" in job["error"]
    assert store.get_project(pid)["qa_review"] == previous, "the previous verdict stays until a run succeeds"

    _script(monkeypatch, json.dumps({"score": 9, "slides": [{"slide": 1, "grammar": "typo"}]}))
    job = _run(client, pid, "qa")
    assert job["status"] == "error" and "returned 1 entries for 2 slides" in job["error"]
    assert store.get_project(pid)["qa_review"] == previous

    _script(monkeypatch, "not json at all")
    job = _run(client, pid, "qa")
    assert job["status"] == "error" and "no 'slides' array" in job["error"]
    assert store.get_project(pid)["qa_review"] == previous

    # The old dict shape is read by its values, in order.
    _script(monkeypatch, json.dumps({"slides": {"a": {"grammar": "x"}, "b": {"grammar": "y"}}}))
    job = _run(client, pid, "qa")
    assert job["status"] == "done" and job["result"]["issues"] == 2
    assert job["result"]["score"] == 7, "no score from the model, issues found: inferred 7"
    assert store.get_project(pid)["qa_review"]["score"] == 7


def test_qa_review_chunks_a_large_deck_and_averages_the_scores(client, monkeypatch):
    notes = tuple(f"Notes for slide {i + 1}." for i in range(25))
    pid = _import_deck(notes=notes)
    fake = _script(monkeypatch, lambda prompt, n: json.dumps({
        "score": 8 if n == 1 else 6,
        "slides": [{"grammar": "None", "tone": "None", "flow": "None", "transitions": "None"}] * (13 if n == 1 else 12),
    }))

    job = _run(client, pid, "qa")
    assert job["status"] == "done", job
    assert job["result"] == {"score": 7, "issues": 0, "slides": 25, "passes": 2, "cancelled": False, "saved": True}
    assert len(fake.calls) == 2 and fake.warmups == 1
    assert fake.calls[0]["prompt"].startswith("Review the following speaker notes across 13 slides.")
    assert fake.calls[1]["prompt"].startswith("Review the following speaker notes across 12 slides.")
    assert "--- Slide 13 ---" in fake.calls[0]["prompt"] and "--- Slide 14 ---" in fake.calls[1]["prompt"]
    review = store.get_project(pid)["qa_review"]
    assert [e["index"] for e in review["slides"]] == list(range(25))

    assert ai_slides.chunked(list(range(38)), 20) == [list(range(19)), list(range(19, 38))]
    assert [len(c) for c in ai_slides.chunked(list(range(41)), 20)] == [14, 14, 13]
    assert [len(c) for c in ai_slides.chunked(list(range(20)), 20)] == [20]
    assert ai_slides.chunked([], 20) == []


def test_qa_passes_are_capped_by_characters_too(client, monkeypatch):
    long = ["x" * 5000, "y" * 5000, "z" * 5000, "w" * 100]
    assert ai_slides.chunked(long, 20, 12_000) == [["x" * 5000, "y" * 5000], ["z" * 5000, "w" * 100]]
    assert ai_slides.chunked(["a" * 20_000, "b"], 20, 12_000) == [["a" * 20_000], ["b"]], "one oversize note is its own pass"
    pairs = [(i, text) for i, text in enumerate(long)]
    assert [[i for i, _ in c] for c in ai_slides.chunked(pairs, 20, 12_000, text=lambda p: p[1])] == [[0, 1], [2, 3]]

    pid = _import_deck(notes=tuple(long))
    fake = _script(monkeypatch, lambda prompt, n: json.dumps({
        "score": 9, "slides": [{"grammar": "None"}] * prompt.count("--- Slide "),
    }))
    job = _run(client, pid, "qa")
    assert job["status"] == "done" and job["result"]["passes"] == 2 and job["result"]["slides"] == 4
    assert [c["prompt"].count("--- Slide ") for c in fake.calls] == [2, 2]


def test_qa_review_cancels_between_passes_and_saves_nothing(client, monkeypatch):
    notes = tuple(f"Notes for slide {i + 1}." for i in range(25))
    pid = _import_deck(notes=notes)
    fake, entered, gate = _gated(monkeypatch, lambda prompt, n: json.dumps({
        "score": 9, "slides": [{"grammar": "None"}] * prompt.count("--- Slide "),
    }))
    job_id = client.post(f"/api/projects/{pid}/ai/qa").json()["job_id"]
    assert entered.wait(5)
    assert client.post(f"/api/jobs/{job_id}/cancel").status_code == 200
    gate.set()
    job = _wait_job(client, job_id)
    assert job["status"] == "done" and job["message"] == "Cancelled"
    assert job["result"] == {"cancelled": True, "passes": 1, "total_passes": 2, "slides": 25, "saved": False}
    assert len(fake.calls) == 1 and store.get_project(pid).get("qa_review") is None


def test_qa_score_is_inferred_when_missing_and_no_issue_spellings_are_the_sentinel(client, monkeypatch):
    pid = _import_deck()
    clean = [{"grammar": "None", "tone": "-", "flow": "no issues", "transitions": "OK"}] * 2
    _script(monkeypatch, json.dumps({"slides": clean}))
    job = _run(client, pid, "qa")
    assert job["result"]["score"] == 10 and job["result"]["issues"] == 0
    assert ai_slides.qa_score("9.6") == 9 and ai_slides.qa_score(42) == 10 and ai_slides.qa_score(None) is None
    assert ai_slides.qa_score("great") is None and ai_slides.qa_score(True) is None
    assert ai_slides.issue_text("  Good ") == "None" and ai_slides.issue_text(7) == "7"
    assert not ai_slides.is_issue("n/a") and ai_slides.is_issue("Missing comma")


def test_qa_fix_saves_the_slide_and_marks_the_criterion_fixed(client, monkeypatch):
    pid = _import_deck()
    review = [
        {"slide": 1, "grammar": "'deck' should be 'Deck'", "tone": "Too casual", "flow": "None", "transitions": "None"},
        {"slide": 3, "grammar": "None", "tone": "None", "flow": "None", "transitions": "None"},
    ]
    _script(monkeypatch, json.dumps({"score": 7, "slides": review}))
    assert _run(client, pid, "qa")["status"] == "done"

    fake = _script(monkeypatch, "Welcome to the Deck.")
    r = client.post(f"/api/projects/{pid}/slides/0/ai/qa-fix", json={"criterion": "grammar", "issue": "'deck' should be 'Deck'"})
    assert r.status_code == 200, r.text
    assert r.json()["speaker_notes"] == "Welcome to the Deck." and r.json()["ai_enhanced"] is True
    assert _notes(pid)[0] == "Welcome to the Deck."
    saved = store.get_project(pid)["qa_review"]
    assert saved["slides"][0]["fixed"] == ["grammar"] and saved["issues"] == 1

    call = fake.calls[0]
    assert call["system"] == ai_slides.QA_FIX_SYSTEM and fake.warmups == 0, "a sync call: no warm-up"
    assert call["prompt"].startswith("QA Review found this issue with Slide 1:\nCategory: Grammar\nIssue: 'deck' should be 'Deck'\n\n")
    assert "Current speaker notes:\nWelcome to the deck.\n\n" in call["prompt"]
    assert ai_slides.QA_CONSTRAINTS["grammar"] + "\nReturn ONLY the improved notes text, nothing else." in call["prompt"]

    # Fixing it again is idempotent on the review; an unknown criterion is a 422; an empty slide a 400.
    client.post(f"/api/projects/{pid}/slides/0/ai/qa-fix", json={"criterion": "grammar", "issue": "again"})
    assert store.get_project(pid)["qa_review"]["slides"][0]["fixed"] == ["grammar"]
    assert client.post(f"/api/projects/{pid}/slides/0/ai/qa-fix", json={"criterion": "length", "issue": "x"}).status_code == 422
    r = client.post(f"/api/projects/{pid}/slides/1/ai/qa-fix", json={"criterion": "tone", "issue": "x"})
    assert r.status_code == 400 and "no notes to fix" in r.json()["detail"]

    _script(monkeypatch, "   ")
    r = client.post(f"/api/projects/{pid}/slides/2/ai/qa-fix", json={"criterion": "flow", "issue": "x"})
    assert r.status_code == 502 and "returned nothing" in r.json()["detail"]
    assert _notes(pid)[2] == "Third slide notes."


def test_qa_fix_on_an_unreviewed_slide_saves_without_a_review(client, monkeypatch):
    pid = _import_deck()
    _script(monkeypatch, "Third slide notes, fixed.")
    r = client.post(f"/api/projects/{pid}/slides/2/ai/qa-fix", json={"criterion": "flow", "issue": "abrupt"})
    assert r.status_code == 200 and r.json()["speaker_notes"] == "Third slide notes, fixed."
    assert store.get_project(pid).get("qa_review") is None, "no review is invented"
    assert ai_slides._mark_fixed(pid, 2, "flow") is None


# -- the per-slide proposal ------------------------------------------------------------

def test_enhance_one_is_a_proposal_that_saves_nothing(client, monkeypatch):
    pid = _import_deck()
    fake = _script(monkeypatch, "A better draft.")
    r = client.post(f"/api/projects/{pid}/slides/0/ai/enhance", json={"notes": "My unsaved draft"})
    assert r.status_code == 200, r.text
    assert r.json() == {"suggestion": "A better draft.", "vision": False, "vision_reason": None}
    assert _notes(pid)[0] == "Welcome to the deck." and _flags(pid)[0] is False
    assert "Existing speaker notes:\nMy unsaved draft\n\n" in fake.calls[0]["prompt"], "the editor's text, not the saved one"

    r = client.post(f"/api/projects/{pid}/slides/2/ai/enhance", json={})
    assert r.status_code == 200 and "Existing speaker notes:\nThird slide notes." in fake.calls[1]["prompt"]

    r = client.post(f"/api/projects/{pid}/slides/1/ai/enhance", json={"use_vision": True})
    assert r.status_code == 200 and "not been rendered" in r.json()["vision_reason"] and fake.calls[2]["images"] == []
    _render(pid, "powerpoint")
    r = client.post(f"/api/projects/{pid}/slides/1/ai/enhance", json={"use_vision": True})
    assert r.json()["vision"] is True and Path(fake.calls[3]["images"][0]).name == "slide_002.png"
    assert fake.warmups == 0

    assert client.post(f"/api/projects/{pid}/slides/9/ai/enhance", json={}).status_code == 400
    monkeypatch.setattr(ollama_client, "generate", FakeOllama(TimeoutError()))
    r = client.post(f"/api/projects/{pid}/slides/0/ai/enhance", json={})
    assert r.status_code == 502 and "did not answer within" in r.json()["detail"]


def test_slide_sources_open_the_inner_project_once(client, monkeypatch):
    pid = _import_deck()
    _render(pid, "powerpoint")
    opened = []
    real_open = slides._open

    def counting_open(p):
        opened.append(p)
        return real_open(p)

    monkeypatch.setattr(slides, "_open", counting_open)
    sources = ai_slides.slide_sources(pid)
    assert [s.index for s in sources] == [0, 1, 2] and all(s.image and s.image.name == f"slide_{s.index + 1:03d}.png" for s in sources)
    assert len(opened) == 1


# -- tone, translate, pacing, Q&A document, analyze ------------------------------------

def test_tone_job_sends_the_engines_prompt_per_slide(client, monkeypatch):
    pid = _import_deck()
    fake = _script(monkeypatch, lambda prompt, n: f"[Executive] {n}")
    job = _run(client, pid, "tone", {"tone": "Executive"})
    assert job["status"] == "done" and job["kind"] == "ai-tone"
    assert (job["result"]["total"], job["result"]["done"], job["result"]["skipped"]) == (2, 2, 0), "only slides with notes"
    assert _notes(pid) == ["[Executive] 1", "", "[Executive] 2"] and _flags(pid) == [True, False, True]
    assert fake.warmups == 1
    assert [c["prompt"] for c in fake.calls] == [
        ai_slides.tone_prompt("Welcome to the deck.", "Executive"), ai_slides.tone_prompt("Third slide notes.", "Executive"),
    ]
    assert all(c["system"] == "" and c["timeout"] == 120.0 and c["model"] == MODEL for c in fake.calls)

    fake = _script(monkeypatch, "Rhymed.")
    job = _run(client, pid, "tone", {"tone": "Custom", "custom_prompt": "  Make it rhyme. "})
    assert job["status"] == "done" and fake.calls[0]["prompt"].startswith("Make it rhyme.\n\nOriginal text:\n[Executive] 1")

    r = client.post(f"/api/projects/{pid}/ai/tone", json={"tone": "Custom"})
    assert r.status_code == 400 and "instruction" in r.json()["detail"]
    assert client.post(f"/api/projects/{pid}/ai/tone", json={"tone": "Pirate"}).status_code == 422


def test_rewrite_loop_counts_failures_and_stops_when_the_server_is_gone(client, monkeypatch):
    pid = _import_deck()
    # A failed call is a failed slide; an answer identical to the note is "unchanged" (and nothing else is).
    _script(monkeypatch, lambda prompt, n: _http_error(500) if n == 1 else "Third slide notes.")
    job = _run(client, pid, "tone", {"tone": "Casual"})
    assert job["status"] == "done"
    assert (job["result"]["done"], job["result"]["failed"], job["result"]["unchanged"]) == (0, 1, 1)
    assert job["result"]["errors"] == ["Slide 1: Ollama answered 500 for model 'gemma3:12b': boom"]
    assert _notes(pid) == list(NOTES)

    _script(monkeypatch, "   ")
    job = _run(client, pid, "translate", {"language": "German", "match_voice": False})
    assert job["result"]["failed"] == 2 and job["result"]["errors"][0] == "Slide 1: the model returned nothing"

    up = {"value": True}
    monkeypatch.setattr(ollama_client, "check_connection", lambda base_url=None, timeout=None: up["value"])

    def answer(prompt, n):
        if n == 2:
            up["value"] = False
            return ConnectionResetError(104, "reset")
        return "Rewritten."

    _script(monkeypatch, answer)
    job = _run(client, pid, "translate", {"language": "German"})
    assert job["status"] == "error" and "Stopped after 1 of 2 slides" in job["error"]
    assert _notes(pid) == ["Rewritten.", "", "Third slide notes."]


def _edge_voices(*locales):
    return {
        "provider": "edge_tts", "error": None, "notice": None,
        "voices": [{"voice_id": f"{loc}-{name}Neural", "name": f"{name} ({loc})", "locale": loc, "gender": None}
                   for loc, name in locales],
    }


def test_translate_job_replaces_notes_and_suggests_a_matching_voice(client, monkeypatch):
    pid = _import_deck()
    fake = _script(monkeypatch, lambda prompt, n: f"French: {n}")
    asked = []

    def fake_voices(provider):
        asked.append(provider)
        return _edge_voices(("en-US", "Aria"), ("fr-CA", "Sylvie"), ("fr-FR", "Denise"))

    monkeypatch.setattr(voice_lists, "voices_for", fake_voices)

    job = _run(client, pid, "translate", {"language": "French"})
    assert job["status"] == "done" and job["kind"] == "ai-translate"
    result = job["result"]
    assert result["done"] == 2 and result["language"] == "French" and result["subtag"] == "fr"
    assert result["provider"] == "edge_tts" and result["match_voice"] is True and asked == ["edge_tts"]
    assert result["suggested_voice_id"] == "fr-CA-SylvieNeural" and result["suggested_voice_name"] == "Sylvie (fr-CA)", \
        "the first voice whose locale's language matches, as suggest_voice_for_language picks it"
    assert _notes(pid) == ["French: 1", "", "French: 2"]
    assert _flags(pid) == [True, False, True]
    assert [c["prompt"] for c in fake.calls] == [
        ai_slides.translate_prompt("Welcome to the deck.", "French"), ai_slides.translate_prompt("Third slide notes.", "French"),
    ]
    assert all(c["system"] == "" for c in fake.calls) and fake.warmups == 1

    # No voice for the language: null, the caller keeps its voice. The provider of the request is asked.
    monkeypatch.setattr(voice_lists, "voices_for", lambda provider: asked.append(provider) or _edge_voices(("en-US", "Aria")))
    job = _run(client, pid, "translate", {"language": "German", "provider": "kokoro"})
    assert job["result"]["suggested_voice_id"] is None and job["result"]["provider"] == "kokoro" and asked[-1] == "kokoro"

    job = _run(client, pid, "translate", {"language": "Spanish", "match_voice": False})
    assert job["result"]["suggested_voice_id"] is None and asked[-1] == "kokoro", "not asked when not wanted"

    assert client.post(f"/api/projects/{pid}/ai/translate", json={"language": "Klingon"}).status_code == 422
    assert client.post(f"/api/projects/{pid}/ai/translate", json={"language": "French", "provider": "polly"}).status_code == 400


def test_suggest_voice_matches_the_locale_language(monkeypatch):
    monkeypatch.setattr(voice_lists, "voices_for", lambda provider: _edge_voices(("en-GB", "Ryan"), ("de-DE", "Katja")))
    assert ai_slides.suggest_voice("de", "edge_tts")["voice_id"] == "de-DE-KatjaNeural"
    assert ai_slides.suggest_voice("fr", "edge_tts") is None
    kokoro = {"provider": "kokoro", "error": None, "notice": None, "voices": [
        {"voice_id": "af_heart", "name": "Heart", "locale": "en-US", "gender": "Female"},
        {"voice_id": "ff_siwis", "name": "Siwis", "locale": "fr-FR", "gender": "Female"},
        {"voice_id": "xx_odd", "name": "Odd", "locale": "", "gender": None},
    ]}
    monkeypatch.setattr(voice_lists, "voices_for", lambda provider: kokoro)
    assert ai_slides.suggest_voice("fr", "kokoro")["voice_id"] == "ff_siwis"
    assert ai_slides.suggest_voice("xx", "kokoro") is None, "an unknown locale never matches by accident"
    monkeypatch.setattr(voice_lists, "voices_for", lambda provider: {"provider": "edge_tts", "voices": [], "error": "offline", "notice": None})
    assert ai_slides.suggest_voice("en", "edge_tts") is None


def test_pacing_rules_run_now_and_ai_pacing_is_a_job(client, monkeypatch):
    pid = _import_deck(notes=("Now, we grew 50 percent. Any questions?", "", "Plain text without triggers"))
    r = client.post(f"/api/projects/{pid}/ai/pacing", json={})
    assert r.status_code == 200, r.text
    result = r.json()
    assert result["done"] == 1 and result["skipped"] == 1 and result["unchanged"] == 1 and result["adjustments"] >= 2
    paced = _notes(pid)[0]
    assert "Now, ..." in paced and "..." in paced and _notes(pid)[2] == "Plain text without triggers"
    assert _flags(pid) == [False, False, False], "rules are nobody's AI rewrite"
    assert jobs.active_for(pid) is None, "no job for the rules"

    def answer(prompt, n):
        note = prompt.split("\n\n", 1)[1]
        return "x" if n == 1 else f"{note} ... end"  # the first answer lost words: the rules pace it instead

    fake = _script(monkeypatch, answer)
    job = _run(client, pid, "pacing", {"use_ai": True})
    assert job["status"] == "done" and job["kind"] == "ai-pacing" and job["result"]["done"] == 2
    assert fake.calls[0]["prompt"] == ai_slides.pacing_prompt(paced) and fake.warmups == 1
    assert _notes(pid)[2] == "Plain text without triggers ... end" and _flags(pid) == [True, False, True]
    assert _notes(pid)[0] != "x" and "..." in _notes(pid)[0], "a short answer never replaces the note; the rules pace it"


def test_qa_doc_job_writes_the_document_and_the_export_route_serves_it(client, monkeypatch):
    pid = _import_deck()
    assert client.get(f"/api/projects/{pid}/export/qa").status_code == 404
    fake = _script(monkeypatch, "unused")  # the warm-up only; the Q&A generator itself is faked
    seen = {}

    def fake_generate_qa(notes, url, model, num_questions=10, on_progress=None):
        seen.update({"notes": list(notes), "url": url, "model": model, "n": num_questions})
        if on_progress:
            on_progress(0.5, "thinking")
        return [{"question": "Why?", "answer": "Because.", "slide_ref": 1}, {"question": "How?", "answer": "So.", "slide_ref": 0}]

    monkeypatch.setattr(qa_generator, "generate_qa", fake_generate_qa)
    job = _run(client, pid, "qa-doc", {"num_questions": 2})
    assert job["status"] == "done" and job["kind"] == "ai-qa-doc"
    assert job["result"] == {"pairs": 2, "qa_doc": "exports/deck-qa.txt", "download": f"/api/projects/{pid}/export/qa"}
    assert seen == {"notes": list(NOTES), "url": URL, "model": MODEL, "n": 2}
    assert fake.warmups == 1 and fake.calls == []
    assert store.get_project(pid)["outputs"] == {"qa_doc": "exports/deck-qa.txt"}
    text = (store.PROJECTS_DIR / pid / "exports" / "deck-qa.txt").read_text(encoding="utf-8")
    assert text.startswith("Q&A Document: deck") and "1. Q: Why?" in text and "(Reference: Slide 1)" in text

    r = client.get(f"/api/projects/{pid}/export/qa")
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/plain")
    assert r.headers["content-disposition"] == 'attachment; filename="deck-qa.txt"'
    assert r.content == (store.PROJECTS_DIR / pid / "exports" / "deck-qa.txt").read_bytes()

    monkeypatch.setattr(qa_generator, "generate_qa", lambda *a, **k: [])
    job = _run(client, pid, "qa-doc", {})
    assert job["status"] == "error" and "no question/answer pairs" in job["error"]
    assert store.get_project(pid)["outputs"] == {"qa_doc": "exports/deck-qa.txt"}, "the earlier document stays"

    for bad in ({"num_questions": 0}, {"num_questions": 51}, {"num_questions": "5"}, {"questions": 5}):
        assert client.post(f"/api/projects/{pid}/ai/qa-doc", json=bad).status_code == 422, bad


def test_export_qa_refuses_a_path_outside_the_project(client):
    pid = _import_deck()
    outside = store.PROJECTS_DIR / "outside.txt"
    outside.parent.mkdir(parents=True, exist_ok=True)
    outside.write_text("secret")
    record = store.get_project(pid)
    record["outputs"] = {"qa_doc": "../outside.txt"}
    store.save_project(record)
    assert client.get(f"/api/projects/{pid}/export/qa").status_code == 404


def test_the_qa_document_survives_a_full_render(client, monkeypatch):
    monkeypatch.setattr(processing, "VideoProcessor", _FakeRenderProcessor)
    monkeypatch.setattr(file_item_module, "FileItem", _FakeFileItem)
    pid = _import_deck()
    record = store.get_project(pid)
    record["outputs"] = {"qa_doc": "exports/deck-qa.txt", "preview": "deck_preview.mp4", "gif": "old.gif"}
    store.save_project(record)

    job = _wait_job(client, client.post(f"/api/projects/{pid}/generate", json={}).json()["job_id"])
    assert job["status"] == "done", job
    assert store.get_project(pid)["outputs"] == {"qa_doc": "exports/deck-qa.txt", "preview": "deck_preview.mp4", "srt": "deck.srt"}


def test_analyze_runs_the_rules_offline_and_adds_the_models_prose_online(client, monkeypatch):
    pid = _import_deck()
    slides.update_slide(pid, 0, speaker_notes="Edited narration for slide one, quite a few words in here now.")
    monkeypatch.setattr(ollama_client, "check_connection", lambda base_url=None, timeout=None: False)
    monkeypatch.setattr(slide_analyzer, "_get_ai_suggestions", lambda *a, **k: pytest.fail("no model offline"))

    r = client.post(f"/api/projects/{pid}/ai/analyze")
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["ai"] is False and body["model"] is None and body["suggestions"] == []
    assert 1 <= body["overall_score"] <= 10 and body["summary"]
    assert [s["index"] for s in body["slides"]] == [0, 1, 2]
    assert any("No speaker notes" in issue for issue in body["slides"][1]["issues"]), "the empty slide"
    assert not any("No speaker notes" in issue for issue in body["slides"][0]["issues"]), "the editor's notes, not the deck's"

    monkeypatch.setattr(ollama_client, "check_connection", lambda base_url=None, timeout=None: True)
    monkeypatch.setattr(slide_analyzer, "_get_ai_suggestions", lambda data, url, model: [f"Add visuals ({model} at {url})"])
    body = client.post(f"/api/projects/{pid}/ai/analyze").json()
    assert body["ai"] is True and body["model"] == MODEL
    assert body["suggestions"] == [f"Add visuals ({MODEL} at {URL})"]

    pdf = store.import_upload("pages.pdf", _pdf_bytes(2))["id"]
    body = client.post(f"/api/projects/{pdf}/ai/analyze").json()
    assert [s["index"] for s in body["slides"]] == [0, 1] and body["summary"]


def test_analyze_never_echoes_the_readers_exception(client, monkeypatch):
    pid = _import_deck()
    (store.PROJECTS_DIR / pid / "deck.pptx").write_bytes(b"not a deck C:\\secret\\path")
    r = client.post(f"/api/projects/{pid}/ai/analyze")
    assert r.status_code == 400 and r.json()["detail"] == "The deck could not be read; the server log has the reason."
    pdf = store.import_upload("pages.pdf", b"%PDF-broken")["id"]
    r = client.post(f"/api/projects/{pdf}/ai/analyze")
    assert r.status_code == 400 and r.json()["detail"] == "The PDF could not be read; the server log has the reason."


def test_pdf_slide_sources_carry_the_page_text(client, monkeypatch):
    pdf = store.import_upload("pages.pdf", _pdf_bytes(2))["id"]
    sources = ai_slides.slide_sources(pdf)
    assert [(s.index, s.title, s.body) for s in sources] == [(0, "Page 1", "Some page text"), (1, "Page 2", "Some page text")]
    assert all(s.image is not None and s.image.name == f"page_{s.index + 1:03d}.png" for s in sources)

    fake = _script(monkeypatch, "Narrate the page.")
    job = _run(client, pdf, "notes", {"use_vision": True})
    assert job["status"] == "done" and job["result"]["done"] == 2 and job["result"]["vision"] is True, "a PDF's page renders are real"
    assert fake.calls[0]["prompt"].startswith(ai_slides.IMAGE_PREAMBLE["notes"] + "\n\nSlide title: Page 1\n\nSlide content:\nSome page text")


# -- guards, schemas, the slides.py extensions ----------------------------------------------

def test_routes_need_a_signed_in_user_and_a_deck_or_pdf(client):
    pid = _import_deck()
    video = store.import_upload("clip.mp4", b"mp4")["id"]
    for method, path, body in (
        ("get", f"/api/projects/{video}/ai/status", None),
        ("post", f"/api/projects/{video}/ai/notes", {}),
        ("post", f"/api/projects/{video}/ai/analyze", None),
        ("post", f"/api/projects/{video}/slides/0/ai/enhance", {}),
        ("get", f"/api/projects/{video}/export/qa", None),
    ):
        r = getattr(client, method)(path, json=body) if body is not None else getattr(client, method)(path)
        assert r.status_code == 400, (path, r.text)
    assert client.get("/api/projects/000000000000/ai/status").status_code == 404
    assert client.post("/api/projects/000000000000/ai/qa").status_code == 404

    with TestClient(client.app) as anonymous:
        assert anonymous.get(f"/api/projects/{pid}/ai/status").status_code == 401
        assert anonymous.post(f"/api/projects/{pid}/ai/enhance", json={}).status_code == 401
        assert anonymous.post("/api/jobs/x/cancel").status_code == 401


def test_request_schemas_bound_and_forbid_unknown_keys():
    assert AiNotesRequest().model_dump() == {"scope": "empty", "slide_indexes": None, "use_vision": False}
    assert AiNotesRequest(scope="all", slide_indexes=[0, 3]).slide_indexes == [0, 3]
    for bad in ({"scope": "some"}, {"slide_indexes": [-1]}, {"slide_indexes": [True]}, {"vision": True}):
        with pytest.raises(ValidationError):
            AiNotesRequest(**bad)
    assert AiEnhanceRequest().model_dump() == {"scope": "all", "slide_indexes": None, "use_vision": False}
    assert AiToneRequest(tone="Casual").custom_prompt is None
    for bad in ({"tone": "Loud"}, {"tone": "Casual", "custom_prompt": "x" * 2001}, {"tone": "Casual", "extra": 1}):
        with pytest.raises(ValidationError):
            AiToneRequest(**bad)
    assert AiTranslateRequest(language="French").model_dump() == {"language": "French", "match_voice": True, "provider": None}
    with pytest.raises(ValidationError):
        AiTranslateRequest(language="Elvish")
    assert AiQaDocRequest().num_questions == 10
    for bad in ({"num_questions": 0}, {"num_questions": 51}, {"num_questions": 2.5}, {"num_questions": True}):
        with pytest.raises(ValidationError):
            AiQaDocRequest(**bad)
    assert AiQaFixRequest(criterion="flow", issue="x").criterion == "flow"
    with pytest.raises(ValidationError):
        AiQaFixRequest(criterion="length", issue="x")
    with pytest.raises(ValidationError):
        AiEnhanceOneRequest(notes="x" * (slides.MAX_NOTES_CHARS + 1))
    with pytest.raises(ValidationError):
        RenderRequest(force=True, again=True)


def test_update_slide_sets_the_ai_flag_and_reset_clears_it():
    pid = _import_deck()
    updated = slides.update_slide(pid, 0, speaker_notes="AI text.", ai_enhanced=True)
    assert updated["ai_enhanced"] is True and updated["speaker_notes"] == "AI text." and updated["notes_history_depth"] == 1
    assert slides.update_slide(pid, 0, ai_enhanced=True)["ai_enhanced"] is True, "idempotent"
    assert slides.update_slide(pid, 0, ai_enhanced=False)["ai_enhanced"] is False
    assert json.loads((store.PROJECTS_DIR / pid / "deck_project" / "project.json").read_text())["slides"][0]["ai_enhanced"] is False
    with pytest.raises(ValueError):
        slides.update_slide(pid, 0, ai_enhanced="yes")

    slides.update_slide(pid, 0, ai_enhanced=True)
    reset = slides.reset_slide(pid, 0)
    assert reset["speaker_notes"] == "Welcome to the deck." and reset["ai_enhanced"] is False
    assert slides.list_slides(pid)[0]["ai_enhanced"] is False, "persisted"
    # A reset of a slide already on the deck's text still drops the flag, with a write of its own.
    slides.update_slide(pid, 2, ai_enhanced=True)
    assert slides.reset_slide(pid, 2)["ai_enhanced"] is False and slides.list_slides(pid)[2]["ai_enhanced"] is False


def test_no_notes_is_a_400_before_any_job(client):
    pid = _import_deck(notes=("", "", ""))
    for path, body in (("qa", None), ("tone", {"tone": "Casual"}), ("translate", {"language": "French"}),
                       ("pacing", {"use_ai": True}), ("qa-doc", {})):
        r = client.post(f"/api/projects/{pid}/ai/{path}", json=body)
        assert r.status_code == 400 and "no speaker notes" in r.json()["detail"], (path, r.text)
    assert jobs.active_for(pid) is None


def test_cancel_requested_here_is_false_outside_a_job():
    assert jobs.current_job_id() is None
    assert jobs.cancel_requested_here() is False
    assert jobs.cancel("missing") is None
