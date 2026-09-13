"""PowerPoint reader for extracting slides and speaker notes."""

from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional
from pptx import Presentation


@dataclass
class SlideInfo:
    """Information about a single slide."""
    index: int  # 0-based index
    speaker_notes: str
    title: Optional[str] = None
    body_text: Optional[str] = None  # all non-title text from the slide
    thumbnail_path: Optional[Path] = None


class PPTXReader:
    """Reads PowerPoint files and extracts slide information."""

    def __init__(self, pptx_path: Path):
        self.pptx_path = Path(pptx_path)
        self._presentation = None
        self._slides_info: List[SlideInfo] = []

    def load(self) -> bool:
        """Load the PowerPoint file."""
        try:
            self._presentation = Presentation(str(self.pptx_path))
            self._extract_slides_info()
            return True
        except Exception as e:
            print(f"Error loading PowerPoint: {e}")
            return False

    def _extract_slides_info(self):
        """Extract information from all slides."""
        self._slides_info = []

        for idx, slide in enumerate(self._presentation.slides):
            # Extract speaker notes
            notes = ""
            if slide.has_notes_slide:
                notes_slide = slide.notes_slide
                if notes_slide.notes_text_frame:
                    notes = notes_slide.notes_text_frame.text.strip()

            # Extract title if available
            title = None
            if slide.shapes.title:
                title = slide.shapes.title.text

            # Extract all non-title text from shapes
            body_parts = []
            for shape in slide.shapes:
                if shape == slide.shapes.title:
                    continue
                if shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    if text:
                        body_parts.append(text)
                if shape.has_table:
                    for row in shape.table.rows:
                        row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_text:
                            body_parts.append(" | ".join(row_text))
            body_text = "\n".join(body_parts) if body_parts else None

            self._slides_info.append(SlideInfo(
                index=idx,
                speaker_notes=notes,
                title=title,
                body_text=body_text,
            ))

    @property
    def slide_count(self) -> int:
        """Get the number of slides."""
        return len(self._slides_info)

    @property
    def slides(self) -> List[SlideInfo]:
        """Get all slide information."""
        return self._slides_info

    def get_slide(self, index: int) -> Optional[SlideInfo]:
        """Get information for a specific slide."""
        if 0 <= index < len(self._slides_info):
            return self._slides_info[index]
        return None

    def get_speaker_notes(self, index: int) -> str:
        """Get speaker notes for a specific slide."""
        slide = self.get_slide(index)
        return slide.speaker_notes if slide else ""

    def has_speaker_notes(self) -> bool:
        """Check if any slide has speaker notes."""
        return any(slide.speaker_notes for slide in self._slides_info)

    @staticmethod
    def export_with_updated_notes(
        source_pptx: Path, updated_notes: List[str], output_path: Path,
    ) -> Path:
        """Save a copy of the PPTX with updated speaker notes.

        Creates a new .pptx file with all slides, formatting, and media
        preserved but speaker notes replaced with the edited text.

        Args:
            source_pptx: Path to the original .pptx file.
            updated_notes: List of note strings, one per slide (index-aligned).
            output_path: Where to save the new .pptx file.

        Returns:
            Path to the saved file.
        """
        prs = Presentation(str(source_pptx))

        for idx, slide in enumerate(prs.slides):
            if idx >= len(updated_notes):
                break
            new_text = updated_notes[idx]

            # Ensure the slide has a notes slide
            if not slide.has_notes_slide:
                slide.notes_slide  # accessing creates it if missing

            notes_slide = slide.notes_slide
            tf = notes_slide.notes_text_frame
            # Clear existing paragraphs and set new text
            tf.clear()
            tf.text = new_text

        output_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(output_path))
        return output_path
