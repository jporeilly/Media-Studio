"""Batch video generation from PPTX templates with variable substitution.

Replaces {{variable}} placeholders in speaker notes and slide text with values
from a CSV/dict, generating personalised videos at scale.
"""

from pathlib import Path
from typing import List, Dict
import csv
import re
import shutil
import logging

logger = logging.getLogger("mediastudio.BATCH")

def extract_variables(pptx_path: Path) -> List[str]:
    """Scan a PPTX file for {{variable}} placeholders in speaker notes and slide text.
    Returns a sorted list of unique variable names found.
    """
    from pptx import Presentation
    variables = set()
    prs = Presentation(str(pptx_path))
    for slide in prs.slides:
        # Check speaker notes
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            text = slide.notes_slide.notes_text_frame.text
            variables.update(re.findall(r'\{\{(\w+)\}\}', text))
        # Check slide text shapes
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        variables.update(re.findall(r'\{\{(\w+)\}\}', run.text))
    return sorted(variables)


def load_csv_data(csv_path: Path) -> List[Dict[str, str]]:
    """Load variable values from a CSV file.
    First row is headers (variable names), each subsequent row is a set of values.
    Returns list of dicts.
    """
    with open(csv_path, 'r', encoding='utf-8-sig') as f:
        reader = csv.DictReader(f)
        return list(reader)


def apply_variables(pptx_path: Path, variables: Dict[str, str], output_path: Path) -> Path:
    """Create a copy of the PPTX with all {{variable}} placeholders replaced.

    Replaces in both speaker notes and slide text shapes.
    Returns path to the new PPTX.
    """
    from pptx import Presentation

    shutil.copy2(str(pptx_path), str(output_path))
    prs = Presentation(str(output_path))

    for slide in prs.slides:
        # Replace in speaker notes
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            for para in slide.notes_slide.notes_text_frame.paragraphs:
                for run in para.runs:
                    for key, value in variables.items():
                        run.text = run.text.replace(f'{{{{{key}}}}}', value)

        # Replace in slide text shapes
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        for key, value in variables.items():
                            run.text = run.text.replace(f'{{{{{key}}}}}', value)

    prs.save(str(output_path))
    return output_path


def generate_batch_plan(pptx_path: Path, csv_path: Path) -> List[Dict]:
    """Generate a batch plan: list of dicts with 'variables', 'output_name' for each row.

    Output name is derived from the first variable value or row index.
    """
    rows = load_csv_data(csv_path)
    template_vars = extract_variables(pptx_path)
    plan = []

    for i, row in enumerate(rows):
        # Use first column value as identifier, fallback to index
        first_val = list(row.values())[0] if row else f"row_{i+1}"
        safe_name = re.sub(r'[^\w\s-]', '', str(first_val)).strip().replace(' ', '_')
        plan.append({
            'index': i,
            'variables': row,
            'output_name': f"{pptx_path.stem}_{safe_name}",
            'missing': [v for v in template_vars if v not in row],
        })

    return plan
