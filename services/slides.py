"""Per-slide state of a deck or PDF project: the engine's own store, exposed.

The engine keeps every slide's state (notes, per-slide overrides, undo
history, rendered image and audio) in the INNER project of a Media Studio
project, ``data/projects/<pid>/<stem>_project/project.json``
(``core.project_manager.ProjectManager``) - but only created it when the first
generate job ran. This module is the accessor layer the slide editor's API
reads and writes it through:

- ``manager(pid)`` opens that inner project and MATERIALISES it when absent:
  a deck's notes through ``PPTXReader`` (no slide images are exported - that
  is ``ensure_images``, a job), a PDF through ``FileItem.load_pdf`` (its page
  images come with the import, as before). A project.json that exists but
  cannot be read is never recreated: ``ProjectManager.load`` raises
  ``ProjectStateError`` (the routes answer 409). The outer record gets its
  ``slide_count`` reconciled and ``engine_project_dir`` (the directory NAME
  under the project, never an absolute path).
- ``list_slides`` / ``update_slide`` / ``undo_slide`` / ``reset_slide`` /
  ``bulk_update`` go through the manager's own mutators (history, the
  ``needs_regeneration`` flag) under a per-project lock: ``project.json`` has
  none of its own and two requests may edit the same project at once. A job
  attached to the project (generate, render) is refused by the routes
  (``services.jobs.active_for``), since it holds its own copy of the state.
- ``ensure_images`` exports a deck's slide images behind ``SLIDE_EXPORT_LOCK``.
  PowerPoint COM is a single instance and ``services.jobs`` runs two workers,
  so every slide export in this process (here and the generate pipeline's,
  ``services.processing``) queues on this one lock and says so in its
  progress while it waits.
- ``image_path`` and ``export_notes_pptx`` derive their files from the
  project directory and the slide index; the absolute paths ``project.json``
  stores are never trusted for serving and never returned (``list_slides``
  carries no paths at all).

``images_source`` on the outer record says where the slide images came from:
``powerpoint`` (real renders), ``pillow`` (the title-only fallback - vision
features must skip it), ``pdf`` (page renders), or None when unknown.
"""

import math
import threading
import uuid
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path

from core.project_manager import ProjectManager, SlideRenderState, get_project_dir
from services import projects as store
from services import studio_settings

# One slide export at a time, process-wide (PowerPoint COM is single-instance).
SLIDE_EXPORT_LOCK = threading.Lock()
WAITING_FOR_EXPORT = "Waiting for PowerPoint to finish another export…"

MAX_NOTES_CHARS = 20_000
MAX_PAUSE_SECONDS = 30.0
SLIDE_KINDS = ("deck", "pdf")
IMAGE_SOURCES = ("powerpoint", "pillow", "pdf")

# The exporters' file names by project kind, 1-based: PPTXExporter writes
# slide_NNN.png, PDFReader page_NNN.png. Derived from the index, never read
# back from the stored absolute string.
_IMAGE_NAME = {"deck": "slide_{n:03d}.png", "pdf": "page_{n:03d}.png"}

_UNSET = object()

_locks: dict[str, threading.Lock] = {}
_locks_guard = threading.Lock()

# The deck's own notes / title / body text per project, keyed on the source
# file's (mtime, size) so a re-read happens only when the file changes.
_deck_cache: dict[str, tuple[tuple[int, int], list]] = {}


class ProjectNotFound(LookupError):
    """No project with that id (the routes answer 404). A dedicated class, so
    an IndexError or KeyError from a malformed record is not mistaken for it."""


def project_lock(pid: str) -> threading.Lock:
    """The lock every read-modify-write of one project's inner project.json takes."""
    with _locks_guard:
        lock = _locks.get(pid)
        if lock is None:
            lock = _locks[pid] = threading.Lock()
        return lock


def forget(pid: str) -> None:
    """Drop the per-project lock and the deck cache (the project was deleted)."""
    with _locks_guard:
        _locks.pop(pid, None)
    _deck_cache.pop(pid, None)


@contextmanager
def slide_export_lock(progress=None, fraction: float = 0.0, label: str = ""):
    """Hold ``SLIDE_EXPORT_LOCK`` for a slide export.

    When another export holds it the wait is reported through ``progress``
    first (``WAITING_FOR_EXPORT``, prefixed with ``label`` when given), so a
    queued job does not look stuck.
    """
    if not SLIDE_EXPORT_LOCK.acquire(blocking=False):
        if progress:
            progress(fraction, f"{label}: {WAITING_FOR_EXPORT}" if label else WAITING_FOR_EXPORT)
        SLIDE_EXPORT_LOCK.acquire()
    try:
        yield
    finally:
        SLIDE_EXPORT_LOCK.release()


# -- opening the inner project --------------------------------------------

def _record(pid: str) -> dict:
    """The outer record of a deck or PDF project. ``ProjectNotFound`` when
    there is no such project, ``ValueError`` for a video project."""
    record = store.get_project(pid)
    if not record:
        raise ProjectNotFound("Project not found.")
    if record.get("kind") not in SLIDE_KINDS:
        raise ValueError("Only deck and PDF projects have slides.")
    return record


def project_record(pid: str) -> dict:
    """The outer record of a deck or PDF project, for the layers built on this
    one (the AI assistant): ``ProjectNotFound`` / ``ValueError`` as ``_record``."""
    return _record(pid)


def source_path(pid: str) -> Path:
    """The project's source file (the deck or the PDF)."""
    record = _record(pid)
    return _source(pid, record)


def _source(pid: str, record: dict) -> Path:
    return store.PROJECTS_DIR / pid / record["source_filename"]


def _materialise(pid: str, record: dict, source: Path, pm: ProjectManager) -> ProjectManager:
    """Create the inner project from the source file, the way the generate
    path would, but without exporting a deck's images. Only ever called when
    there is no project.json (an unreadable one raised before this)."""
    if record["kind"] == "pdf":
        # The PDF path renders the page images as part of reading the file
        # (there is nothing else to render them from) and creates the project.
        from services.file_item import FileItem

        item = FileItem(source, projects_base=store.PROJECTS_DIR / pid)
        if not item.load_pdf() or item.project_manager is None or item.project_manager.state is None:
            raise ValueError("The PDF could not be read.")
        record_images_source(pid, "pdf")
        return item.project_manager

    from core.pptx_reader import PPTXReader

    reader = PPTXReader(source)
    if not reader.load():
        raise ValueError("The deck could not be read.")
    pm.create_project(
        pptx_path=source,
        slide_notes=[slide.speaker_notes for slide in reader.slides],
        voice_id="",
    )
    return pm


def _reconcile(pid: str, record: dict, pm: ProjectManager) -> None:
    """Keep the outer record's ``slide_count`` and ``engine_project_dir`` in
    step with the inner project (saved onto the record as it is now)."""
    count = len(pm.state.slides)
    inner = pm.project_dir.name
    if record.get("slide_count") == count and record.get("engine_project_dir") == inner:
        return
    current = store.get_project(pid) or record
    current["slide_count"] = count
    current["engine_project_dir"] = inner
    store.save_project(current)
    record.update(current)


def _open(pid: str) -> tuple[dict, ProjectManager]:
    """The outer record and the loaded (or freshly materialised) inner project.
    Callers hold ``project_lock(pid)``. Raises ``ProjectStateError`` for a
    project.json that exists but cannot be read - it is never recreated."""
    record = _record(pid)
    source = _source(pid, record)
    if not source.is_file():
        raise ValueError("The project's source file is missing.")
    pm = ProjectManager(get_project_dir(source, store.PROJECTS_DIR / pid))
    if pm.project_file.exists():
        pm.load()
    if pm.state is None:
        pm = _materialise(pid, record, source, pm)
    _reconcile(pid, record, pm)
    _adopt_images(record, pm)
    return record, pm


def manager(pid: str) -> ProjectManager:
    """The inner project of a deck or PDF project, materialised if absent."""
    with project_lock(pid):
        return _open(pid)[1]


# -- reading ---------------------------------------------------------------

def _deck_info(pid: str, record: dict) -> list:
    """The deck's own slides (``SlideInfo``: notes, title, body text), read
    fresh when the source file changed and cached otherwise. [] for a PDF."""
    if record["kind"] != "deck":
        return []
    source = _source(pid, record)
    stat = source.stat()
    key = (stat.st_mtime_ns, stat.st_size)
    cached = _deck_cache.get(pid)
    if cached and cached[0] == key:
        return cached[1]

    from core.pptx_reader import PPTXReader

    reader = PPTXReader(source)
    info = list(reader.slides) if reader.load() else []
    _deck_cache[pid] = (key, info)
    return info


def _image_file(record: dict, pm: ProjectManager, index: int) -> Path:
    return pm.images_dir / _IMAGE_NAME[record["kind"]].format(n=index + 1)


def _images_ready(record: dict, pm: ProjectManager) -> bool:
    slides = pm.state.slides
    return bool(slides) and all(_image_file(record, pm, s.index).is_file() for s in slides)


def _adopt_images(record: dict, pm: ProjectManager) -> None:
    """Point a slide's stored ``image_path`` at the derived file when that file
    exists and the stored path does not (a data directory moved between
    machines), so the generate pipeline's own "images already present" check
    agrees with ``slides_ready`` instead of exporting again. Saves only when
    something changed."""
    changed = False
    for slide in pm.state.slides:
        if slide.image_path and Path(slide.image_path).is_file():
            continue
        derived = _image_file(record, pm, slide.index)
        if derived.is_file() and slide.image_path != str(derived):
            slide.image_path = str(derived)
            changed = True
    if changed:
        pm.save()


def _serialise(pid: str, record: dict, pm: ProjectManager) -> list[dict]:
    info = _deck_info(pid, record)
    out = []
    for slide in pm.state.slides:
        deck = info[slide.index] if slide.index < len(info) else None
        out.append({
            "index": slide.index,
            "speaker_notes": slide.speaker_notes or "",
            "original_notes": deck.speaker_notes if deck else "",
            "title": deck.title if deck else None,
            "body_text": deck.body_text if deck else None,
            "has_image": _image_file(record, pm, slide.index).is_file(),
            "has_audio": bool(slide.audio_path) and Path(slide.audio_path).is_file(),
            "ai_enhanced": bool(slide.ai_enhanced),
            "needs_regeneration": bool(slide.needs_regeneration),
            "voice_override": slide.voice_override or None,
            "pause_override": slide.pause_override,
            "alt_text": slide.alt_text or None,
            "notes_history_depth": len(slide.notes_history or []),
            "has_animation": bool(slide.has_animation),
        })
    return out


def list_slides(pid: str) -> list[dict]:
    """Every slide's state merged with the deck's title / body text and the
    deck's own notes (``original_notes``). No paths, absolute or otherwise."""
    with project_lock(pid):
        record, pm = _open(pid)
        return _serialise(pid, record, pm)


def slides_payload(pid: str) -> dict:
    """What ``GET /slides`` answers: the slides plus where the images came from
    (``powerpoint`` / ``pillow`` / ``pdf``, None = unknown), whether every
    slide has one, and the last QA review (``services.ai_slides``; None until
    one has run)."""
    with project_lock(pid):
        record, pm = _open(pid)
        source = record.get("images_source")
        review = record.get("qa_review")
        return {
            "slides": _serialise(pid, record, pm),
            "images_source": source if source in IMAGE_SOURCES else None,
            "images_rendered_at": record.get("images_rendered_at"),
            "slides_ready": _images_ready(record, pm),
            "qa_review": review if isinstance(review, dict) else None,
        }


def images_ready(pid: str) -> bool:
    with project_lock(pid):
        record, pm = _open(pid)
        return _images_ready(record, pm)


def slide_context(pid: str) -> list[dict]:
    """Server-side only (the AI prompts): every slide as ``list_slides`` gives
    it plus ``image``, the rendered file's path when it exists under the
    project (None otherwise) - one open of the inner project for the whole
    deck. The paths never leave the server."""
    with project_lock(pid):
        record, pm = _open(pid)
        base = (store.PROJECTS_DIR / pid).resolve()
        out = []
        for item in _serialise(pid, record, pm):
            path = _image_file(record, pm, item["index"]).resolve()
            item["image"] = path if base in path.parents and path.is_file() else None
            out.append(item)
        return out


# -- editing ----------------------------------------------------------------

def _slide(pm: ProjectManager, index: int):
    count = len(pm.state.slides)
    if isinstance(index, bool) or not isinstance(index, int) or not 0 <= index < count:
        raise ValueError(f"No slide with index {index!r}: the project has {count} slides (0 to {count - 1}).")
    return pm.state.slides[index]


def _notes(value) -> str:
    if not isinstance(value, str):
        raise ValueError("Speaker notes must be text.")
    if len(value) > MAX_NOTES_CHARS:
        raise ValueError(f"Speaker notes are limited to {MAX_NOTES_CHARS:,} characters.")
    return value


def _pause(value):
    """A pause override in seconds (0 to 30), or None for the studio default."""
    if value is None:
        return None
    if isinstance(value, bool):
        raise ValueError("The pause after a slide must be a number of seconds between 0 and 30, or empty for the default.")
    try:
        seconds = float(value)
    except (TypeError, ValueError):
        raise ValueError("The pause after a slide must be a number of seconds between 0 and 30, or empty for the default.") from None
    if not math.isfinite(seconds) or not 0 <= seconds <= MAX_PAUSE_SECONDS:
        raise ValueError("The pause after a slide must be a number of seconds between 0 and 30, or empty for the default.")
    return round(seconds, 3)


def _voice(value, provider) -> str | None:
    """A per-slide voice id checked against ``provider`` (None = the studio's
    configured provider) the way generate checks its voice, so an id from the
    other provider is refused now rather than silently dropped at render
    time (``core.tts_provider.effective_voice``). None/"" clears the override."""
    if value is None:
        return None
    if not isinstance(value, str):
        raise ValueError("The voice override must be a voice id.")
    if not value.strip():
        return None
    return studio_settings.check_voice_for_provider(value, studio_settings.resolve_provider(provider))


def _text_or_none(value, what: str, limit: int = 2000):
    if value is None:
        return None
    if not isinstance(value, str):
        raise ValueError(f"{what} must be text.")
    if len(value) > limit:
        raise ValueError(f"{what} is limited to {limit:,} characters.")
    return value.strip() or None


def update_slide(
    pid: str, index: int, *,
    speaker_notes=_UNSET, voice_override=_UNSET, pause_override=_UNSET, alt_text=_UNSET,
    ai_enhanced=_UNSET, provider: str | None = None,
) -> dict:
    """Change one slide. A keyword left out is left alone; ``None`` clears an
    override (voice, pause, alt text) - notes are text only. Notes go through
    ``update_slide_notes`` (undo history, ``needs_regeneration``).
    ``ai_enhanced`` (a bool) is the "an AI wrote this text" flag the AI
    assistant sets beside the notes it writes. Raises ``ValueError`` for a bad
    index or value; nothing is written then."""
    with project_lock(pid):
        record, pm = _open(pid)
        slide = _slide(pm, index)
        # Validate everything before the first write.
        notes = _notes(speaker_notes) if speaker_notes is not _UNSET and speaker_notes is not None else _UNSET
        voice = _voice(voice_override, provider) if voice_override is not _UNSET else _UNSET
        pause = _pause(pause_override) if pause_override is not _UNSET else _UNSET
        alt = _text_or_none(alt_text, "The alt text") if alt_text is not _UNSET else _UNSET
        if ai_enhanced is not _UNSET and not isinstance(ai_enhanced, bool):
            raise ValueError("ai_enhanced must be true or false.")

        # The flag is set before the other writes, which save the file anyway;
        # it gets a write of its own only when nothing else saved.
        flag_changed = ai_enhanced is not _UNSET and slide.ai_enhanced != ai_enhanced
        if flag_changed:
            slide.ai_enhanced = ai_enhanced
        saved = False
        if notes is not _UNSET:
            saved = slide.speaker_notes != notes
            pm.update_slide_notes(index, notes)  # saves only when the text changed
        if pause is not _UNSET:
            pm.update_slide_pause(index, pause)
            saved = True
        if voice is not _UNSET and slide.voice_override != voice:
            slide.voice_override = voice
            slide.needs_regeneration = True
            pm.save()
            saved = True
        if alt is not _UNSET and slide.alt_text != alt:
            slide.alt_text = alt
            pm.save()
            saved = True
        if flag_changed and not saved:
            pm.save()
        return _serialise(pid, record, pm)[index]


def undo_slide(pid: str, index: int) -> dict | None:
    """Restore the notes before the last edit. None when there is nothing to undo."""
    with project_lock(pid):
        record, pm = _open(pid)
        _slide(pm, index)
        if pm.undo_slide_notes(index) is None:
            return None
        return _serialise(pid, record, pm)[index]


def reset_slide(pid: str, index: int) -> dict:
    """Back to the deck's own notes ("" for a PDF); the current text goes onto
    the undo history like any edit. The deck's text is nobody's AI rewrite,
    so the ``ai_enhanced`` flag comes off with it."""
    with project_lock(pid):
        record, pm = _open(pid)
        slide = _slide(pm, index)
        info = _deck_info(pid, record)
        original = info[index].speaker_notes if index < len(info) else ""
        flagged = slide.ai_enhanced
        slide.ai_enhanced = False
        changed = slide.speaker_notes != original
        pm.update_slide_notes(index, original)  # saves only when the text changed
        if flagged and not changed:
            pm.save()
        return _serialise(pid, record, pm)[index]


def bulk_update(pid: str, items: list[dict]) -> list[dict]:
    """Save the notes of several slides (``[{index, speaker_notes}]``) at once.
    Every item is checked before the first write, so a bad one refuses the
    whole batch. Returns every slide afterwards."""
    with project_lock(pid):
        record, pm = _open(pid)
        checked = []
        for item in items:
            index = item.get("index")
            _slide(pm, index)
            checked.append((index, _notes(item.get("speaker_notes"))))
        for index, notes in checked:
            pm.update_slide_notes(index, notes)
        return _serialise(pid, record, pm)


# -- images ------------------------------------------------------------------

def record_images_source(pid: str, source: str) -> None:
    """Note where a project's slide images came from (``powerpoint``,
    ``pillow`` or ``pdf``) and when, onto the outer record as it is now. The
    generate route calls this after a render that exported slides; the
    editor's notice and (later) the vision features read it."""
    if source not in IMAGE_SOURCES:
        raise ValueError(f"Unknown image source '{source}'.")
    current = store.get_project(pid)
    if current is None:
        return
    current["images_source"] = source
    current["images_rendered_at"] = datetime.now(timezone.utc).isoformat()
    store.save_project(current)


def ensure_images(pid: str, progress=None, force: bool = False) -> dict:
    """Render every slide's image: a deck through ``PPTXExporter`` (PowerPoint,
    else the title-only Pillow fallback) behind ``SLIDE_EXPORT_LOCK``, a PDF
    through its page renderer. Idempotent: ``{"cached": True}`` when every
    image is already there - checked again once the export lock is held, in
    case another export of the same deck just finished - unless ``force``
    asks for a fresh export (previews rendered before ``images_source``
    existed have no recorded source, and vision will not use them until one
    is). Records ``images_source`` and ``images_rendered_at``. Meant to run
    in a job."""

    def report(fraction: float, message: str) -> None:
        if progress:
            progress(fraction, message)

    def cached_or_none():
        if force:
            return None
        record, pm = _open(pid)
        if _images_ready(record, pm):
            return {"cached": True, "images_source": record.get("images_source")}
        return None

    with project_lock(pid):
        cached = cached_or_none()
        if cached:
            return cached
        record, pm = _open(pid)
        source = _source(pid, record)
        kind = record["kind"]
        if kind == "pdf":
            # Fast, no COM, and it rewrites project.json (image paths), so it
            # stays under the project lock like any other write.
            report(0.1, "Rendering the PDF pages…")
            from services.file_item import FileItem

            item = FileItem(source, projects_base=store.PROJECTS_DIR / pid)
            if not item.load_pdf():
                raise RuntimeError("The PDF pages could not be rendered.")
            record_images_source(pid, "pdf")
            report(1.0, "Slide previews ready")
            return {"cached": False, "images_source": "pdf", "slides": item.slide_count}

    # The export itself runs outside the project lock - a minute or two for a
    # large deck, and notes saves must not queue behind it - and inside the
    # process-wide export lock, since PowerPoint takes one export at a time.
    with slide_export_lock(progress, 0.05):
        with project_lock(pid):
            cached = cached_or_none()
        if cached:
            return cached
        report(0.1, "Exporting slides…")
        from core.pptx_exporter import PPTXExporter

        exporter = PPTXExporter(source, pm.images_dir)
        exported = exporter.export_slides_as_images(
            progress_callback=lambda done, total: report(
                0.1 + 0.8 * done / max(total, 1), f"Exporting slide {done}/{total}…",
            ),
        )
        images_source = "powerpoint" if exporter.backend == "powerpoint" else "pillow"

    with project_lock(pid):
        # Re-read: notes may have been saved while the export ran.
        record, pm = _open(pid)
        for exp in exported:
            pm.update_slide_image(exp.index, exp.image_path)
            if exp.video_path:
                pm.update_slide_video(exp.index, exp.video_path, exp.has_animation)
        record_images_source(pid, images_source)
    report(1.0, "Slide previews ready")
    return {"cached": False, "images_source": images_source, "slides": len(exported)}


def image_path(pid: str, index: int) -> Path | None:
    """The rendered image of one slide, derived from the project directory and
    the index and proven to lie under the project; None when not rendered or
    out of range."""
    with project_lock(pid):
        record, pm = _open(pid)
        if not isinstance(index, int) or not 0 <= index < len(pm.state.slides):
            return None
        base = (store.PROJECTS_DIR / pid).resolve()
        path = _image_file(record, pm, index).resolve()
        if base not in path.parents or not path.is_file():
            return None
        return path


# -- a still becomes a slide (T3) ------------------------------------------------

#: The size a new slide's image is rendered at when the deck has no rendered
#: slide yet to match: what the PowerPoint export produces.
DEFAULT_SLIDE_IMAGE_SIZE = (1920, 1080)
#: The background a still is letterboxed onto when the deck has no rendered
#: slide to take its background from - white, a new deck's own default.
LETTERBOX_BACKGROUND = (255, 255, 255)
#: How deep a band of the slide's edge is read for its background: this share
#: of the shorter side, at least one pixel.
BORDER_SHARE = 0.01


def letterbox(image, size: tuple[int, int], background=LETTERBOX_BACKGROUND):
    """``image`` fitted inside ``size`` with its aspect kept and the rest
    filled with ``background`` (a transparent still is flattened onto it).
    A still already of that aspect fills the frame; one wider or taller gets
    bars - never a stretch, which is what the render would otherwise do to
    an image whose aspect differs from the slide's."""
    from PIL import Image

    target_w, target_h = size
    src = image.convert("RGBA")
    scale = min(target_w / src.width, target_h / src.height)
    w, h = max(1, round(src.width * scale)), max(1, round(src.height * scale))
    fitted = src.resize((w, h), Image.LANCZOS) if (w, h) != src.size else src
    canvas = Image.new("RGB", (target_w, target_h), background)
    x, y = (target_w - w) // 2, (target_h - h) // 2
    canvas.paste(fitted, (x, y), fitted)
    return canvas


def border_colour(image) -> tuple[int, int, int]:
    """The colour of an image's edge: the per-channel MEDIAN of a band round
    its border (``BORDER_SHARE`` of the shorter side, at least one pixel) -
    a slide's background where its content does not reach, and the median so
    a logo or a footer touching one edge does not tint it."""
    import numpy as np

    a = np.asarray(image.convert("RGB"))
    h, w = a.shape[:2]
    d = max(1, int(min(w, h) * BORDER_SHARE))
    ring = np.concatenate([
        a[:d].reshape(-1, 3), a[-d:].reshape(-1, 3),
        a[d:-d, :d].reshape(-1, 3), a[d:-d, -d:].reshape(-1, 3),
    ])
    return tuple(int(v) for v in np.median(ring, axis=0))


def _deck_look(record: dict, pm: ProjectManager) -> tuple[tuple[int, int], tuple[int, int, int]]:
    """The size and background a new slide's image takes from the deck: the
    first rendered slide image's pixel size and its border colour
    (``border_colour``); 1920x1080 on white when no slide is rendered."""
    from PIL import Image

    for slide in pm.state.slides:
        p = _image_file(record, pm, slide.index)
        if p.is_file():
            try:
                with Image.open(p) as img:
                    return img.size, border_colour(img)
            except Exception:
                continue
    return DEFAULT_SLIDE_IMAGE_SIZE, LETTERBOX_BACKGROUND


def _append_picture_slide(pptx_path: Path, image: Path) -> None:
    """Add a slide at the end of the deck holding ``image`` fitted to the
    slide (aspect kept, centred), on the blank layout. The deck is rewritten
    atomically: saved beside itself, then moved over."""
    from pptx import Presentation

    from utils.helpers import replace_with_retry

    prs = Presentation(str(pptx_path))
    layouts = list(prs.slide_layouts)
    layout = layouts[6] if len(layouts) > 6 else layouts[-1]
    slide = prs.slides.add_slide(layout)
    # The blank layout can still carry placeholders; a picture slide wants none.
    for shape in list(slide.placeholders):
        shape._element.getparent().remove(shape._element)
    from PIL import Image

    with Image.open(image) as img:
        iw, ih = img.size
    sw, sh = prs.slide_width, prs.slide_height
    scale = min(sw / iw, sh / ih)
    w, h = int(iw * scale), int(ih * scale)
    slide.shapes.add_picture(str(image), int((sw - w) / 2), int((sh - h) / 2), w, h)
    tmp = pptx_path.with_name(f"{pptx_path.stem}.{uuid.uuid4().hex[:8]}.tmp.pptx")
    prs.save(str(tmp))
    replace_with_retry(tmp, pptx_path)


def append_image_slide(pid: str, image: Path) -> dict:
    """Make ``image`` (a capture's still) the deck's new LAST slide: the
    slide's rendered image is the still letterboxed to the deck's rendered
    size, on the deck's own background - both read from its first rendered
    slide (``_deck_look``; 1920x1080 on white when none is rendered) - so the
    render never stretches it; the deck's ``.pptx`` gains a picture slide so a later export from
    PowerPoint agrees; the inner project gains the slide with empty notes.
    Decks only - a PDF's pages are fixed. The caller holds the project idle
    (``jobs.require_idle``); this takes the project lock. Returns
    ``{"index", "slide_count"}``."""
    from PIL import Image

    with project_lock(pid):
        record, pm = _open(pid)
        if record["kind"] != "deck":
            raise ValueError("Only a deck can take a new slide; a PDF's pages are fixed.")
        image = Path(image)
        if not image.is_file():
            raise ValueError("The capture's image is missing.")
        index = len(pm.state.slides)
        size, background = _deck_look(record, pm)
        target = _image_file(record, pm, index)
        target.parent.mkdir(parents=True, exist_ok=True)
        with Image.open(image) as img:
            letterbox(img, size, background).save(target)
        _append_picture_slide(_source(pid, record), image)
        pm.state.slides.append(SlideRenderState(index=index, speaker_notes=""))
        if pm.state.slide_order is not None:
            pm.state.slide_order.append(index)
        pm.update_slide_image(index, target)
        _reconcile(pid, record, pm)
        return {"index": index, "slide_count": index + 1}


# -- export -------------------------------------------------------------------

def export_notes_pptx(pid: str) -> Path:
    """A copy of the deck with every slide's notes replaced by the edited ones,
    at ``<project>/exports/<name>-notes.pptx`` (index-aligned, all slides).
    Decks only: a PDF has no deck to write into."""
    with project_lock(pid):
        record, pm = _open(pid)
        if record["kind"] != "deck":
            raise ValueError("Only a deck can be exported as a .pptx with notes.")
        from core.pptx_reader import PPTXReader

        notes = [slide.speaker_notes or "" for slide in pm.state.slides]
        base = (store.PROJECTS_DIR / pid).resolve()
        target = (base / "exports" / f"{record['name']}-notes.pptx").resolve()
        # The name comes from the record: a tampered one must not write outside the project.
        if base not in target.parents:
            raise ValueError("The project name is not a valid file name.")
        return PPTXReader.export_with_updated_notes(_source(pid, record), notes, target)
